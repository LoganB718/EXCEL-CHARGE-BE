# -*- coding: utf-8 -*-
"""Génère le classeur « Plan de charge BE » : onglets Charge, Data, Dashboard et Calculs.

Usage : python scripts/build_dashboard.py [chemin_sortie.xlsx]
"""
import re
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.data_source import NumFmt, StrRef
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.marker import Marker
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

OUT = sys.argv[1] if len(sys.argv) > 1 else "Plan_de_charge_BE_Dashboard.xlsx"

FONT = "Arial"
NAVY = "1F2A44"
INK = "1F2937"
GREY = "6B7280"
MUTED = "9CA3AF"
LINE = "D6DCE4"
BG = "EEF1F5"
WHITE = "FFFFFF"
ACCENT = "2A78D6"
AMBER = "EDA100"
INPUT_FILL = "EAF2FD"
INPUT_INK = "1F4E9A"
RED_T, GREEN_T = "B42318", "1E7B34"   # textes d'alerte (surcharge) et de confort
# Palette catégorielle validée (ordre fixe) + gris pour « Autres projets »
SLOT_COLORS = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "6250D6", "E34948", "9CA3AF"]
N_SLOTS = len(SLOT_COLORS)

FIRST, LAST = 4, 503            # lignes de données de l'onglet Charge (500 affaires max)
N_MONTHS = 12                   # ne pas augmenter : les blocs de l'onglet Calculs sont placés pour 12 mois
assert N_MONTHS <= 12
N_TABLE = 100                   # lignes du tableau des affaires sur le Dashboard

NA_LABEL = "(non affecté)"          # ressource d'une ligne sans dessinateur ni ressource externe
PERIM_ALL = "Tous les projets"
PERIM_RUN = "Projets en cours"
REP_SMOOTH = "Lissée jusqu'à la date de besoin"
REP_START = "Tout sur le mois de début"


def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def side(color=LINE, style="thin"):
    return Side(style=style, color=color)


def lit(text):
    """Texte en constante de formule, découpé en morceaux de 250 caractères (Excel limite chaque constante à 255)."""
    parts = [text[i:i + 250] for i in range(0, len(text), 250)] or [""]
    return "&".join('"' + p.replace('"', '""') + '"' for p in parts)


def help_formula(toggle, content):
    """Aide affichée seulement si l'état d'aide de l'onglet (nom défini `toggle` : case « Afficher l'aide » ou,
    sur Dashboard et Dessinateurs, cellule de Calculs basculée par le bouton) vaut Oui.
    `content` : texte simple, ou formule commençant par « = »."""
    expr = content[1:] if content.startswith("=") else lit(content)
    return f'=IF({toggle}="Oui",{expr},"")'


HELP_CHOICES = '"Oui,Non"'
MOIS = ("janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre", "novembre",
        "décembre")


def mois_fr(ref):
    """Nom du mois en français (TEXTE dépend de la langue d'Excel)."""
    return f'CHOOSE(MONTH({ref}),' + ",".join(f'"{m}"' for m in MOIS) + ")"


def rng(col, r1=FIRST, r2=LAST, sheet=None):
    ref = f"${col}${r1}:${col}${r2}"
    return f"{sheet}!{ref}" if sheet else ref


wb = Workbook()

# ---------------------------------------------------------------- Data
ws_data = wb.active
ws_data.title = "Data"
ws_charge = wb.create_sheet("Charge", 0)
ws_dash = wb.create_sheet("Dashboard")
ws_dess = wb.create_sheet("Dessinateurs")
ws_calc = wb.create_sheet("Calculs")

LISTS = [  # (colonne, en-tête, nom de la liste, valeurs d'exemple)
    ("A", "Chefs de projet", "lst_ChefsProjet",
     ["Julie Martin", "Thomas Bernard", "Sophie Lefèvre", "Nicolas Moreau", "Camille Laurent"]),
    ("B", "Noms de projet", "lst_Projets",
     ["ATLAS", "BORÉAL", "CASSIOPÉE", "DELTA", "ORION", "PÉGASE", "VEGA", "TITAN"]),
    ("C", "Types", "lst_Types",
     ["Étude de conception", "Modification", "Mise à jour de plans", "Industrialisation",
      "Chiffrage / avant-projet"]),
    ("D", "Coordinateurs", "lst_Coordinateurs", ["Marc Dubois", "Isabelle Roux"]),
    ("E", "Dessinateurs projeteurs", "lst_Dessinateurs",
     ["Lucas Garnier", "Emma Faure", "Hugo Mercier", "Léa Chevalier", "Antoine Blanc"]),
    ("H", "Ressources externes (personnes)", "lst_Prestataires",
     ["Paul Girard", "Inès Morel", "Karim Haddad"]),
]
# Réglages propres à chaque ressource (même ligne que son nom) ; vide = valeur par défaut du Dashboard
DESS_COLS = [  # (colonne, en-tête, valeurs d'exemple alignées sur la liste correspondante)
    ("F", "Capacité mensuelle (h)\nvide = défaut", [None, None, None, 96, None]),
    ("G", "Charge fixe hors projet (h/mois)\nvide = défaut", [None, None, None, None, None]),
    ("I", "Capacité mensuelle (h)\nvide = défaut", [None, None, None]),
    ("J", "Charge fixe hors projet (h/mois)\nvide = défaut", [None, None, None]),
]
DATA_LAST = 203

ws_data["A1"] = "LISTES DE RÉFÉRENCE"
ws_data["A1"].font = font(14, True, NAVY)
ws_data["A2"] = help_formula("aide_Data", "Saisir une valeur par cellule, sans ligne vide, sous chaque en-tête : les listes déroulantes "
                 "de l'onglet Charge se mettent à jour automatiquement (200 valeurs max par liste). "
                 "Colonnes F-G et I-J : capacité et charge fixe propres à un dessinateur ou à une ressource "
                 "externe (vide = valeurs par défaut réglées en haut du Dashboard). Colonne K (facultative) : jours "
                 "non travaillés communs à tout le bureau (fériés, fermeture), retirés du lissage comme les week-ends ; "
                 "les congés individuels ne se saisissent pas.  Les valeurs ci-dessous sont des exemples à remplacer "
                 "(ex. Léa Chevalier à 96 h = temps partiel 80 %).")
ws_data["A2"].font = font(9, color=GREY, italic=True)
ws_data["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws_data.merge_cells("A2:K2")
ws_data.row_dimensions[2].height = 26
ws_data.row_dimensions[3].height = 42
for col, title, values in DESS_COLS:
    ws_data.column_dimensions[col].width = 17
    h = ws_data[f"{col}3"]
    h.value = title
    h.font = font(9, True, WHITE)
    h.fill = fill("2F5C8F")
    h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in range(4, DATA_LAST + 1):
        c = ws_data[f"{col}{r}"]
        c.font = font(10, color=INPUT_INK)
        c.fill = fill("F5F8FD")
        c.number_format = 'General" h"'
        c.alignment = Alignment(horizontal="center")
        c.border = Border(bottom=side("E5E7EB"), left=side("E5E7EB"), right=side("E5E7EB"))
    for j, v in enumerate(values):
        if v is not None:
            ws_data[f"{col}{4 + j}"] = v
    is_cap = col in ("F", "I")
    dv = DataValidation(type="decimal", operator="between", formula1="1" if is_cap else "0", formula2="744",
                        allow_blank=True, showErrorMessage=True, errorTitle="Heures par mois",
                        error="Saisir un nombre d'heures par mois (ou laisser vide pour la valeur par défaut)."
                              + (" La capacité doit être d'au moins 1 h." if is_cap else ""))
    dv.add(f"{col}4:{col}{DATA_LAST}")
    ws_data.add_data_validation(dv)
for col, title, name, values in LISTS:
    ws_data.column_dimensions[col].width = 30
    h = ws_data[f"{col}3"]
    h.value = title
    h.font = font(10, True, WHITE)
    h.fill = fill(NAVY)
    h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in range(4, DATA_LAST + 1):
        c = ws_data[f"{col}{r}"]
        c.font = font(10)
        c.border = Border(bottom=side("E5E7EB"), left=side("E5E7EB"), right=side("E5E7EB"))
    for j, v in enumerate(values):
        ws_data[f"{col}{4 + j}"] = v
    wb.defined_names[name] = DefinedName(
        name, attr_text=f"OFFSET(Data!${col}$3,1,0,MAX(1,COUNTA(Data!${col}$4:${col}${DATA_LAST})),1)")
HOLIDAYS = f"Data!$K$4:$K${DATA_LAST}"   # jours non travaillés communs (dates), retirés du lissage
ws_data.column_dimensions["K"].width = 22
h = ws_data["K3"]
h.value = "Jours non travaillés\n(fériés, fermetures)\nfacultatif"
h.font = font(9, True, WHITE)
h.fill = fill("6B7280")
h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for r in range(4, DATA_LAST + 1):
    c = ws_data[f"K{r}"]
    c.font = font(10)
    c.number_format = "DD/MM/YYYY"
    c.alignment = Alignment(horizontal="center")
    c.border = Border(bottom=side("E5E7EB"), left=side("E5E7EB"), right=side("E5E7EB"))
dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True, showErrorMessage=True,
                    errorTitle="Date attendue", error="Saisir uniquement des dates (jj/mm/aaaa), une par cellule.")
dv.add(f"K4:K{DATA_LAST}")
ws_data.add_data_validation(dv)
ws_data.freeze_panes = "A4"
ws_data.sheet_properties.tabColor = "6B7280"
ws_data.print_area = f"A2:K{DATA_LAST}"


MACRO_WARNING = ("ATTENTION : macros désactivées, les boutons et le masquage des lignes vides ne fonctionnent pas. "
                 "Fermer le fichier, clic droit > Propriétés > cocher « Débloquer », rouvrir puis « Activer le "
                 "contenu ».")


def macro_warning(sheet, cells):
    """Voyant rouge hors zone d'impression : visible tant que les macros ne tournent pas (nom macros_actives, mis à
    Oui par la macro d'ouverture du classeur, remis à Non avant chaque enregistrement)."""
    first = cells.split(":")[0]
    sheet.merge_cells(cells)
    c = sheet[first]
    c.value = f'=IF(macros_actives="Oui","",{lit(MACRO_WARNING)})'
    c.font = font(9, True, RED_T)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


ws_data.row_dimensions[1].height = 26
macro_warning(ws_data, "D1:H1")

# ---------------------------------------------------------------- Charge
CHARGE_COLS = [
    # (lettre, en-tête, largeur, liste, format)
    ("A", "Chef de projet", 20, "lst_ChefsProjet", None),
    ("B", "Nom du projet", 18, "lst_Projets", None),
    ("C", "Sous-projet", 22, None, None),
    ("D", "Produit", 22, None, None),
    ("E", "Type", 20, "lst_Types", None),
    ("F", "Date de besoin", 13, None, "DD/MM/YYYY"),
    ("G", "Description du projet", 38, None, None),
    ("H", "Contraintes, spécifications, délais, budget", 34, None, None),
    ("I", "Autres informations pertinentes / suggestions d'optimisation des ressources", 34, None, None),
    ("J", "Coordinateur", 18, "lst_Coordinateurs", None),
    ("K", "Dessinateur projeteur (interne)", 20, "lst_Dessinateurs", None),
    ("L", "OU ressource externe", 20, "lst_Prestataires", None),
    ("M", "Heures de cette ressource (h)", 13, None, None),
    ("N", "Date de début planifiée", 13, None, "DD/MM/YYYY"),
    ("O", "Taux d'avancement", 13, None, "0%"),
]

# Données d'exemple : une ligne = une ressource (dessinateur K OU ressource externe L) avec SES heures en M
_ATL1 = ("Julie Martin", "ATLAS", "Lot 1 – Structure", "Châssis principal", "Étude de conception", date(2027, 3, 31),
         "Conception du châssis mécano-soudé de la nouvelle gamme", "Masse < 450 kg ; budget étude 60 k€",
         "Mutualiser les calculs RDM avec le lot 2", "Marc Dubois")
_CAS2 = ("Sophie Lefèvre", "CASSIOPÉE", "Rétrofit ligne 2", "Poste de contrôle", "Étude de conception",
         date(2027, 2, 26), "Nouveau poste de contrôle vision", None, None, "Marc Dubois")
_DEL1 = ("Sophie Lefèvre", "DELTA", "Lot unique", "Bâti machine", "Industrialisation", date(2027, 6, 30),
         "Industrialisation du bâti", "Cible coût -10 %", None, "Isabelle Roux")
_PEG1 = ("Camille Laurent", "PÉGASE", "Lot 1", "Ensemble complet", "Étude de conception", date(2027, 9, 30),
         "Étude complète si commande", "Sous réserve de commande", "Renfort externe prévu", "Marc Dubois")
_DEL2 = ("Thomas Bernard", "DELTA", "Lot unique", "Habillage", "Étude de conception", date(2027, 8, 27),
         "Habillage design", None, None, "Isabelle Roux")
SAMPLE = [
    _ATL1 + ("Lucas Garnier", None, 300, date(2026, 11, 2), 0.15),
    _ATL1 + (None, "Paul Girard", 120, date(2026, 11, 2), 0),
    ("Julie Martin", "ATLAS", "Lot 2 – Carénage", "Capots latéraux", "Étude de conception", date(2027, 5, 28),
     "Dessin des capots et de leurs fixations", "Matière ABS, outillage existant", None, "Marc Dubois",
     "Emma Faure", None, 260, date(2027, 1, 11), 0),
    ("Julie Martin", "ATLAS", "Lot 3 – Intégration", "Faisceau électrique", "Industrialisation", date(2027, 7, 30),
     "Routage des faisceaux et supports", "Normes NF C 15-100", None, "Marc Dubois",
     None, "Inès Morel", 180, date(2027, 4, 5), 0),
    ("Thomas Bernard", "BORÉAL", "Phase APD", "Groupe froid", "Étude de conception", date(2027, 1, 29),
     "Avant-projet détaillé du groupe froid", "Délai client ferme", None, "Isabelle Roux", "Léa Chevalier",
     None, 310, date(2026, 10, 5), 0.35),
    ("Thomas Bernard", "BORÉAL", "Phase EXE", "Groupe froid", "Mise à jour de plans", date(2027, 4, 30),
     "Plans d'exécution", "Budget 25 k€", None, "Isabelle Roux", "Léa Chevalier", None, 220,
     date(2027, 2, 1), 0),
    ("Thomas Bernard", "BORÉAL", "Phase EXE", "Armoire de commande", "Modification", date(2027, 3, 15),
     "Reprise de l'armoire existante", None, None, "Isabelle Roux", None, "Karim Haddad", 90,
     date(2027, 1, 18), 0),
    ("Sophie Lefèvre", "CASSIOPÉE", "Rétrofit ligne 2", "Convoyeur", "Modification", date(2026, 12, 18),
     "Adaptation du convoyeur au nouveau format", "Arrêt de ligne limité à 2 semaines",
     "Réutiliser les plans de la ligne 1", "Marc Dubois", "Hugo Mercier", None, 150, date(2026, 10, 12), 0.5),
    _CAS2 + ("Emma Faure", None, 160, date(2026, 12, 1), 0.05),
    _CAS2 + (None, "Paul Girard", 80, date(2026, 12, 1), 0),
    _DEL1 + ("Antoine Blanc", None, 240, date(2027, 2, 15), 0),
    _DEL1 + ("Hugo Mercier", None, 120, date(2027, 2, 15), 0),
    ("Nicolas Moreau", "ORION", "Tranche 1", "Support moteur", "Étude de conception", date(2026, 11, 27),
     "Conception du support moteur", None, None, "Isabelle Roux", "Lucas Garnier", None, 120,
     date(2026, 9, 14), 0.8),
    ("Nicolas Moreau", "ORION", "Tranche 1", "Carter de protection", "Mise à jour de plans", date(2026, 10, 9),
     "Mise à jour indice C", None, None, "Isabelle Roux", "Lucas Garnier", None, 40, date(2026, 9, 21), 1),
    ("Nicolas Moreau", "ORION", "Tranche 2", "Support moteur", "Modification", date(2027, 5, 14),
     "Évolution pour moteur 30 kW", None, None, "Isabelle Roux", "Emma Faure", None, 160, date(2027, 3, 1), 0),
    ("Nicolas Moreau", "PÉGASE", "Avant-projet", "Ensemble complet", "Chiffrage / avant-projet",
     date(2026, 12, 11), "Chiffrage et esquisses pour l'appel d'offres", "Remise d'offre le 15/12", None,
     "Marc Dubois", "Léa Chevalier", None, 70, date(2026, 11, 16), 0.1),
    _PEG1 + ("Hugo Mercier", None, 300, date(2027, 3, 15), 0),
    _PEG1 + (None, "Karim Haddad", 180, date(2027, 3, 15), 0),
    ("Camille Laurent", "VEGA", "Mise en conformité", "Garde-corps", "Modification", date(2027, 1, 15),
     "Mise en conformité des garde-corps", "Norme NF E 85-015", None, "Isabelle Roux", "Antoine Blanc", None,
     85, date(2026, 11, 23), 0),
    ("Camille Laurent", "VEGA", "Mise en conformité", "Passerelle", "Mise à jour de plans", date(2027, 3, 31),
     "Plans de recollement", None, None, "Isabelle Roux", "Antoine Blanc", None, 60, date(2027, 2, 8), 0),
    ("Camille Laurent", "VEGA", "Mise en conformité", "Escalier d'accès", "Modification", date(2027, 4, 30),
     "Mise en conformité de l'escalier d'accès", None, "Ressource à définir", "Isabelle Roux", None, None, 50,
     date(2027, 3, 1), 0),
    ("Julie Martin", "TITAN", "Lot 1", "Poutre de levage", "Étude de conception", date(2026, 9, 30),
     "Note de calcul et plans", None, None, "Marc Dubois", "Lucas Garnier", None, 95, date(2026, 8, 24), 1),
    _DEL2 + ("Emma Faure", None, 120, date(2027, 5, 3), 0),
    _DEL2 + (None, "Inès Morel", 80, date(2027, 5, 3), 0),
]
assert all(len(row) == 15 for row in SAMPLE)

ws = ws_charge
ws["A1"] = "PLAN DE CHARGE — BUREAU D'ÉTUDES"
ws["A1"].font = font(14, True, NAVY)
ws["A2"] = help_formula("aide_Charge", "UNE LIGNE = UNE RESSOURCE sur un sous-projet : un dessinateur (K) OU une ressource externe (L), avec SES "
            "heures en M.  Pour mettre plusieurs personnes sur un sous-projet, recopier la ligne (colonnes A à J à "
            "l'identique) et changer la ressource et les heures.  Couleurs : rouge = dessinateur ET externe sur la "
            "même ligne (à corriger, seul K est compté) ; orange = aucune ressource (charge à affecter).  "
            "Insérer, supprimer ou trier des lignes est possible.  Les lignes 4 à "
            f"{FIRST + len(SAMPLE) - 1} sont des EXEMPLES FICTIFS à remplacer par vos données.")
ws["A2"].font = font(9, color=GREY, italic=True)
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A2:I2")
ws.row_dimensions[2].height = 28
ws.row_dimensions[1].height = 26
macro_warning(ws, "F1:I1")
ws.print_area = f"A2:O{LAST}"
ws.print_title_rows = "3:3"
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A3
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.row_dimensions[3].height = 48
thin = side("D9DEE5")
for col, title, width, _, fmt in CHARGE_COLS:
    ws.column_dimensions[col].width = width
    h = ws[f"{col}3"]
    h.value = title
    h.font = font(10, True, WHITE)
    h.fill = fill(NAVY if col <= "I" else "2F5C8F")
    h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    h.border = Border(left=side(WHITE), right=side(WHITE))
    wrap = col in "GHI"
    for r in range(FIRST, LAST + 1):
        c = ws[f"{col}{r}"]
        c.font = font(10)
        c.border = Border(bottom=thin, left=thin, right=thin)
        c.alignment = Alignment(vertical="top", wrap_text=wrap,
                                horizontal="center" if col in "FMNO" else None)
        if fmt:
            c.number_format = fmt

for i, row in enumerate(SAMPLE):
    r = FIRST + i
    for (col, *_), v in zip(CHARGE_COLS, row):
        if v is not None:
            ws[f"{col}{r}"] = v

ws.freeze_panes = "C4"
ws.auto_filter.ref = f"A3:O{LAST}"
ws.sheet_view.zoomScale = 90

for col, _, _, lst, _ in CHARGE_COLS:
    if lst:
        dv = DataValidation(type="list", formula1=lst, allow_blank=True, errorStyle="warning",
                            showErrorMessage=True, errorTitle="Valeur hors liste",
                            error="Cette valeur n'est pas dans la liste de l'onglet Data. Continuer quand même ?")
        dv.add(f"{col}{FIRST}:{col}{LAST}")
        ws.add_data_validation(dv)
for col in "FN":
    dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                        showErrorMessage=True, errorTitle="Date attendue", error="Saisir une date (jj/mm/aaaa).")
    dv.add(f"{col}{FIRST}:{col}{LAST}")
    ws.add_data_validation(dv)
dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                    showErrorMessage=True, errorTitle="Nombre d'heures",
                    error="Saisir un nombre d'heures positif (heures de la personne de cette ligne).")
dv.add(f"M{FIRST}:M{LAST}")
ws.add_data_validation(dv)
dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                    showErrorMessage=True, errorTitle="Taux d'avancement",
                    error="Saisir un pourcentage entre 0 % et 100 %.")
dv.add(f"O{FIRST}:O{LAST}")
ws.add_data_validation(dv)

# Contrôles visuels : charge ou date de début manquante, début après la date de besoin
miss = PatternFill("solid", start_color="FDE2E1", end_color="FDE2E1", bgColor="FDE2E1")
warn = PatternFill("solid", start_color="FFF1CC", end_color="FFF1CC", bgColor="FFF1CC")
ws.conditional_formatting.add(f"M{FIRST}:N{LAST}",
                              FormulaRule(formula=[f'AND($B{FIRST}<>"",M{FIRST}="")'], fill=miss))
ws.conditional_formatting.add(f"N{FIRST}:N{LAST}",
                              FormulaRule(formula=[f"AND(ISNUMBER($N{FIRST}),ISNUMBER($F{FIRST}),"
                                                   f"$N{FIRST}>$F{FIRST})"], fill=warn))
ws.conditional_formatting.add(f"O{FIRST}:O{LAST}",
                              DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                                          color="86B6EF", showValue=True))
# une ressource par ligne : rouge si dessinateur ET ressource externe, orange si aucune ressource
ws.conditional_formatting.add(f"K{FIRST}:L{LAST}",
                              FormulaRule(formula=[f'AND(TRIM($K{FIRST})<>"",TRIM($L{FIRST})<>"")'], stopIfTrue=True,
                                          fill=miss, font=Font(color=RED_T, bold=True)))
ws.conditional_formatting.add(f"K{FIRST}:L{LAST}",
                              FormulaRule(formula=[f'AND(TRIM($B{FIRST})<>"",TRIM($K{FIRST})="",TRIM($L{FIRST})="")'],
                                          fill=warn))

# ---------------------------------------------------------------- Dashboard : paramètres (nommés)
GRID0 = 2          # colonne Excel de la 1re colonne de la grille (B)
NG = 59            # nombre de colonnes de la grille
MARGIN_W = 2       # largeur des colonnes de marge A et BI (unité Excel) = 14 pixels
MARGIN_H = 10.5    # hauteur des lignes de marge (points) = 14 pixels : mêmes marges en haut, en bas et sur les côtés
LOGO = (51, 57)    # emplacement du logo de l'entreprise : cases fusionnées BA2:BG2
# boutons des macros : hors zone d'impression, à droite de la colonne BI (placés par build_xlsm.py en BL2:BM4 et
# BL6:BM7) ; l'aide de ces onglets s'affiche ou se masque par le bouton (état dans Calculs, noms aide_...).
# Sous les boutons : voyant rouge tant que les macros ne tournent pas.
MACRO_WARNING_CELLS = "BL9:BQ14"


def gc(i):
    return GRID0 + i


def gl(i):
    return L(gc(i))


P_ROW = 6
PARAMS = {  # nom: (col début input, col fin input)
    "p_Perimetre": (4, 10),
    "p_Repartition": (16, 28),
    "p_PremierMois": (35, 39),
    "p_HeuresETP": (45, 48),     # capacité mensuelle par défaut d'une personne
    "p_ChargeFixe": (54, 58),    # charge fixe hors projet par défaut (h/mois/personne)
}
for name, (c1, _) in PARAMS.items():
    wb.defined_names[name] = DefinedName(name, attr_text=f"Dashboard!${gl(c1)}${P_ROW}")

# ---------------------------------------------------------------- Calculs
ws = ws_calc
ws.sheet_properties.tabColor = "9CA3AF"
ws["A1"] = "CALCULS INTERMÉDIAIRES — alimentent le Dashboard, ne pas modifier"
ws["A1"].font = font(12, True, NAVY)

LINE_COLS = {
    "A": "Ligne (onglet Charge)", "B": "Affaire valide (projet renseigné)", "C": "Avancement (0-1)",
    "D": "Incluse dans le périmètre", "E": "Charge incluse (h)", "F": "Chef de projet", "G": "Projet",
    "H": "Type", "I": "Début planifié", "J": "Fin de répartition", "K": "Jours ouvrés de la période",
    "L": "Rang affaire (charge décroissante)",
    "M": "1re occurrence chef", "N": "Charge du chef", "O": "Rang chef",
    "P": "1re occurrence projet", "Q": "Charge du projet", "R": "Rang projet",
    "S": "1re occurrence type", "T": "Charge du type", "U": "Rang type",
    "V": "Groupe couleur (projet)",
}
MONTH_COL0 = 23  # W
MCOLS = [L(MONTH_COL0 + m) for m in range(N_MONTHS)]  # W..AH
SUM_COL = L(MONTH_COL0 + N_MONTHS)                    # AI

hdr_font = font(9, True, WHITE)
for col, title in list(LINE_COLS.items()) + [(SUM_COL, "Charge sur la période (h)")]:
    c = ws[f"{col}3"]
    c.value = title
    c.font = hdr_font
    c.fill = fill(NAVY)
    c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.column_dimensions[col].width = 12
ws.row_dimensions[3].height = 54
ws["V2"] = "Fin de mois →"
ws["V2"].font = font(9, color=GREY, italic=True)

R1, R2 = FIRST, LAST
for r in range(R1, R2 + 1):
    f = {
        "A": f"=ROW()",
        "B": f'=IF(TRIM(Charge!B{r})<>"",1,0)',
        "C": f"=IF(B{r}=1,MIN(1,MAX(0,N(Charge!O{r}))),0)",
        "D": f'=IF(B{r}=1,IF(OR(p_Perimetre="{PERIM_ALL}",C{r}<1),1,0),0)',
        "E": f"=IF(D{r}=1,MAX(0,N(Charge!M{r})),0)",
        "F": f'=IF(D{r}=1,IF(TRIM(Charge!A{r})="","(non renseigné)",TRIM(Charge!A{r})),"")',
        "G": f'=IF(D{r}=1,TRIM(Charge!B{r}),"")',
        "H": f'=IF(D{r}=1,IF(TRIM(Charge!E{r})="","(non renseigné)",TRIM(Charge!E{r})),"")',
        "I": f'=IF(AND(D{r}=1,ISNUMBER(Charge!N{r})),Charge!N{r},"")',
        "J": f'=IF(I{r}="","",IF(AND(p_Repartition="{REP_SMOOTH}",ISNUMBER(Charge!F{r})),'
             f'IF(Charge!F{r}>=I{r},Charge!F{r},I{r}),I{r}))',
        "K": f'=IF(I{r}="","",NETWORKDAYS(I{r},J{r},{HOLIDAYS}))',
        "L": f'=IF(D{r}=1,COUNTIFS($D${R1}:$D${R2},1,$E${R1}:$E${R2},">"&E{r})'
             f'+COUNTIFS($D${R1}:$D{r},1,$E${R1}:$E{r},E{r}),"")',
        "V": f'=IF(D{r}=1,MIN(INDEX($R${R1}:$R${R2},MATCH(G{r},$G${R1}:$G${R2},0)),'
             f'IF($AL$7>{N_SLOTS},{N_SLOTS},999)),"")',
    }
    # chefs (F), projets (G), types (H) : 1re occurrence, charge cumulée, rang décroissant
    for key, first, tot, rank in (("F", "M", "N", "O"), ("G", "P", "Q", "R"), ("H", "S", "T", "U")):
        f[first] = f'=IF({key}{r}="",0,IF(COUNTIF(${key}${R1}:{key}{r},{key}{r})=1,1,0))'
        f[tot] = f'=IF({first}{r}=1,ROUND(SUMIF(${key}${R1}:${key}${R2},{key}{r},$E${R1}:$E${R2}),2),"")'
        f[rank] = (f'=IF({first}{r}=1,COUNTIFS(${first}${R1}:${first}${R2},1,${tot}${R1}:${tot}${R2},">"&{tot}{r})'
                   f'+COUNTIFS(${first}${R1}:{first}{r},1,${tot}${R1}:{tot}{r},{tot}{r}),"")')
    for col, formula in f.items():
        ws[f"{col}{r}"] = formula
    for mc in MCOLS:
        # heures réparties au prorata des jours ouvrés ; période sans aucun jour ouvré : tout sur le mois de début
        ws[f"{mc}{r}"] = (f'=IF(OR($I{r}="",$E{r}=0),0,IF(OR($J{r}<{mc}$3,$I{r}>{mc}$2),0,'
                          f'IF($K{r}=0,IF(AND($I{r}>={mc}$3,$I{r}<={mc}$2),$E{r},0),'
                          f'$E{r}*NETWORKDAYS(MAX($I{r},{mc}$3),MIN($J{r},{mc}$2),{HOLIDAYS})/$K{r})))')
        ws[f"{mc}{r}"].number_format = "0.0"
    ws[f"{SUM_COL}{r}"] = f"=SUM({MCOLS[0]}{r}:{MCOLS[-1]}{r})"
    ws[f"{SUM_COL}{r}"].number_format = "0.0"
    for col in "IJ":
        ws[f"{col}{r}"].number_format = "DD/MM/YYYY"

# en-têtes de mois (début en ligne 3, fin en ligne 2)
for m, mc in enumerate(MCOLS):
    ws[f"{mc}3"] = "=$AL$14" if m == 0 else f"=EDATE({MCOLS[m - 1]}3,1)"
    ws[f"{mc}3"].number_format = "MMM YYYY"
    ws[f"{mc}2"] = f"=EOMONTH({mc}3,0)"
    ws[f"{mc}2"].number_format = "DD/MM/YYYY"
    ws[f"{mc}2"].font = font(9, color=GREY)

# Indicateurs (AK libellé / AL valeur)
K = {}
INDICATORS = [
    ("total", "Charge totale identifiée (h)", f"=SUM(E{R1}:E{R2})", "#,##0.0"),
    ("nbAff", "Nombre d'affaires (sous-projets regroupés)", f"=SUM(EO{R1}:EO{R2})", "0"),
    ("nbChefs", "Nombre de chefs de projet", f"=SUM(M{R1}:M{R2})", "0"),
    ("nbProj", "Nombre de projets", f"=SUM(P{R1}:P{R2})", "0"),
    ("nbTypes", "Nombre de types", f"=SUM(S{R1}:S{R2})", "0"),
    ("realise", "Charge réalisée (selon avancement)", f"=SUMPRODUCT(E{R1}:E{R2},C{R1}:C{R2})", "#,##0.0"),
    ("reste", "Reste à faire (h)", "=AL4-AL9", "#,##0.0"),
    ("avgSimple", "Avancement moyen (moyenne simple des affaires)",
     f"=IF(AL5=0,0,SUM(EU{R1}:EU{R2})/AL5)", "0.0%"),
    ("avgPond", "Avancement moyen pondéré par les heures (toutes les lignes des affaires du périmètre)",
     f"=IF(SUM(FH{R1}:FH{R2})=0,0,SUM(FI{R1}:FI{R2})/SUM(FH{R1}:FH{R2}))", "0.0%"),
    ("autoStart", "1er mois automatique (1er début planifié)",
     f"=IF(COUNT(I{R1}:I{R2})=0,DATE(YEAR(TODAY()),MONTH(TODAY()),1),"
     f"DATE(YEAR(MIN(I{R1}:I{R2})),MONTH(MIN(I{R1}:I{R2})),1))", "MMM YYYY"),
    ("start", "1er mois affiché",
     "=IF(ISNUMBER(p_PremierMois),DATE(YEAR(p_PremierMois),MONTH(p_PremierMois),1),AL13)", "MMM YYYY"),
    ("periode", "Charge répartie sur la période affichée (h)", f"=SUM(AP4:AP{3 + N_MONTHS})", "#,##0.0"),
    ("hors", "Charge hors période ou sans date de début (h)", "=MAX(0,ROUND(AL4-AL15,1))", "#,##0.0"),
    ("pic", "Pic de charge mensuel (h)", f"=MAX(AP4:AP{3 + N_MONTHS})", "#,##0.0"),
    ("picIdx", "Rang du mois de pic", f"=MATCH(AL17,AP4:AP{3 + N_MONTHS},0)", "0"),
    ("picMois", "Mois le plus chargé", f"=INDEX(AO4:AO{3 + N_MONTHS},AL18)", "MMM YYYY"),
    ("picActives", "Affaires actives sur le mois de pic", f"=INDEX(AQ4:AQ{3 + N_MONTHS},AL18)", "0"),
    ("picETP", "ETP nécessaires au pic (capacité projet = capacité - charge fixe)",
     "=IF(N(p_HeuresETP)-N(p_ChargeFixe)<=0,0,AL17/(p_HeuresETP-p_ChargeFixe))", "0.0"),
    ("totalETP", "Charge totale en ETP-mois", "=IF(N(p_HeuresETP)-N(p_ChargeFixe)<=0,0,AL4/(p_HeuresETP-p_ChargeFixe))",
     "0.0"),
    ("bigIdx", "Position de l'affaire au plus gros reste à faire",
     f'=IFERROR(MATCH(MAX(FJ{R1}:FJ{R2}),FJ{R1}:FJ{R2},0),"")', "0"),
    ("nbTermines", "Lignes terminées (avancement 100 %)", f"=COUNTIFS(B{R1}:B{R2},1,C{R1}:C{R2},\">=1\")", "0"),
    ("finPeriode", "Dernier mois affiché", f"=EDATE(AL14,{N_MONTHS - 1})", "MMM YYYY"),
    ("picPart", "Part du mois de pic dans la charge de la période", "=IF(AL15=0,0,AL17/AL15)", "0%"),
    ("periodeTxt", "Période affichée (bandeau des onglets Dashboard et Dessinateurs)",
     f'="Période "&{mois_fr("AL14")}&" "&YEAR(AL14)&" – "&{mois_fr("AL25")}&" "&YEAR(AL25)', "General"),
    ("aideDashboard", "Aide affichée sur le Dashboard (Oui / Non, bouton « Afficher / masquer l'aide »)", "Oui",
     "General"),
    ("aideDessinateurs", "Aide affichée sur Dessinateurs (Oui / Non, bouton « Afficher / masquer l'aide »)", "Oui",
     "General"),
    ("aideCharge", "Aide affichée sur Charge (Oui / Non, bouton « Afficher / masquer l'aide »)", "Oui", "General"),
    ("aideData", "Aide affichée sur Data (Oui / Non, bouton « Afficher / masquer l'aide »)", "Oui", "General"),
    ("macrosActives", "Macros actives (Oui écrit à l'ouverture par la macro ; Non dans le fichier enregistré)", "Non",
     "General"),
]
_n = len(INDICATORS)
_cap, _sur = 4 + _n, 5 + _n
INDICATORS += [
    ("capEquipe", "Capacité projets de l'équipe (h/mois) : somme des (capacité - charge fixe) des ressources Data",
     f"=SUM(FM4:FM{3 + 2 * (DATA_LAST - 3)})", "#,##0"),
    ("nbSurcharge", "Mois où la charge dépasse la capacité projets de l'équipe",
     f"=IF(AL{_cap}<=0,0,SUMPRODUCT(--(ROUND(AP4:AP{3 + N_MONTHS},0)>ROUND(AL{_cap},0))))", "0"),
    ("premSurcharge", "Premier mois en surcharge",
     f'=IF(AL{_sur}=0,"",INDEX(AO4:AO{3 + N_MONTHS},MATCH(1,INDEX(--(ROUND(AP4:AP{3 + N_MONTHS},0)>'
     f'ROUND(AL{_cap},0)),0),0)))', "MMM YYYY"),
    ("picCap", "Pic de charge rapporté à la capacité de l'équipe", f"=IF(AL{_cap}<=0,0,AL17/AL{_cap})", "0%"),
    ("nbRetard", "Affaires en retard", f'=COUNTIF(FL{R1}:FL{R2},">0")', "0"),
    ("nbDepasse", "dont date de besoin dépassée sans être à 100 %", f"=COUNTIF(FL{R1}:FL{R2},2)", "0"),
    ("nbEcart", "dont avancement inférieur de plus de 15 points à l'attendu", f"=COUNTIF(FL{R1}:FL{R2},1)", "0"),
    ("reglagesTxt", "Rappel des réglages (bandeau du Dashboard)",
     f'=p_Perimetre&"  ·  "&IF(p_Repartition="{REP_SMOOTH}","charge lissée en jours ouvrés jusqu\'à la date de '
     f'besoin","charge sur le mois de début planifié")&"  ·  "&p_HeuresETP&" h par personne dont "&p_ChargeFixe'
     f'&" h hors projet"', "General"),
]
ws["AK3"], ws["AL3"] = "Indicateur", "Valeur"
for c in ("AK3", "AL3"):
    ws[c].font = hdr_font
    ws[c].fill = fill(NAVY)
ws.column_dimensions["AK"].width = 44
ws.column_dimensions["AL"].width = 14
for i, (key, label, formula, fmt) in enumerate(INDICATORS):
    r = 4 + i
    K[key] = f"Calculs!$AL${r}"
    ws[f"AK{r}"] = label
    ws[f"AL{r}"] = formula
    ws[f"AL{r}"].number_format = fmt
# contrôle de cohérence des références codées en dur (AL4, AL5, AL7, AL9, AL13...)
expected = {"total": 4, "nbAff": 5, "nbProj": 7, "realise": 9, "autoStart": 13, "start": 14,
            "periode": 15, "pic": 17, "picIdx": 18}
for k, row in expected.items():
    assert K[k] == f"Calculs!$AL${row}", (k, K[k])
for sheet_title in ("Dashboard", "Dessinateurs", "Charge", "Data"):
    ref = K["aide" + sheet_title]
    wb.defined_names["aide_" + sheet_title] = DefinedName("aide_" + sheet_title, attr_text=ref)
    v = DataValidation(type="list", formula1=HELP_CHOICES, allow_blank=False)
    v.add(ref.split("!")[1].replace("$", ""))
    ws.add_data_validation(v)
wb.defined_names["macros_actives"] = DefinedName("macros_actives", attr_text=K["macrosActives"])

# Charge par mois (AN..AQ)
for c, t in zip(("AN", "AO", "AP", "AQ"), ("N° mois", "Mois", "Charge (h)", "Affaires actives")):
    ws[f"{c}3"] = t
    ws[f"{c}3"].font = hdr_font
    ws[f"{c}3"].fill = fill(NAVY)
for m, mc in enumerate(MCOLS):
    r = 4 + m
    ws[f"AN{r}"] = m + 1
    ws[f"AO{r}"] = f"={mc}$3"
    ws[f"AO{r}"].number_format = "MMM YYYY"
    ws[f"AP{r}"] = f"=SUM({mc}{R1}:{mc}{R2})"
    ws[f"AP{r}"].number_format = "#,##0.0"
    ws[f"AQ{r}"] = (f"=SUMPRODUCT(({mc}{R1}:{mc}{R2}>0)/(COUNTIFS($EN${R1}:$EN${R2},$EN${R1}:$EN${R2},"
                     f"{mc}{R1}:{mc}{R2},\">0\")+({mc}{R1}:{mc}{R2}<=0)))")

# Matrice des groupes de couleur (AS..BH) : séries des graphiques
SM = [L(49 + m) for m in range(N_MONTHS)]  # AW..BH
for c, t in zip(("AS", "AT", "AU", "AV"), ("Groupe", "Projet", "Charge (h)", "Part")):
    ws[f"{c}3"] = t
    ws[f"{c}3"].font = hdr_font
    ws[f"{c}3"].fill = fill(NAVY)
ws.column_dimensions["AT"].width = 22
for m, c in enumerate(SM):
    ws[f"{c}3"] = f"={MCOLS[m]}$3"
    ws[f"{c}3"].number_format = "MMM YY"
    ws[f"{c}3"].font = hdr_font
    ws[f"{c}3"].fill = fill(NAVY)
SLOT_R1 = 4
SLOT_R2 = SLOT_R1 + N_SLOTS - 1
TOT_ROW = SLOT_R2 + 1
for k in range(1, N_SLOTS + 1):
    r = SLOT_R1 + k - 1
    ws[f"AS{r}"] = k
    if k == N_SLOTS:
        ws[f"AT{r}"] = (f'=IF(AS{r}>$AL$7,"",IF($AL$7>{N_SLOTS},"Autres projets ("&($AL$7-{N_SLOTS - 1})&")",'
                        f'INDEX($G${R1}:$G${R2},MATCH(AS{r},$R${R1}:$R${R2},0))))')
    else:
        ws[f"AT{r}"] = f'=IF(AS{r}>$AL$7,"",INDEX($G${R1}:$G${R2},MATCH(AS{r},$R${R1}:$R${R2},0)))'
    ws[f"AU{r}"] = f"=SUMIF($V${R1}:$V${R2},AS{r},$E${R1}:$E${R2})"
    ws[f"AV{r}"] = f"=IF($AL$4=0,0,AU{r}/$AL$4)"
    ws[f"AV{r}"].number_format = "0.0%"
    for m, c in enumerate(SM):
        ws[f"{c}{r}"] = f"=SUMIF($V${R1}:$V${R2},$AS{r},{MCOLS[m]}${R1}:{MCOLS[m]}${R2})"
        ws[f"{c}{r}"].number_format = "0.0"
ws[f"AT{TOT_ROW}"] = "Total"
ws[f"AU{TOT_ROW}"] = f"=SUM(AU{SLOT_R1}:AU{SLOT_R2})"
for c in SM:
    ws[f"{c}{TOT_ROW}"] = f"=ROUND(SUM({c}{SLOT_R1}:{c}{SLOT_R2}),0)"
    ws[f"{c}{TOT_ROW}"].number_format = "#,##0"
CAP_ROW = TOT_ROW + 1   # capacité projets de l'équipe (ligne pointillée du graphique)
ws[f"AT{CAP_ROW}"] = "Capacité projets de l'équipe"
for c in SM:
    ws[f"{c}{CAP_ROW}"] = f"={K['capEquipe'].split('!')[1]}"
    ws[f"{c}{CAP_ROW}"].number_format = "#,##0"


def top_list(c_k, c_name, c_h, n, count_ref, key_col, tot_col, rank_col, other_label, title):
    """Liste des n premiers éléments par charge décroissante ; la dernière ligne regroupe le reste."""
    for c, t in zip((c_k, c_name, c_h), ("Rang", title, "Charge (h)")):
        ws[f"{c}3"] = t
        ws[f"{c}3"].font = hdr_font
        ws[f"{c}3"].fill = fill(NAVY)
    ws.column_dimensions[c_name].width = 22
    for k in range(1, n + 1):
        r = 3 + k
        ws[f"{c_k}{r}"] = k
        look = f"MATCH({c_k}{r},${rank_col}${R1}:${rank_col}${R2},0)"
        if k == n:
            ws[f"{c_name}{r}"] = (f'=IF({c_k}{r}>{count_ref},"",IF({count_ref}>{n},'
                                  f'"+ "&({count_ref}-{n - 1})&" {other_label}",'
                                  f'INDEX(${key_col}${R1}:${key_col}${R2},{look})))')
            ws[f"{c_h}{r}"] = (f'=IF({c_name}{r}="","",IF({count_ref}>{n},$AL$4-SUM({c_h}4:{c_h}{r - 1}),'
                               f'INDEX(${tot_col}${R1}:${tot_col}${R2},{look})))')
        else:
            ws[f"{c_name}{r}"] = f'=IF({c_k}{r}>{count_ref},"",INDEX(${key_col}${R1}:${key_col}${R2},{look}))'
            ws[f"{c_h}{r}"] = f'=IF({c_name}{r}="","",INDEX(${tot_col}${R1}:${tot_col}${R2},{look}))'


top_list("BK", "BL", "BM", 8, "$AL$6", "F", "N", "O", "autres", "Chef de projet")
top_list("BO", "BP", "BQ", 8, "$AL$7", "G", "Q", "R", "autres projets", "Projet")
top_list("BS", "BT", "BU", 4, "$AL$8", "H", "T", "U", "autres types", "Type")

# Tableau des affaires : position (dans Charge) de l'affaire de rang k
for c, t in zip(("BW", "BX", "BY"), ("Rang", "Position dans Charge", "Retard (2 date dépassée, 1 écart)")):
    ws[f"{c}3"] = t
    ws[f"{c}3"].font = hdr_font
    ws[f"{c}3"].fill = fill(NAVY)
for k in range(1, N_TABLE + 1):
    r = 3 + k
    ws[f"BW{r}"] = k
    ws[f"BX{r}"] = f'=IFERROR(MATCH(BW{r},$EQ${R1}:$EQ${R2},0),"")'
    ws[f"BY{r}"] = f'=IF(BX{r}="","",INDEX($FL${R1}:$FL${R2},BX{r}))'

# ---- Affaires = sous-projets regroupés (projet + sous-projet + produit), toutes ressources confondues
AFF_COLS = [("FD", "Ligne avec dessinateur ET externe (1 = à corriger)"),
            ("FE", "1re ligne de cette ressource sur l'affaire"), ("EM", "Type de ressource de la ligne"), ("EN", "Clé affaire (projet | sous-projet | produit)"),
            ("EO", "1re ligne de l'affaire"), ("EP", "Charge de l'affaire (h)"), ("EQ", "Rang de l'affaire"),
            ("ER", "Personnes internes (distinctes)"), ("ES", "Personnes externes (distinctes)"),
            ("ET", "Ressources (texte)"),
            ("EU", "Avancement de l'affaire (pondéré par les heures)"), ("EV", "Début de la ligne (31/12/9999 = vide)"),
            ("EW", "Début de l'affaire (le plus tôt)"), ("EX", "Besoin de la ligne (0 = vide)"),
            ("EY", "Besoin de l'affaire (le plus tard)"),
            ("FF", "Clé affaire, toutes lignes valides (périmètre ignoré)"), ("FG", "Heures de la ligne (périmètre ignoré)"),
            ("FH", "Heures de l'affaire (toutes ses lignes)"), ("FI", "Heures réalisées de l'affaire (selon avancement)"),
            ("FJ", "Reste à faire de l'affaire (h)"), ("FK", "Avancement attendu (prorata des jours ouvrés écoulés)"),
            ("FL", "Retard : 2 = besoin dépassé, 1 = écart > 15 pts, 0 = à l'heure")]
for c, t in AFF_COLS:
    ws[f"{c}3"] = t
    ws[f"{c}3"].font = hdr_font
    ws[f"{c}3"].fill = fill(NAVY)
    ws[f"{c}3"].alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.column_dimensions[c].width = 13
ws.column_dimensions["EN"].width = 30
A_RNG = {c: f"${c}${R1}:${c}${R2}" for c in ("C", "E", "EM", "EN", "EO", "EP", "EV", "EX", "FF", "FG")}
for r in range(R1, R2 + 1):
    ws[f"EM{r}"] = (f'=IF(D{r}=1,IF(TRIM(Charge!K{r})<>"","Interne",IF(TRIM(Charge!L{r})<>"","Externe",'
                    f'"Non affecté")),"")')
    part = lambda c: f'TRIM(SUBSTITUTE(Charge!{c}{r},CHAR(160)," "))'
    ws[f"EN{r}"] = (f'=IF(D{r}=1,LEFT(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({part("B")}&" | "&{part("C")}&" | "&{part("D")},'
                    f'"~","-"),"*","+"),"?","!"),240),"")')
    ws[f"EO{r}"] = f'=IF(EN{r}="",0,IF(COUNTIF($EN${R1}:EN{r},EN{r})=1,1,0))'
    ws[f"EP{r}"] = f'=IF(EO{r}=1,ROUND(SUMIF({A_RNG["EN"]},EN{r},{A_RNG["E"]}),2),"")'
    ws[f"EQ{r}"] = (f'=IF(EO{r}=1,COUNTIFS({A_RNG["EO"]},1,{A_RNG["EP"]},">"&EP{r})'
                    f'+COUNTIFS($EO${R1}:EO{r},1,$EP${R1}:EP{r},EP{r}),"")')
    ws[f"FD{r}"] = f'=IF(AND(B{r}=1,TRIM(Charge!K{r})<>"",TRIM(Charge!L{r})<>""),1,0)'
    ws[f"FE{r}"] = (f'=IF(OR(EN{r}="",BZ{r}="{NA_LABEL}"),0,'
                    f'IF(COUNTIFS($EN${R1}:EN{r},EN{r},$BZ${R1}:BZ{r},BZ{r})=1,1,0))')
    ws[f"ER{r}"] = f'=IF(EO{r}=1,COUNTIFS({A_RNG["EN"]},EN{r},{A_RNG["EM"]},"Interne",$FE${R1}:$FE${R2},1),"")'
    ws[f"ES{r}"] = f'=IF(EO{r}=1,COUNTIFS({A_RNG["EN"]},EN{r},{A_RNG["EM"]},"Externe",$FE${R1}:$FE${R2},1),"")'
    ws[f"ET{r}"] = (f'=IF(EO{r}=1,IF(ER{r}+ES{r}=0,"à affecter",IF(ER{r}=0,ES{r}&" ext.",'
                    f'(ER{r}+ES{r})&IF(ES{r}>0," dont "&ES{r}&" ext.",""))'
                    f'&IF(COUNTIFS({A_RNG["EN"]},EN{r},{A_RNG["EM"]},"Non affecté")>0," + à affecter","")),"")')
    # avancement, réalisé et reste à faire de l'affaire : TOUTES ses lignes, même celles hors périmètre (terminées)
    ws[f"FF{r}"] = (f'=IF(B{r}=1,LEFT(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({part("B")}&" | "&{part("C")}&" | "&{part("D")},'
                    f'"~","-"),"*","+"),"?","!"),240),"")')
    ws[f"FG{r}"] = f"=IF(B{r}=1,MAX(0,N(Charge!M{r})),0)"
    ws[f"FH{r}"] = f'=IF(EO{r}=1,ROUND(SUMIF({A_RNG["FF"]},EN{r},{A_RNG["FG"]}),2),"")'
    ws[f"FI{r}"] = f'=IF(EO{r}=1,SUMPRODUCT(({A_RNG["FF"]}=EN{r})*{A_RNG["FG"]}*{A_RNG["C"]}),"")'
    ws[f"FJ{r}"] = f'=IF(EO{r}=1,MAX(0,ROUND(FH{r}-FI{r},2)),"")'
    ws[f"EU{r}"] = (f'=IF(EO{r}=1,IF(FH{r}>0,FI{r}/FH{r},AVERAGEIF({A_RNG["FF"]},EN{r},{A_RNG["C"]})),"")')
    ws[f"EV{r}"] = f"=IF(ISNUMBER(I{r}),I{r},2958465)"
    ws[f"EW{r}"] = (f'=IF(EO{r}=1,SUMPRODUCT(MIN(({A_RNG["EN"]}=EN{r})*{A_RNG["EV"]}'
                    f'+({A_RNG["EN"]}<>EN{r})*2958465)),"")')
    ws[f"EX{r}"] = f"=IF(AND(D{r}=1,ISNUMBER(Charge!F{r})),Charge!F{r},0)"
    ws[f"EY{r}"] = f'=IF(EO{r}=1,SUMPRODUCT(MAX(({A_RNG["EN"]}=EN{r})*{A_RNG["EX"]})),"")'
    # retard : besoin dépassé sans être à 100 %, ou avancement < avancement attendu - 15 points
    ws[f"FK{r}"] = (f'=IF(EO{r}=1,IF(OR(EW{r}>=2958465,EY{r}=0,EY{r}<=EW{r}),"",IF(TODAY()<=EW{r},0,'
                    f'IF(TODAY()>=EY{r},1,IFERROR(NETWORKDAYS(EW{r},TODAY(),{HOLIDAYS})'
                    f'/NETWORKDAYS(EW{r},EY{r},{HOLIDAYS}),"")))),"")')
    ws[f"FL{r}"] = (f'=IF(EO{r}=1,IF(AND(EY{r}>0,EY{r}<TODAY(),ROUND(EU{r},4)<1),2,'
                    f'IF(FK{r}="",0,IF(ROUND(EU{r}-FK{r},4)<-0.15,1,0))),"")')
    for c in ("FH", "FI", "FJ"):
        ws[f"{c}{r}"].number_format = "0.0"
    ws[f"FK{r}"].number_format = "0%"
    ws[f"EP{r}"].number_format = "0.0"
    ws[f"EU{r}"].number_format = "0%"
    for c in ("EV", "EW", "EX", "EY"):
        ws[f"{c}{r}"].number_format = "DD/MM/YYYY"

# ---- Charge par ressource (dessinateur interne ou ressource externe)
N_DESS = 16          # ressources affichées dans le tableau d'occupation de l'équipe
N_DESS_AFF = 15      # affaires listées pour le dessinateur sélectionné


def hdr(col, text, width=None):
    ws[f"{col}3"] = text
    ws[f"{col}3"].font = hdr_font
    ws[f"{col}3"].fill = fill("2F5C8F")
    ws[f"{col}3"].alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    if width:
        ws.column_dimensions[col].width = width


# par ligne : dessinateur (affaires du périmètre) et rang parmi les affaires du dessinateur sélectionné
hdr("BZ", "Ressource de la ligne (K sinon L)", 16)
hdr("CA", "Rang affaire du dessinateur sélectionné", 12)
for r in range(R1, R2 + 1):
    ws[f"BZ{r}"] = (f'=IF(D{r}=1,IF(TRIM(Charge!K{r})<>"",TRIM(Charge!K{r}),IF(TRIM(Charge!L{r})<>"",'
                    f'TRIM(Charge!L{r}),"{NA_LABEL}")),"")')
    ws[f"CA{r}"] = (f'=IF(AND($CD$4<>"",BZ{r}=$CD$4),COUNTIFS($BZ${R1}:$BZ${R2},$CD$4,$L${R1}:$L${R2},"<"&L{r})+1,"")')

# noms des dessinateurs de la liste Data, sans espaces superflus (EJ)
hdr("EJ", "Dessinateurs Data (nettoyés)", 18)
for d in range(4, DATA_LAST + 1):
    ws[f"EJ{d}"] = f"=TRIM(INDEX(Data!$E:$E,ROW()))"
DATA_NAMES = f"$EJ$4:$EJ${DATA_LAST}"
hdr("EK", "Ressources externes Data (nettoyées)", 18)
for d in range(4, DATA_LAST + 1):
    ws[f"EK{d}"] = f"=TRIM(INDEX(Data!$H:$H,ROW()))"
DATA_EXT = f"$EK$4:$EK${DATA_LAST}"
N_RES = 2 * (DATA_LAST - 3)   # capacité de la liste commune (dessinateurs puis externes)
hdr("FA", "Ressources : dessinateurs puis externes (liste déroulante)", 22)
hdr("FB", "Type", 10)
hdr("FM", "Capacité projets (capacité - charge fixe, h/mois)", 14)
for k in range(1, N_RES + 1):
    r = 3 + k
    ws[f"FA{r}"] = (f'=IF({k}<=$CD$23,INDEX({DATA_NAMES},{k}),IF({k}-$CD$23<=$CD$26,'
                    f'INDEX({DATA_EXT},{k}-$CD$23),""))')
    ws[f"FB{r}"] = f'=IF(FA{r}="","",IF({k}<=$CD$23,"Interne","Externe"))'
    dup_fa = f"COUNTIF($FA$4:FA{r - 1},FA{r})>0" if r > 4 else "FALSE"
    v = {}
    for key, d_int, d_ext, dflt in (("cap", "F", "I", "p_HeuresETP"), ("fixe", "G", "J", "p_ChargeFixe")):
        vi, ve = f"INDEX(Data!${d_int}:${d_int},{k + 3})", f"INDEX(Data!${d_ext}:${d_ext},{k + 3}-$CD$23)"
        v[key] = (f'IF(FB{r}="Interne",IF(ISNUMBER({vi}),{vi},{dflt}),IF(ISNUMBER({ve}),{ve},{dflt}))')
    ws[f"FM{r}"] = f'=IF(OR(FA{r}="",{dup_fa}),0,MAX(0,{v["cap"]}-{v["fixe"]}))'
wb.defined_names["lst_Ressources"] = DefinedName(
    "lst_Ressources", attr_text="OFFSET(Calculs!$FA$4,0,0,MAX(1,Calculs!$CD$23+Calculs!$CD$26),1)")

# dessinateur sélectionné : réglages et indicateurs (CC libellé / CD valeur)
DS = {}
DESS_IND = [
    ("nom", "Ressource sélectionnée", "=TRIM(p_Dessinateur)", "General"),
    ("pos", "Position dans la liste des dessinateurs (0 = absent)",
     f"=IF(CD4=\"\",0,IFERROR(MATCH(CD4,{DATA_NAMES},0),0))", "0"),
    ("cap", "Capacité mensuelle (h)",
     f"=IF(CD4=\"\",0,IF(CD5>0,IF(ISNUMBER(INDEX(Data!$F:$F,CD5+3)),"
     f"INDEX(Data!$F:$F,CD5+3),p_HeuresETP),IF(CD24>0,IF(ISNUMBER(INDEX(Data!$I:$I,CD24+3)),"
     f"INDEX(Data!$I:$I,CD24+3),p_HeuresETP),p_HeuresETP)))", "#,##0.0"),
    ("fixe", "Charge fixe hors projet (h/mois)",
     f"=IF(CD4=\"\",0,IF(CD5>0,IF(ISNUMBER(INDEX(Data!$G:$G,CD5+3)),"
     f"INDEX(Data!$G:$G,CD5+3),p_ChargeFixe),IF(CD24>0,IF(ISNUMBER(INDEX(Data!$J:$J,CD24+3)),"
     f"INDEX(Data!$J:$J,CD24+3),p_ChargeFixe),p_ChargeFixe)))", "#,##0.0"),
    ("capProj", "Capacité disponible pour les projets (h/mois)", "=MAX(0,CD6-CD7)", "#,##0.0"),
    ("chargeTot", "Charge projets affectée, toutes périodes (h)",
     f'=IF(CD4="",0,SUMIF($BZ${R1}:$BZ${R2},CD4,$E${R1}:$E${R2}))', "#,##0.0"),
    ("nbAff", "Nombre d'affaires (sous-projets) de la ressource",
     f'=IF(CD4="",0,COUNTIFS($BZ${R1}:$BZ${R2},CD4,$FE${R1}:$FE${R2},1))', "0"),
    ("periode", "Charge projets sur la période affichée (h)", "=SUM(CI4:CI15)", "#,##0.0"),
    ("pic", "Mois le plus chargé : charge totale (h)", "=MAX(CJ4:CJ15)", "#,##0.0"),
    ("picIdx", "Rang du mois le plus chargé", "=MATCH(CD12,CJ4:CJ15,0)", "0"),
    ("picMois", "Mois le plus chargé", "=INDEX(CG4:CG15,CD13)", "MMM YYYY"),
    ("picOcc", "Occupation du mois le plus chargé", "=INDEX(CO4:CO15,CD13)", "0%"),
    ("nbSur", "Nombre de mois en surcharge", "=COUNT(CQ4:CQ15)", "0"),
    ("depass", "Dépassement cumulé (h)", "=SUMPRODUCT((CN4:CN15<0)*(-CN4:CN15))", "#,##0"),
    ("dispo", "Heures encore disponibles sur la période (h)", "=SUMPRODUCT((CN4:CN15>0)*CN4:CN15)", "#,##0"),
    ("occMoy", "Occupation moyenne sur la période", "=AVERAGE(CO4:CO15)", "0%"),
    ("premSur", "1er mois en surcharge", '=IF(CD16=0,"",INDEX(CG4:CG15,MIN(CQ4:CQ15)))', "MMM YYYY"),
    ("perim", "Périmètre (rappel)", "=p_Perimetre", "General"),
    ("debut", "1er mois affiché (rappel)", "=$AL$14", "MMM YYYY"),
    ("nbListe", "Nombre de dessinateurs dans la liste Data", f"=COUNTA(Data!$E$4:$E${DATA_LAST})", "0"),
    ("posExt", "Position dans la liste des ressources externes (0 = absent)",
     f"=IF(OR(CD4=\"\",CD5>0),0,IFERROR(MATCH(CD4,{DATA_EXT},0),0))", "0"),
    ("type", "Type de ressource", '=IF(CD4="","",IF(CD5>0,"Interne",IF(CD24>0,"Externe","Hors liste")))', "General"),
    ("nbExt", "Nombre de ressources externes dans la liste Data", f"=COUNTA(Data!$H$4:$H${DATA_LAST})", "0"),
    ("nbDouble", "Lignes avec un dessinateur ET une ressource externe (à corriger)", f"=SUM(FD{R1}:FD{R2})", "0"),
    ("nbLignes", "Nombre de lignes de la ressource", f'=IF(CD4="",0,COUNTIF($BZ${R1}:$BZ${R2},CD4))', "0"),
]
hdr("CC", "Dessinateur sélectionné", 40)
hdr("CD", "Valeur", 14)
for i, (key, label, formula, fmt) in enumerate(DESS_IND):
    r = 4 + i
    DS[key] = f"Calculs!$CD${r}"
    ws[f"CC{r}"] = label
    ws[f"CD{r}"] = formula
    ws[f"CD{r}"].number_format = fmt
for k, row in {"nom": 4, "pos": 5, "cap": 6, "fixe": 7, "pic": 12, "picIdx": 13, "nbSur": 16, "nbListe": 23,
               "posExt": 24, "nbExt": 26}.items():
    assert DS[k] == f"Calculs!$CD${row}", k

# charge mensuelle du dessinateur sélectionné (CF..CQ, une ligne par mois)
DM = ["CF", "CG", "CH", "CI", "CJ", "CK", "CL", "CM", "CN", "CO", "CP", "CQ"]
for c, t in zip(DM, ["N° mois", "Mois", "Charge fixe (h)", "Charge projets (h)", "Charge totale (h)",
                     "Capacité (h)", "Projets dans la capacité (h)", "Dépassement (h)",
                     "Marge (+) / dépassement (-) (h, total arrondi)", "Occupation (total arrondi)",
                     "Charge totale arrondie (h)", "N° mois si surcharge"]):
    hdr(c, t, 11)
for m, mc in enumerate(MCOLS):
    r = 4 + m
    ws[f"CF{r}"] = m + 1
    ws[f"CG{r}"] = f"={mc}$3"
    ws[f"CG{r}"].number_format = "MMM YY"
    ws[f"CH{r}"] = '=IF($CD$4="",0,$CD$7)'
    ws[f"CI{r}"] = f'=IF($CD$4="",0,SUMIF($BZ${R1}:$BZ${R2},$CD$4,{mc}${R1}:{mc}${R2}))'
    ws[f"CJ{r}"] = f"=CH{r}+CI{r}"
    ws[f"CK{r}"] = '=IF($CD$4="",0,$CD$6)'
    ws[f"CL{r}"] = f"=IF(CN{r}<0,MAX(0,MIN(CI{r},CK{r}-CH{r})),CI{r})"   # projets dans la capacité
    ws[f"CM{r}"] = f"=CI{r}-CL{r}"                                        # dépassement (rouge)
    ws[f"CN{r}"] = f"=CK{r}-CP{r}"                                        # marge sur le total arrondi
    ws[f"CO{r}"] = f"=IF(CK{r}<=0,IF(CP{r}>0,9.99,0),CP{r}/CK{r})"
    ws[f"CP{r}"] = f"=ROUND(CJ{r},0)"
    ws[f"CQ{r}"] = f'=IF(CN{r}<0,CF{r},"")'
    for c in ("CH", "CI", "CJ", "CK", "CL", "CM", "CN"):
        ws[f"{c}{r}"].number_format = "0.0"
    ws[f"CO{r}"].number_format = "0%"

# occupation de l'équipe : dessinateurs de la liste Data (CS..DV)
TEAM_NAME, TEAM_CAP, TEAM_FIX = "CT", "CU", "CV"
TEAM_H = [L(101 + m) for m in range(N_MONTHS)]          # CW..DH : heures projets
TEAM_O = [L(101 + N_MONTHS + m) for m in range(N_MONTHS)]  # DI..DT : occupation
TEAM_MAX = L(101 + 2 * N_MONTHS)                        # DU
TEAM_TOT = L(102 + 2 * N_MONTHS)                        # DV
hdr("CS", "N°", 6)
hdr(TEAM_NAME, "Ressource", 18)
hdr(TEAM_CAP, "Capacité (h)", 10)
hdr(TEAM_FIX, "Charge fixe (h)", 10)
for m in range(N_MONTHS):
    ws[f"{TEAM_H[m]}3"] = f"={MCOLS[m]}$3"
    ws[f"{TEAM_O[m]}3"] = f"={MCOLS[m]}$3"
    for c in (TEAM_H[m], TEAM_O[m]):
        ws[f"{c}3"].number_format = "MMM YY"
        ws[f"{c}3"].font = hdr_font
        ws[f"{c}3"].fill = fill("2F5C8F")
ws[f"{TEAM_H[0]}2"] = "Heures projets par mois →"
ws[f"{TEAM_O[0]}2"] = "Occupation (charge fixe + projets) / capacité →"
hdr(TEAM_MAX, "Occupation maximale", 11)
hdr(TEAM_TOT, "Charge projets sur la période (h)", 11)
TEAM_R1 = 4
TEAM_R2 = TEAM_R1 + N_DESS - 1
NA_ROW = TEAM_R2 + 1
OTHER_ROW = TEAM_R2 + 2
for k in range(1, N_DESS + 1):
    r = TEAM_R1 + k - 1
    ws[f"CS{r}"] = k   # même rang que la liste commune FA/FB (ligne r)
    dup = f"COUNTIF($FA$4:FA{r - 1},FA{r})>0" if r > 4 else "FALSE"
    ws[f"{TEAM_NAME}{r}"] = f'=IF(OR(FA{r}="",{dup}),"",FA{r})'
    for col_t, d_int, d_ext, dflt in ((TEAM_CAP, "F", "I", "p_HeuresETP"), (TEAM_FIX, "G", "J", "p_ChargeFixe")):
        v_int = f"INDEX(Data!${d_int}:${d_int},{k + 3})"
        v_ext = f"INDEX(Data!${d_ext}:${d_ext},{k + 3}-$CD$23)"
        ws[f"{col_t}{r}"] = (f'=IF({TEAM_NAME}{r}="","",IF(FB{r}="Interne",IF(ISNUMBER({v_int}),{v_int},{dflt}),'
                             f'IF(ISNUMBER({v_ext}),{v_ext},{dflt})))')
    for m in range(N_MONTHS):
        h, o, mc = TEAM_H[m], TEAM_O[m], MCOLS[m]
        ws[f"{h}{r}"] = f'=IF(${TEAM_NAME}{r}="","",SUMIF($BZ${R1}:$BZ${R2},${TEAM_NAME}{r},{mc}${R1}:{mc}${R2}))'
        tot = f"ROUND(${TEAM_FIX}{r}+{h}{r},0)"
        ws[f"{o}{r}"] = (f'=IF(${TEAM_NAME}{r}="","",IF(${TEAM_CAP}{r}<=0,IF({tot}>0,9.99,0),'
                         f'{tot}/${TEAM_CAP}{r}))')
        ws[f"{h}{r}"].number_format = "0.0"
        ws[f"{o}{r}"].number_format = "0%"
    ws[f"{TEAM_MAX}{r}"] = f'=IF({TEAM_NAME}{r}="","",MAX({TEAM_O[0]}{r}:{TEAM_O[-1]}{r}))'
    ws[f"{TEAM_MAX}{r}"].number_format = "0%"
    ws[f"{TEAM_TOT}{r}"] = f'=IF({TEAM_NAME}{r}="","",SUM({TEAM_H[0]}{r}:{TEAM_H[-1]}{r}))'
    ws[f"{TEAM_TOT}{r}"].number_format = "#,##0.0"
ws[f"{TEAM_NAME}{NA_ROW}"] = NA_LABEL
ws[f"{TEAM_NAME}{OTHER_ROW}"] = f"Autres (absents des listes Data ou au-delà du {N_DESS}e)"
for m in range(N_MONTHS):
    h, mc = TEAM_H[m], MCOLS[m]
    ws[f"{h}{NA_ROW}"] = f'=SUMIF($BZ${R1}:$BZ${R2},"{NA_LABEL}",{mc}${R1}:{mc}${R2})'
    ws[f"{h}{OTHER_ROW}"] = f"=MAX(0,ROUND(AP{4 + m}-SUM({h}{TEAM_R1}:{h}{NA_ROW}),1))"
    ws[f"{h}{NA_ROW}"].number_format = ws[f"{h}{OTHER_ROW}"].number_format = "0.0"
for r in (NA_ROW, OTHER_ROW):
    ws[f"{TEAM_TOT}{r}"] = f"=SUM({TEAM_H[0]}{r}:{TEAM_H[-1]}{r})"
    ws[f"{TEAM_TOT}{r}"].number_format = "#,##0.0"

# affaires du dessinateur sélectionné (DX..EH), par charge décroissante
DA = ["DX", "DY", "DZ", "EA", "EB", "EC", "ED", "EE", "EF", "EG", "EH"]
for c, t in zip(DA, ["Rang", "Position dans Charge", "Projet", "Sous-projet", "Produit", "Chef de projet",
                     "Charge (h)", "Heures sur la période", "Début planifié", "Date de besoin", "Avancement"]):
    hdr(c, t, 13)
for k in range(1, N_DESS_AFF + 1):
    r = 3 + k
    i = f"DY{r}"
    ws[f"DX{r}"] = k
    ws[i] = f'=IFERROR(MATCH(DX{r},$CA${R1}:$CA${R2},0),"")'
    for c, col in (("DZ", "B"), ("EA", "C"), ("EB", "D"), ("EC", "A")):
        ws[f"{c}{r}"] = f'=IF({i}="","",INDEX(Charge!${col}:${col},{i}+{R1 - 1})&"")'
    ws[f"ED{r}"] = f'=IF({i}="","",INDEX($E${R1}:$E${R2},{i}))'
    ws[f"EE{r}"] = f'=IF({i}="","",INDEX(${SUM_COL}${R1}:${SUM_COL}${R2},{i}))'
    for c, col in (("EF", "N"), ("EG", "F")):
        src = f"INDEX(Charge!${col}:${col},{i}+{R1 - 1})"
        ws[f"{c}{r}"] = f'=IF({i}="","",IF(ISNUMBER({src}),{src},"—"))'
        ws[f"{c}{r}"].number_format = "DD/MM/YYYY"
    ws[f"EH{r}"] = f'=IF({i}="","",INDEX($C${R1}:$C${R2},{i}))'
    ws[f"EH{r}"].number_format = "0%"

# Chaque ligne de Calculs lit la même ligne de Charge par INDEX(colonne entière ; ROW()) : insérer, supprimer
# ou trier des lignes dans l'onglet Charge ne décale ni ne casse les calculs.
_rel = re.compile(r"Charge!([A-O])(\d+)(?![\d:])")
for row in ws.iter_rows(min_row=R1, max_row=R2):
    for cell in row:
        if isinstance(cell.value, str) and "Charge!" in cell.value:
            cell.value = _rel.sub(lambda m, rr=cell.row: (f"INDEX(Charge!${m.group(1)}:${m.group(1)},ROW())"
                                                          if int(m.group(2)) == rr else m.group(0)), cell.value)

ws.freeze_panes = "B4"

# ---------------------------------------------------------------- Dashboard
ws = ws_dash
ws.sheet_view.showGridLines = False
ws.sheet_view.zoomScale = 90
ws.column_dimensions["A"].width = MARGIN_W
for i in range(NG):
    ws.column_dimensions[gl(i)].width = 2.71
ws.column_dimensions[L(gc(NG))].width = MARGIN_W

# ligne 1 = marge haute, de même taille que les colonnes de marge A et BI
ROW_H = {1: MARGIN_H, 2: 30, 3: 16, 4: 16, 5: 8, 6: 22, 7: 14, 8: 24, 9: 6,
         10: 6, 11: 18, 12: 34, 13: 15, 14: 6, 23: 8, 24: 14, 25: 20, 26: 6,
         27: 6, 28: 18, 29: 34, 30: 15, 31: 15, 32: 8, 33: 14, 34: 20, 35: 6,
         36: 6, 37: 18, 38: 15, 55: 8, 56: 14, 57: 20, 58: 6, 59: 30}
for r in range(15, 23):
    ROW_H[r] = 16
for r in range(39, 55):
    ROW_H[r] = 15
TABLE_R1 = 60
TABLE_R2 = TABLE_R1 + N_TABLE - 1
for r in range(TABLE_R1, TABLE_R2 + 1):
    ROW_H[r] = 16.5
LAST_ROW = TABLE_R2 + 8   # dernière ligne = marge basse
ROW_H[LAST_ROW] = MARGIN_H
for r, h in ROW_H.items():
    ws.row_dimensions[r].height = h

# fond de page
for r in range(1, LAST_ROW + 1):
    for c in range(1, gc(NG) + 1):
        ws.cell(r, c).fill = fill(BG)


def box(r1, r2, c1, c2, bg=WHITE, border=LINE, top=None):
    """Carte : fond + bordure extérieure (c1, c2 = indices de grille)."""
    for r in range(r1, r2 + 1):
        for i in range(c1, c2 + 1):
            cell = ws.cell(r, gc(i))
            cell.fill = fill(bg)
            cell.border = Border(
                left=side(border) if i == c1 else None,
                right=side(border) if i == c2 else None,
                top=(top or side(border)) if r == r1 else None,
                bottom=side(border) if r == r2 else None)


def put(r, c1, c2, value, f=None, h="left", v="center", fmt=None, r2=None, bg=None, wrap=False, indent=0,
        shrink=False):
    r2 = r2 or r
    if c2 > c1 or r2 > r:
        ws.merge_cells(start_row=r, start_column=gc(c1), end_row=r2, end_column=gc(c2))
    cell = ws.cell(r, gc(c1))
    cell.value = value
    cell.font = f or font()
    cell.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent,
                               shrink_to_fit=shrink and not wrap)
    if fmt:
        cell.number_format = fmt
    if bg:
        cell.fill = fill(bg)
    return cell


def section(r, text):
    for i in range(NG):
        ws.cell(r, gc(i)).border = Border(bottom=side("B8C2CF"))
    c = ws.cell(r, gc(0))
    c.value = text
    c.font = font(11, True, NAVY)
    c.alignment = Alignment(vertical="bottom")


def banner(title, subtitle, edite_name, left4=None):
    """Bandeau des lignes 2 à 4 : titre, logo (BA2:BG2), « Édité le » et période affichée.
    « Édité le » est une date fixe, écrite par la macro « Éditer PDF » (nom défini `edite_name`)."""
    for r in (2, 3, 4):
        for i in range(NG):
            ws.cell(r, gc(i)).fill = fill(NAVY)
    put(2, 1, 40, title, font(17, True, WHITE), bg=NAVY)
    put(2, LOGO[0], LOGO[1], None, bg=NAVY)
    put(3, 1, 42, subtitle, font(9, color="C9D3E3"), bg=NAVY)
    put(3, 44, LOGO[1], None, font(10, color="C9D3E3"), h="right", fmt='"Édité le "DD/MM/YYYY', bg=NAVY)
    wb.defined_names[edite_name] = DefinedName(edite_name, attr_text=f"{ws.title}!${gl(44)}$3")
    put(4, 30, LOGO[1], f"={K['periodeTxt']}", font(10, True, WHITE), h="right", bg=NAVY)
    if left4:
        put(4, 1, 28, left4, font(9, color="C9D3E3"), bg=NAVY, shrink=True)


def auto_hide_rows(sheet_title, *areas):
    """Noms lignes_masquables_<onglet>_1, _2... (une zone par nom) : cellules repères dont la ligne est masquée
    par les macros quand elles sont vides (lignes inutilisées des listes, aide masquée, alertes absentes), à
    l'écran comme dans le PDF."""
    for i, area in enumerate(areas, 1):
        name = f"lignes_masquables_{sheet_title}_{i}"
        wb.defined_names[name] = DefinedName(name, attr_text=f"{sheet_title}!{area}")


# Bandeau titre
banner("TABLEAU DE BORD — PLAN DE CHARGE DU BUREAU D'ÉTUDES",
       "Synthèse automatique de l'onglet Charge : mettre à jour Charge et Data, le tableau de bord se recalcule seul.",
       "edite_Dashboard", left4=f"={K['reglagesTxt']}")
# la barre de réglages (lignes 5 à 7) n'est pas imprimée par « Éditer PDF » : rappel des réglages en ligne 4
wb.defined_names["lignes_pdf_Dashboard"] = DefinedName("lignes_pdf_Dashboard", attr_text="Dashboard!$A$5:$A$7")

# Barre de paramètres
box(P_ROW, P_ROW + 1, 0, NG - 1)
labels = {"p_Perimetre": (0, 3, "Périmètre"), "p_Repartition": (12, 15, "Répartition"),
          "p_PremierMois": (29, 34, "1er mois affiché"), "p_HeuresETP": (41, 44, "Capacité"),
          "p_ChargeFixe": (50, 53, "Charge fixe")}
defaults = {"p_Perimetre": PERIM_RUN, "p_Repartition": REP_SMOOTH, "p_PremierMois": None, "p_HeuresETP": 120,
            "p_ChargeFixe": 20}
hints = {"p_Perimetre": ("En cours = avancement < 100 %", None),
         "p_Repartition": (f'=IF(p_Repartition="{REP_START}","Toute la charge sur le mois de début planifié",'
                           f'"Prorata des jours ouvrés entre début planifié et date de besoin")', None),
         "p_PremierMois": ('=IF(ISNUMBER(p_PremierMois),"Saisi manuellement",Calculs!$AL$13)',
                           '"Vide = auto ("MMM YYYY")"'),
         "p_HeuresETP": ("par personne · propre : Data F / I", None),
         "p_ChargeFixe": ("hors projet · propre : Data G / J", None)}
inp_border = side("8FB3E6")
for name, (c1, c2) in PARAMS.items():
    lc1, lc2, text = labels[name]
    put(P_ROW, lc1, lc2, text, font(9, True, GREY), h="right", indent=1)
    for i in range(c1, c2 + 1):
        cell = ws.cell(P_ROW, gc(i))
        cell.fill = fill(INPUT_FILL)
        cell.border = Border(top=inp_border, bottom=inp_border,
                             left=inp_border if i == c1 else None, right=inp_border if i == c2 else None)
    cell = put(P_ROW, c1, c2, defaults[name], font(10, True, INPUT_INK), h="center", bg=INPUT_FILL,
               fmt={"p_PremierMois": "MMMM YYYY", "p_HeuresETP": '#,##0" h"',
                    "p_ChargeFixe": '#,##0" h"'}.get(name))
    put(P_ROW + 1, c1, c2, help_formula("aide_Dashboard", hints[name][0]), font(7, color=MUTED, italic=True), h="center",
        v="top", fmt=hints[name][1], shrink=True)

dv = DataValidation(type="list", formula1=f'"{PERIM_RUN},{PERIM_ALL}"', allow_blank=False)
dv.add(f"{gl(PARAMS['p_Perimetre'][0])}{P_ROW}")
ws.add_data_validation(dv)
dv = DataValidation(type="list", formula1=f'"{REP_SMOOTH},{REP_START}"', allow_blank=False)
dv.add(f"{gl(PARAMS['p_Repartition'][0])}{P_ROW}")
ws.add_data_validation(dv)
dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                    showErrorMessage=True, error="Saisir une date (le 1er du mois par exemple) ou laisser vide.")
dv.add(f"{gl(PARAMS['p_PremierMois'][0])}{P_ROW}")
ws.add_data_validation(dv)
dv = DataValidation(type="decimal", operator="between", formula1="1", formula2="744", allow_blank=False,
                    showErrorMessage=True, error="Saisir la capacité mensuelle d'une personne en heures (ex. 120).")
dv.add(f"{gl(PARAMS['p_HeuresETP'][0])}{P_ROW}")
ws.add_data_validation(dv)
dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="744", allow_blank=False,
                    showErrorMessage=True, error="Saisir un nombre d'heures par mois (ex. 20).")
dv.add(f"{gl(PARAMS['p_ChargeFixe'][0])}{P_ROW}")
ws.add_data_validation(dv)

# ---- Section 1 : projets en cours
section(8, f'=IF(p_Perimetre="{PERIM_ALL}","TOUS LES PROJETS","PROJETS EN COURS")')
put(8, 24, NG - 1, f'=IF(Calculs!$CD$27>0,"Attention : "&Calculs!$CD$27&" ligne(s) de Charge ont un dessinateur ET une ressource externe (en rouge dans Charge)","")', font(9, True, RED_T), h="right", v="bottom", shrink=True)
CARDS = [(0, 13), (15, 28), (30, 43), (45, 58)]
H_FMT = '#,##0" h"'
title_f = font(9, True, GREY)
big_f = font(24, True, NAVY)
sub_f = font(8.5, color=GREY)
item_f = font(9.5, color=INK)
item_h = font(9, color=GREY)
for idx, (c1, c2) in enumerate(CARDS):
    if idx == 3:
        box(10, 23, c1, c2, bg=NAVY, border=NAVY, top=side(AMBER, "thick"))
    else:
        box(10, 23, c1, c2)
    if idx < 3:
        for i in range(c1 + 1, c2):
            ws.cell(14, gc(i)).border = Border(bottom=side("E5E7EB"))

# Carte 1 : chefs de projet
c1, c2 = CARDS[0]
put(11, c1 + 1, c2 - 1, "CHEFS DE PROJET", title_f)
put(12, c1 + 1, c2 - 1, f"={K['nbChefs']}", big_f, fmt="0")
put(13, c1 + 1, c2 - 1, f'=IF(p_Perimetre="{PERIM_ALL}","sur l\'ensemble des projets","sur les projets en cours")',
    sub_f)
for k in range(8):
    r = 15 + k
    put(r, c1 + 1, c2 - 4, f"=Calculs!BL{4 + k}", item_f)
    put(r, c2 - 3, c2 - 1, f"=Calculs!BM{4 + k}", item_h, h="right", fmt=H_FMT)

# Carte 2 : programmes (projets) suivis
c1, c2 = CARDS[1]
put(11, c1 + 1, c2 - 1, "PROGRAMMES SUIVIS", title_f)
put(12, c1 + 1, c2 - 1, f"={K['nbProj']}", big_f, fmt="0")
put(13, c1 + 1, c2 - 1, "projets distincts · classés par charge", sub_f)
for k in range(8):
    r = 15 + k
    sw = ws.cell(r, gc(c1 + 1))
    sw.fill = fill(SLOT_COLORS[k])
    sw.border = Border(top=side(WHITE, "thick"), bottom=side(WHITE, "thick"), right=side(WHITE, "medium"))
    put(r, c1 + 2, c2 - 4, f"=Calculs!BP{4 + k}", item_f, indent=0)
    put(r, c2 - 3, c2 - 1, f"=Calculs!BQ{4 + k}", item_h, h="right", fmt=H_FMT)
    name_ref = f"{gl(c1 + 2)}{r}"
    ws.conditional_formatting.add(
        f"{gl(c1 + 1)}{r}",
        FormulaRule(formula=[f'OR({name_ref}="",LEFT({name_ref},1)="+")'],
                    fill=PatternFill("solid", start_color=WHITE, end_color=WHITE, bgColor=WHITE)))

# Carte 3 : charge identifiée
c1, c2 = CARDS[2]
put(11, c1 + 1, c2 - 1, "CHARGE IDENTIFIÉE", title_f)
put(12, c1 + 1, c2 - 1, f"={K['total']}", big_f, fmt=H_FMT)
put(13, c1 + 1, c2 - 1, f"={K['totalETP']}", sub_f, fmt='"soit ≈ "0.0" ETP-mois"')
put(15, c1 + 1, c2 - 4, "Réalisé (selon avancement)", item_f)
put(15, c2 - 3, c2 - 1, f"={K['realise']}", font(9.5, True, INK), h="right", fmt=H_FMT)
put(16, c1 + 1, c2 - 4, "Reste à faire", item_f)
put(16, c2 - 3, c2 - 1, f"={K['reste']}", font(9.5, True, INK), h="right", fmt=H_FMT)
put(18, c1 + 1, c2 - 1, "PAR TYPE D'AFFAIRE", font(8, True, MUTED))
for k in range(4):
    r = 19 + k
    put(r, c1 + 1, c2 - 4, f"=Calculs!BT{4 + k}", item_f)
    put(r, c2 - 3, c2 - 1, f"=Calculs!BU{4 + k}", item_h, h="right", fmt=H_FMT)

# Carte 4 : plus grosse affaire (mise en avant)
c1, c2 = CARDS[3]
IDX = K["bigIdx"]
pop_lab = font(8.5, color="AEB9CC")
pop_val = font(9.5, True, WHITE)
put(11, c1 + 1, c2 - 1, "★  PLUS GROS RESTE À FAIRE", font(9, True, AMBER), bg=NAVY)
put(12, c1 + 1, c2 - 1, f'=IF({IDX}="",0,INDEX(Calculs!$FJ${R1}:$FJ${R2},{IDX}))', font(24, True, WHITE),
    fmt=H_FMT, bg=NAVY)
put(13, c1 + 1, c2 - 1, f'=IF({IDX}="",0,INDEX(Calculs!$FH${R1}:$FH${R2},{IDX}))',
    font(8.5, color="C9D3E3"), fmt='"sur "#,##0" h au total pour cette affaire"', bg=NAVY)
for i in range(c1 + 1, c2):
    ws.cell(14, gc(i)).border = Border(bottom=side("3B4A6B"))
POP = [  # (libellé, colonne source, onglet source, type) — affaire = sous-projet regroupé
    ("Projet", "B", "Charge", "text"), ("Sous-projet", "C", "Charge", "text"), ("Produit", "D", "Charge", "text"),
    ("Chef de projet", "A", "Charge", "text"), ("Ressources", "ET", "Calculs", "text"),
    ("Début planifié", "EW", "Calculs", "debut"), ("Date de besoin", "EY", "Calculs", "besoin"),
    ("Avancement", "EU", "Calculs", "pct"),
]
for k, (lab, col, sheet, kind) in enumerate(POP):
    r = 15 + k
    put(r, c1 + 1, c1 + 5, lab, pop_lab, bg=NAVY)
    src = (f"INDEX(Charge!${col}:${col},{IDX}+{R1 - 1})" if sheet == "Charge"
           else f"INDEX({sheet}!${col}${R1}:${col}${R2},{IDX})")
    if kind == "text":
        formula, fmt = f'=IF({IDX}="","—",{src}&"")', None
    elif kind == "debut":
        formula, fmt = f'=IF({IDX}="","",IF({src}>=2958465,"—",{src}))', "DD/MM/YYYY"
    elif kind == "besoin":
        formula, fmt = f'=IF({IDX}="","",IF({src}=0,"—",{src}))', "DD/MM/YYYY"
    else:
        formula, fmt = f'=IF({IDX}="","",{src})', "0%"
    put(r, c1 + 6, c2 - 1, formula, pop_val, fmt=fmt, bg=NAVY, shrink=True)

# ---- Section 2 : indicateurs clés
section(25, "INDICATEURS CLÉS")
TILES = [(i * 12, i * 12 + 10) for i in range(5)]
kpi_val = font(22, True, NAVY)
kpi_month = font(16, True, NAVY)
KPIS = [
    ("MOIS EN SURCHARGE", f"={K['nbSurcharge']}", '0" mois"',
     (f"={K['capEquipe']}", '"Capacité projets : "#,##0" h/mois"'),
     (f'=IF({K["nbSurcharge"]}>0,{K["premSurcharge"]},{K["picCap"]})',
      '[>2]"Dès "MMMM YYYY;"Pic à "0%" de la capacité"')),
    ("PIC DE CHARGE", f"={K['pic']}", H_FMT,
     (f"={K['picMois']}", '"en "MMMM YYYY'),
     (f'={K["picActives"]}&" affaire(s) active(s)  ·  ≈ "&ROUND({K["picETP"]},1)&" ETP"', None)),
    ("AFFAIRES EN RETARD", f"={K['nbRetard']}", "0",
     (f"={K['nbDepasse']}", '"dont "0" date de besoin dépassée"'),
     (f"={K['nbEcart']}", '0" en retard de plus de 15 pts"')),
    ("NOMBRE D'AFFAIRES", f"={K['nbAff']}", "0",
     (f"={K['nbProj']}", '"réparties sur "0" projet(s)"'),
     (f'=IF(p_Perimetre="{PERIM_ALL}","dont "&{K["nbTermines"]}&" ligne(s) terminée(s)",'
      f'{K["nbTermines"]}&" ligne(s) terminée(s) non comptée(s)")', None)),
    ("TAUX D'AVANCEMENT MOYEN", f"={K['avgPond']}", "0%",
     ("pondéré par les heures des affaires", None),
     (f"={K['avgSimple']}", '"Moyenne simple par affaire : "0%')),
]
for (c1, c2), (title, formula, fmt, s1, s2) in zip(TILES, KPIS):
    box(27, 32, c1, c2, top=side(ACCENT, "thick"))
    put(28, c1 + 1, c2 - 1, title, title_f)
    put(29, c1 + 1, c2 - 1, formula, kpi_month if "MMMM" in fmt else kpi_val, fmt=fmt, shrink=True)
    put(30, c1 + 1, c2 - 1, s1[0], sub_f, fmt=s1[1], shrink=True)
    put(31, c1 + 1, c2 - 1, s2[0], sub_f, fmt=s2[1], shrink=True)
# surcharge de l'équipe et affaires en retard : valeur en rouge dès qu'il y en a
for (c1, c2) in (TILES[0], TILES[2]):
    absref = f"${gl(c1 + 1)}$29"
    ws.conditional_formatting.add(f"{gl(c1 + 1)}29", FormulaRule(formula=[f"AND(ISNUMBER({absref}),{absref}>0)"],
                                                              font=Font(color=RED_T, bold=True)))

# ---- Section 3 : graphiques
section(34, "RÉPARTITION DE LA CHARGE")
LP = (0, 30)
RP = (32, 58)
box(36, 55, *LP)
box(36, 55, *RP)
put(37, LP[0] + 1, LP[0] + 16, "CHARGE EN HEURES PAR MOIS ET PAR PROJET", title_f)
CAP_LINE_C = "1F2937"
for i in (LP[0] + 17, LP[0] + 18):
    ws.cell(37, gc(i)).border = Border(bottom=Side(style="mediumDashed", color=CAP_LINE_C))
put(37, LP[0] + 19, LP[1] - 1, f"={K['capEquipe']}", font(8.5, color=INK), h="right",
    fmt='"Capacité projets de l\'équipe : "#,##0" h/mois"', shrink=True)
put(38, LP[0] + 1, LP[0] + 14, f"={K['start']}", sub_f, fmt='"12 mois à partir de "MMMM YYYY')
put(38, LP[0] + 15, LP[1] - 1, f"={K['hors']}", sub_f, h="right",
    fmt='"Hors période ou sans date de début : "#,##0" h";;"Toute la charge est visible sur la période"')
put(37, RP[0] + 1, RP[1] - 1, "RÉPARTITION DE LA CHARGE PAR PROJET", title_f)
put(38, RP[0] + 1, RP[1] - 1, help_formula("aide_Dashboard", "Total au centre · une couleur par projet, commune "
                                                               "aux deux graphiques"), sub_f)

CH_R1, CH_R2 = 39, 54
DN = (33, 45)       # colonnes de grille du donut
# total au centre du donut (le graphique est transparent)
put(45, DN[0], DN[1], f"={K['total']}", font(18, True, NAVY), h="center", v="bottom", fmt='#,##0" h"', r2=46)
put(47, DN[0], DN[1], "au total", font(9, color=GREY), h="center", v="top", r2=48)

# légende commune
LG = {"sw": 47, "name": (48, 52), "h": (53, 55), "p": (56, 57)}
put(42, *LG["name"], "Projet", font(8, True, MUTED))
put(42, *LG["h"], "Heures", font(8, True, MUTED), h="right")
put(42, *LG["p"], "Part", font(8, True, MUTED), h="right")
for k in range(N_SLOTS):
    r = 43 + k
    sw = ws.cell(r, gc(LG["sw"]))
    sw.fill = fill(SLOT_COLORS[k])
    sw.border = Border(top=side(WHITE, "thick"), bottom=side(WHITE, "thick"), right=side(WHITE, "medium"))
    put(r, *LG["name"], f"=Calculs!AT{SLOT_R1 + k}", font(9, color=INK))
    put(r, *LG["h"], f'=IF(Calculs!AT{SLOT_R1 + k}="","",Calculs!AU{SLOT_R1 + k})', font(9, color=GREY),
        h="right", fmt="#,##0")
    put(r, *LG["p"], f'=IF(Calculs!AT{SLOT_R1 + k}="","",Calculs!AV{SLOT_R1 + k})', font(9, color=GREY),
        h="right", fmt="0%")
    ws.conditional_formatting.add(
        f"{gl(LG['sw'])}{r}",
        FormulaRule(formula=[f'{gl(LG["name"][0])}{r}=""'],
                    fill=PatternFill("solid", start_color=WHITE, end_color=WHITE, bgColor=WHITE)))


def text_props(size=800, color="595959", bold=False):
    cp = CharacterProperties(latin=DFont(typeface=FONT), sz=size, b=bold, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def no_fill():
    return GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))


def anchor(c1, r1, c2, r2):
    a = TwoCellAnchor()
    a._from = AnchorMarker(col=gc(c1) - 1, row=r1 - 1, colOff=0, rowOff=0)
    a.to = AnchorMarker(col=gc(c2), row=r2, colOff=0, rowOff=0)
    return a


# Histogramme empilé : heures par mois et par projet + total au-dessus de chaque colonne
bar = BarChart()
bar.type = "col"
bar.grouping = "stacked"
bar.overlap = 100
bar.gapWidth = 55
for k in range(N_SLOTS):
    r = SLOT_R1 + k
    bar.add_data(Reference(ws_calc, min_col=49, max_col=48 + N_MONTHS, min_row=r, max_row=r),
                 from_rows=True, titles_from_data=False)
    s = bar.series[-1]
    s.tx = SeriesLabel(strRef=StrRef(f"Calculs!$AT${r}"))
    s.graphicalProperties = GraphicalProperties(solidFill=SLOT_COLORS[k])
    s.graphicalProperties.line = LineProperties(solidFill=WHITE, w=9525)
cats = Reference(ws_calc, min_col=49, max_col=48 + N_MONTHS, min_row=3, max_row=3)
bar.set_categories(cats)

line = LineChart()
line.add_data(Reference(ws_calc, min_col=49, max_col=48 + N_MONTHS, min_row=TOT_ROW, max_row=TOT_ROW),
              from_rows=True, titles_from_data=False)
ls = line.series[0]
ls.tx = SeriesLabel(strRef=StrRef(f"Calculs!$AT${TOT_ROW}"))
ls.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
ls.marker = Marker(symbol="none")
ls.smooth = False
ls.dLbls = DataLabelList(numFmt='#,##0;;;', dLblPos="t", showVal=True, showLegendKey=False,
                         showCatName=False, showSerName=False, showPercent=False, showBubbleSize=False,
                         showLeaderLines=False)
ls.dLbls.txPr = text_props(800, "374151", True)
line.add_data(Reference(ws_calc, min_col=49, max_col=48 + N_MONTHS, min_row=CAP_ROW, max_row=CAP_ROW),
              from_rows=True, titles_from_data=False)
cs = line.series[1]
cs.tx = SeriesLabel(strRef=StrRef(f"Calculs!$AT${CAP_ROW}"))
cs.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=CAP_LINE_C, w=19050, prstDash="dash"))
cs.marker = Marker(symbol="none")
cs.smooth = False
line.grouping = "standard"
line.set_categories(cats)
bar += line

bar.legend = None
bar.title = None
bar.roundedCorners = False
bar.graphical_properties = no_fill()
bar.plot_area.graphicalProperties = no_fill()
bar.y_axis.delete = False
bar.x_axis.delete = False
bar.y_axis.title = None
bar.y_axis.numFmt = NumFmt(formatCode='#,##0" h"', sourceLinked=False)
bar.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E5E7EB", w=6350))
bar.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
bar.y_axis.txPr = text_props(800, "6B7280")
bar.y_axis.majorTickMark = "none"
bar.x_axis.numFmt = NumFmt(formatCode="mmm yy", sourceLinked=False)
bar.x_axis.auto = False
bar.x_axis.txPr = text_props(800, "6B7280")
bar.x_axis.majorTickMark = "none"
bar.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="BFC6D1", w=9525))
bar.y_axis.scaling.min = 0
bar.anchor = anchor(LP[0] + 1, CH_R1, LP[1] - 1, CH_R2)
ws.add_chart(bar)

# Anneau : répartition par projet (sans légende ni titre ; total affiché dans la cellule sous le trou)
donut = DoughnutChart(holeSize=62)
donut.add_data(Reference(ws_calc, min_col=47, min_row=SLOT_R1, max_row=SLOT_R2), titles_from_data=False)
donut.set_categories(Reference(ws_calc, min_col=46, min_row=SLOT_R1, max_row=SLOT_R2))
ds = donut.series[0]
ds.tx = SeriesLabel(v="Charge (h)")
for k, color in enumerate(SLOT_COLORS):
    pt = DataPoint(idx=k)
    pt.graphicalProperties = GraphicalProperties(solidFill=color)
    pt.graphicalProperties.line = LineProperties(solidFill=WHITE, w=19050)
    ds.dPt.append(pt)
donut.firstSliceAng = 0
donut.varyColors = True
donut.legend = None
donut.title = None
donut.roundedCorners = False
donut.graphical_properties = no_fill()
donut.plot_area.graphicalProperties = no_fill()
donut.layout = Layout(manualLayout=ManualLayout(layoutTarget="inner", xMode="edge", yMode="edge",
                                                x=0.04, y=0.04, w=0.92, h=0.92))
donut.anchor = anchor(DN[0], CH_R1, DN[1], CH_R2)
ws.add_chart(donut)

# ---- Section 4 : liste des affaires
section(57, "LISTE DES AFFAIRES (un sous-projet par ligne, toutes ressources regroupées, par charge décroissante)")
TCOLS = [  # (titre, nb colonnes de grille, alignement, format)
    ("N°", 2, "center", "0"),
    ("Projet", 8, "left", None),
    ("Sous-projet", 8, "left", None),
    ("Produit", 8, "left", None),
    ("Chef de projet", 7, "left", None),
    ("Ressources", 4, "center", None),
    ("Charge (h)", 5, "right", "#,##0"),
    ("% de la charge totale", 6, "right", "0.0%"),
    ("Avancement", 6, "right", "0%"),
    ("Date de début planifiée", 5, "center", "DD/MM/YYYY"),
]
spans = []
c = 0
for title, w, al, fmt in TCOLS:
    spans.append((c, c + w - 1))
    put(59, c, c + w - 1, title, font(9, True, WHITE), h="center", bg=NAVY, wrap=True)
    for i in range(c, c + w):
        ws.cell(59, gc(i)).fill = fill(NAVY)
    c += w
assert c == NG

row_line = side("E5E7EB")
for k in range(1, N_TABLE + 1):
    r = TABLE_R1 + k - 1
    idx = f"Calculs!$BX${3 + k}"
    for i in range(NG):
        cell = ws.cell(r, gc(i))
        cell.fill = fill(BG)
    vals = [
        f'=IF({idx}="","",Calculs!$BW${3 + k})',
        f'=IF({idx}="","",INDEX(Charge!$B:$B,{idx}+{R1 - 1})&"")',
        f'=IF({idx}="","",INDEX(Charge!$C:$C,{idx}+{R1 - 1})&"")',
        f'=IF({idx}="","",INDEX(Charge!$D:$D,{idx}+{R1 - 1})&"")',
        f'=IF({idx}="","",INDEX(Charge!$A:$A,{idx}+{R1 - 1})&"")',
        f'=IF({idx}="","",INDEX(Calculs!$ET${R1}:$ET${R2},{idx}))',
        f'=IF({idx}="","",INDEX(Calculs!$EP${R1}:$EP${R2},{idx}))',
        f'=IF({idx}="","",IF({K["total"]}=0,0,INDEX(Calculs!$EP${R1}:$EP${R2},{idx})/{K["total"]}))',
        f'=IF({idx}="","",INDEX(Calculs!$EU${R1}:$EU${R2},{idx}))',
        f'=IF({idx}="","",IF(INDEX(Calculs!$EW${R1}:$EW${R2},{idx})>=2958465,"à planifier",'
        f'INDEX(Calculs!$EW${R1}:$EW${R2},{idx})))',
    ]
    for (title, w, al, fmt), (s1, s2), formula in zip(TCOLS, spans, vals):
        bold = title in ("Projet",)
        put(r, s1, s2, formula, font(9, bold, INK if title != "N°" else GREY), h=al, fmt=fmt,
            indent=1 if al != "center" else 0, shrink=title == "Ressources")

first_cell = f"{gl(0)}{TABLE_R1}"
tbl_range = f"{gl(0)}{TABLE_R1}:{gl(NG - 1)}{TABLE_R2}"
ws.conditional_formatting.add(
    tbl_range,
    FormulaRule(formula=[f'AND(${gl(0)}{TABLE_R1}<>"",MOD(ROW(),2)=0)'], border=Border(bottom=row_line),
                fill=PatternFill("solid", start_color="F5F7FA", end_color="F5F7FA", bgColor="F5F7FA")))
ws.conditional_formatting.add(
    tbl_range,
    FormulaRule(formula=[f'${gl(0)}{TABLE_R1}<>""'], border=Border(bottom=row_line),
                fill=PatternFill("solid", start_color=WHITE, end_color=WHITE, bgColor=WHITE)))
pct_col = gl(spans[7][0])
adv_col = gl(spans[8][0])
assert TCOLS[7][0].startswith("%") and TCOLS[8][0] == "Avancement"
ws.conditional_formatting.add(f"{pct_col}{TABLE_R1}:{pct_col}{TABLE_R2}",
                              DataBarRule(start_type="num", start_value=0, end_type="max",
                                          color="86B6EF", showValue=True))
ws.conditional_formatting.add(f"{adv_col}{TABLE_R1}:{adv_col}{TABLE_R2}",
                              DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                                          color="7FD1AE", showValue=True))
# affaire en retard (Calculs!BY, même ordre que le tableau) : avancement en rouge
ws.conditional_formatting.add(f"{adv_col}{TABLE_R1}:{adv_col}{TABLE_R2}",
                              FormulaRule(formula=["AND(ISNUMBER(Calculs!$BY4),Calculs!$BY4>0)"],
                                          font=Font(color=RED_T, bold=True)))

NOTE_R = TABLE_R2 + 2
put(NOTE_R, 0, NG - 1,
    f'=TRIM(IF({K["nbAff"]}>{N_TABLE},"Attention : seules les {N_TABLE} premières affaires (par charge) sont '
    f'listées sur "&{K["nbAff"]}&".  ","")&IF(Calculs!$CD$27>0,"Attention : "&Calculs!$CD$27&" ligne(s) de '
    f'l\'onglet Charge ont à la fois un dessinateur (K) et une ressource externe (L) : une seule ressource par '
    f'ligne.",""))', font(9, True, "B45309"))
NOTES = [
    "Affaire = un sous-projet (projet + sous-projet + produit) : ses lignes sont additionnées ; chef = celui de sa "
    "1re ligne ; avancement = toutes ses lignes pondérées par leurs heures ; reste à faire = heures × (1 − avancement).",
    "En cours = ligne dont l'avancement (O) est inférieur à 100 % : une ressource qui a terminé sa part sort du "
    "périmètre « Projets en cours », même si d'autres ressources continuent l'affaire.",
    "Charge par mois : les heures (M) sont réparties au prorata des jours ouvrés (lundi au vendredi, hors jours "
    "non travaillés de l'onglet Data) entre le début planifié (N) et la date de besoin (F) ; sans date de besoin "
    "valable, tout va sur le mois de début.",
    "Affaire en retard (avancement en rouge) : date de besoin dépassée sans être à 100 %, ou avancement inférieur de "
    "plus de 15 points à l'avancement attendu au prorata des jours ouvrés écoulés depuis le début planifié.",
    "Pointillés : capacité projets de l'équipe = somme des (capacité − charge fixe) des ressources des listes Data.  "
    "ETP = heures ÷ (capacité − charge fixe) d'une personne.  Au-delà de 9 projets : « Autres projets » (gris).",
]
for i, t in enumerate(NOTES):
    put(NOTE_R + 1 + i, 0, NG - 1, help_formula("aide_Dashboard", t), font(8, color=GREY, italic=True))

ws.print_area = f"A1:{L(gc(NG))}{LAST_ROW}"
auto_hide_rows("Dashboard", f"${gl(0)}${TABLE_R1}:${gl(0)}${TABLE_R2}",
               f"${gl(0)}${NOTE_R}:${gl(0)}${NOTE_R + len(NOTES)}")
macro_warning(ws, MACRO_WARNING_CELLS)
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A3
ws.page_setup.scale = 90
ws.row_breaks.append(Break(id=56))
ws.page_margins.left = ws.page_margins.right = 0.3
ws.page_margins.top = ws.page_margins.bottom = 0.4
ws.print_title_rows = "1:1"                 # ligne de marge répétée en haut de chaque page
ws.print_options.horizontalCentered = True
ws.sheet_properties.tabColor = ACCENT

# ---------------------------------------------------------------- Dessinateurs
ws = ws_dess
ws.sheet_view.showGridLines = False
ws.sheet_view.zoomScale = 90
ws.sheet_properties.tabColor = "1BAF7A"
ws.column_dimensions["A"].width = MARGIN_W
for i in range(NG):
    ws.column_dimensions[gl(i)].width = 2.71
ws.column_dimensions[L(gc(NG))].width = MARGIN_W

D_ROW_H = {1: MARGIN_H, 2: 30, 3: 16, 4: 16, 5: 8, 6: 22, 7: 14, 8: 24, 9: 6,
           10: 6, 11: 18, 12: 34, 13: 15, 14: 15, 15: 8, 16: 14, 17: 20, 18: 6,
           19: 6, 20: 18, 21: 16, 38: 6, 39: 18, 46: 8, 47: 14, 48: 20, 49: 6, 50: 30,
           66: 14, 67: 14, 68: 20, 69: 6, 70: 6, 71: 15, 72: 18, 91: 8, 92: 14}
for r in range(22, 38):
    D_ROW_H[r] = 15
for r in range(40, 46):
    D_ROW_H[r] = 16
for r in range(51, 66):
    D_ROW_H[r] = 16.5
for r in range(73, 91):
    D_ROW_H[r] = 16
D_LAST = 97   # dernière ligne = marge basse
D_ROW_H[D_LAST] = MARGIN_H
for r, h in D_ROW_H.items():
    ws.row_dimensions[r].height = h
for r in range(1, D_LAST + 1):
    for c in range(1, gc(NG) + 1):
        ws.cell(r, c).fill = fill(BG)

banner("CHARGE PAR RESSOURCE (DESSINATEURS ET EXTERNES)",
       "Charge fixe hors projet + heures de ses lignes dans l'onglet Charge, comparées à la capacité mensuelle de la "
       "personne.", "edite_Dessinateurs",
       left4=f'=IF({DS["nom"]}="","",{DS["nom"]}&IF({DS["type"]}="Externe"," (ressource externe)","")&"  ·  ")'
             f'&"Périmètre : "&{DS["perim"]}')
# la barre de choix de la ressource (lignes 5 à 7) n'est pas imprimée par « Éditer PDF » : rappel en ligne 4
wb.defined_names["lignes_pdf_Dessinateurs"] = DefinedName("lignes_pdf_Dessinateurs",
                                                          attr_text="Dessinateurs!$A$5:$A$7")

# barre (lignes 6-7, comme sur le Dashboard) : choix de la ressource + rappel de ses réglages
B = P_ROW
box(B, B + 1, 0, NG - 1)
put(B, 0, 4, "Ressource", font(9, True, GREY), h="right", indent=1)
for i in range(5, 17):
    cell = ws.cell(B, gc(i))
    cell.fill = fill(INPUT_FILL)
    cell.border = Border(top=inp_border, bottom=inp_border,
                         left=inp_border if i == 5 else None, right=inp_border if i == 16 else None)
put(B, 5, 16, LISTS[4][3][0], font(11, True, INPUT_INK), h="center", bg=INPUT_FILL)
wb.defined_names["p_Dessinateur"] = DefinedName("p_Dessinateur", attr_text=f"Dessinateurs!${gl(5)}${B}")
dv = DataValidation(type="list", formula1="lst_Ressources", allow_blank=True, showErrorMessage=True,
                    errorTitle="Ressource inconnue",
                    error="Choisir un nom de la liste (onglet Data : dessinateurs colonne E, externes colonne H).")
dv.add(f"{gl(5)}{B}")
ws.add_data_validation(dv)
put(B + 1, 5, 16, f'=IF({DS["type"]}="Hors liste","Absent des listes Data : réglages par défaut",'
             f'IF(aide_Dessinateurs<>"Oui","",IF({DS["type"]}="Externe","Ressource externe (Data, colonne H)",'
             f'"Dessinateur interne (Data, colonne E)")))', font(7, color=MUTED, italic=True), h="center", v="top",
    shrink=True)
ws.conditional_formatting.add(f"{gl(5)}{B + 1}", FormulaRule(formula=[f'LEFT(${gl(5)}${B + 1},7)="Absent "'],
                                                       font=Font(color=RED_T, bold=True, italic=False)))
info_lab = font(9, True, GREY)
info_val = font(10, True, NAVY)
put(B, 18, 21, "Capacité", info_lab, h="right", indent=1)
put(B, 22, 25, f"={DS['cap']}", info_val, fmt='#,##0" h"')
put(B, 27, 31, "Charge fixe", info_lab, h="right", indent=1)
put(B, 32, 34, f"={DS['fixe']}", info_val, fmt='#,##0" h"')
put(B, 36, 42, "Dispo. projets", info_lab, h="right", indent=1)
put(B, 43, 46, f"={DS['capProj']}", info_val, fmt='#,##0" h/mois"')
put(B + 1, 18, 46, help_formula("aide_Dessinateurs", "par mois · valeurs propres à la personne (Data, colonnes F-G ou "
                                                  "I-J) ou valeurs par défaut du Dashboard"),
    font(7, color=MUTED, italic=True), h="center", v="top", shrink=True)
put(B, 48, 57, f"={DS['debut']}", font(9, True, GREY), h="right", fmt='"12 mois dès "MMMM YYYY')
put(B + 1, 48, 57, f'="Périmètre : "&{DS["perim"]}', font(7, color=MUTED, italic=True), h="right", v="top")

# ---- synthèse
section(8, f'=IF({DS["nom"]}="","SYNTHÈSE — choisir une ressource","SYNTHÈSE — "&UPPER({DS["nom"]})'
            f'&IF({DS["type"]}="Externe","  (RESSOURCE EXTERNE)",""))')
put(8, 30, NG - 1, f'=IF(Calculs!$CD$27>0,"Attention : "&Calculs!$CD$27&" ligne(s) de Charge ont un dessinateur ET une ressource externe (en rouge dans Charge)","")', font(9, True, RED_T), h="right", v="bottom", shrink=True)
D_KPIS = [
    ("CAPACITÉ MENSUELLE", f"={DS['cap']}", H_FMT,
     (f"={DS['fixe']}", '"dont "#,##0" h de charge fixe"'),
     (f'=IF(AND({DS["nom"]}<>"",{DS["fixe"]}>={DS["cap"]}),"Attention : charge fixe ≥ capacité",{DS["capProj"]})',
      '"soit "#,##0" h/mois pour les projets"')),
    ("CHARGE PROJETS AFFECTÉE", f"={DS['chargeTot']}", H_FMT,
     (f"={DS['nbAff']}", '"sur "0" affaire(s)"'),
     (f"={DS['periode']}", '"dont "#,##0" h sur 12 mois"')),
    ("MOIS LE PLUS CHARGÉ", f'=IF(OR({DS["nom"]}="",{DS["periode"]}=0),"—",{DS["picMois"]})', "MMMM YYYY",
     (f"={DS['pic']}", '#,##0" h au total (fixe + projets)"'),
     (f"={DS['picOcc']}", '0%" de la capacité"')),
    ("MOIS EN SURCHARGE", f"={DS['nbSur']}", '0" mois"',
     (f"={DS['depass']}", '"Dépassement cumulé : "#,##0" h"'),
     (f'=IF({DS["nbSur"]}=0,"Aucun dépassement prévu",{DS["premSur"]})', '"Dès "MMMM YYYY')),
    ("OCCUPATION MOYENNE", f"={DS['occMoy']}", "0%",
     (f"={DS['dispo']}", '#,##0" h encore disponibles"'),
     ("sur les 12 mois affichés", None)),
]
for (c1, c2), (title, formula, fmt, s1, s2) in zip(TILES, D_KPIS):
    box(10, 15, c1, c2, top=side(ACCENT, "thick"))
    put(11, c1 + 1, c2 - 1, title, title_f)
    put(12, c1 + 1, c2 - 1, formula, kpi_month if "MMMM" in fmt else kpi_val, fmt=fmt, shrink=True)
    put(13, c1 + 1, c2 - 1, s1[0], sub_f, fmt=s1[1], shrink=True)
    put(14, c1 + 1, c2 - 1, s2[0], sub_f, fmt=s2[1], shrink=True)
# alertes (formules limitées à l'onglet) : mois le plus chargé > 100 %, mois en surcharge, charge fixe >= capacité
SEL = f"TRIM(${gl(5)}${B})"
occ_ref = f"${gl(TILES[2][0] + 1)}$14"
for r in (12, 14):
    ws.conditional_formatting.add(f"{gl(TILES[2][0] + 1)}{r}",
                                  FormulaRule(formula=[f"AND(ISNUMBER({occ_ref}),{occ_ref}>1)"],
                                              font=Font(color=RED_T, bold=True)))
nb_ref = f"${gl(TILES[3][0] + 1)}$12"
ws.conditional_formatting.add(f"{gl(TILES[3][0] + 1)}12",
                              FormulaRule(formula=[f"AND(ISNUMBER({nb_ref}),{nb_ref}>0)"], font=Font(color=RED_T)))
ws.conditional_formatting.add(f"{gl(TILES[3][0] + 1)}12",
                              FormulaRule(formula=[f'AND({SEL}<>"",{nb_ref}=0)'], font=Font(color=GREEN_T)))
warn_ref = f"${gl(TILES[0][0] + 1)}$14"
ws.conditional_formatting.add(f"{gl(TILES[0][0] + 1)}14",
                              FormulaRule(formula=[f'LEFT({warn_ref},9)="Attention"'],
                                          font=Font(color=RED_T, bold=True)))

# ---- charge mensuelle : graphique + tableau
section(17, "CHARGE MENSUELLE")
box(19, 46, 0, NG - 1)
put(20, 1, 40, f'=IF({DS["nom"]}="","CHARGE MENSUELLE (HEURES)","CHARGE MENSUELLE DE "&UPPER({DS["nom"]})&" (HEURES)")',
    title_f)
FIX_C, PROJ_C, OVER_C, CAP_C = "B8C0CC", "2A78D6", "D03B3B", "1F2937"
legend = [(1, FIX_C, "Charge fixe hors projet"), (12, PROJ_C, "Charge projets"),
          (23, OVER_C, "Dépassement de capacité")]
for c, color, text in legend:
    sw = ws.cell(21, gc(c))
    sw.fill = fill(color)
    sw.border = Border(top=side(WHITE, "thick"), bottom=side(WHITE, "thick"), right=side(WHITE, "medium"))
    put(21, c + 1, c + 9, text, font(8.5, color=INK))
for i in (35, 36):
    ws.cell(21, gc(i)).border = Border(bottom=Side(style="mediumDashed", color=CAP_C))
put(21, 37, 47, "Capacité mensuelle", font(8.5, color=INK))

DCH_R1, DCH_R2 = 22, 37
dbar = BarChart()
dbar.type = "col"
dbar.grouping = "stacked"
dbar.overlap = 100
dbar.gapWidth = 55
for col, color, title in (("CH", FIX_C, "Charge fixe hors projet"), ("CL", PROJ_C, "Charge projets"),
                          ("CM", OVER_C, "Dépassement de capacité")):
    dbar.add_data(Reference(ws_calc, range_string=f"Calculs!${col}$4:${col}${3 + N_MONTHS}"),
                  titles_from_data=False)
    sr = dbar.series[-1]
    sr.tx = SeriesLabel(v=title)
    sr.graphicalProperties = GraphicalProperties(solidFill=color)
    sr.graphicalProperties.line = LineProperties(solidFill=WHITE, w=9525)
d_cats = Reference(ws_calc, range_string=f"Calculs!$CG$4:$CG${3 + N_MONTHS}")
dbar.set_categories(d_cats)
dline = LineChart()
dline.add_data(Reference(ws_calc, range_string=f"Calculs!$CK$4:$CK${3 + N_MONTHS}"), titles_from_data=False)
cap_s = dline.series[0]
cap_s.tx = SeriesLabel(v="Capacité mensuelle")
cap_s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=CAP_C, w=19050, prstDash="dash"))
cap_s.marker = Marker(symbol="none")
cap_s.smooth = False
dline.set_categories(d_cats)
dbar += dline
dbar.legend = None
dbar.title = None
dbar.roundedCorners = False
dbar.graphical_properties = no_fill()
dbar.plot_area.graphicalProperties = no_fill()
dbar.y_axis.delete = False
dbar.x_axis.delete = False
dbar.y_axis.numFmt = NumFmt(formatCode='#,##0" h"', sourceLinked=False)
dbar.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E5E7EB", w=6350))
dbar.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
dbar.y_axis.txPr = text_props(800, "6B7280")
dbar.y_axis.majorTickMark = "none"
dbar.y_axis.scaling.min = 0
dbar.x_axis.numFmt = NumFmt(formatCode="mmm yy", sourceLinked=False)
dbar.x_axis.auto = False
dbar.x_axis.txPr = text_props(800, "6B7280")
dbar.x_axis.majorTickMark = "none"
dbar.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="BFC6D1", w=9525))
dbar.anchor = anchor(1, DCH_R1, NG - 2, DCH_R2)
ws.add_chart(dbar)

# tableau mensuel sous le graphique (mêmes mois que le graphique)
MT = [(40, "Charge fixe hors projet", "CH", False), (41, "Charge projets", "CI", False),
      (42, "Charge totale", "CJ", True), (43, "Capacité", "CK", False),
      (44, "Marge (+) / dépassement (−)", "CN", False), (45, "Occupation", "CO", True)]
for i in range(1, NG - 1):
    ws.cell(39, gc(i)).border = Border(bottom=side("B8C2CF"))
put(39, 1, 9, "Heures par mois", font(8, True, MUTED))
for m in range(N_MONTHS):
    c1 = 10 + 4 * m
    put(39, c1, c1 + 3, f"=Calculs!CG{4 + m}", font(8, True, GREY), h="right", fmt="mmm yy")
for r, label, col, bold in MT:
    put(r, 1, 9, label, font(9, bold, INK))
    for m in range(N_MONTHS):
        c1 = 10 + 4 * m
        fmt = {"CO": "0%", "CN": '+#,##0;-#,##0;0'}.get(col, "#,##0")
        put(r, c1, c1 + 3, f"=Calculs!{col}{4 + m}", font(9, bold, INK if col != "CK" else GREY), h="right",
            fmt=fmt)
    if r == 42:
        for i in range(1, NG - 1):
            ws.cell(r, gc(i)).border = Border(top=side("D6DCE4"))
row_rng = lambda r: f"{gl(10)}{r}:{gl(10 + 4 * N_MONTHS - 1)}{r}"
first = gl(10)
ws.conditional_formatting.add(row_rng(44), FormulaRule(formula=[f"AND(ISNUMBER({first}44),{first}44<0)"],
                                                       font=Font(color=RED_T, bold=True)))
ws.conditional_formatting.add(row_rng(44), FormulaRule(formula=[f"AND(ISNUMBER({first}44),{first}44>0)"],
                                                       font=Font(color=GREEN_T)))
OCC_RULES = [  # (condition, fond, texte) — occupation calculée sur le total arrondi à l'heure
    ("{c}>1", "F8D7D7", RED_T),
    ("{c}>=0.9", "FFF1CC", "8A5A00"),
    ("{c}>=0", "E3F4E8", GREEN_T),
]


def occ_cf(rng_txt, ref):
    for cond, bg, fg in OCC_RULES:
        ws.conditional_formatting.add(
            rng_txt, FormulaRule(formula=[f"AND(ISNUMBER({ref}),{cond.format(c=ref)})"], stopIfTrue=True,
                                 fill=PatternFill("solid", start_color=bg, end_color=bg, bgColor=bg),
                                 font=Font(color=fg, bold=True)))


occ_cf(row_rng(45), f"{first}45")

# ---- affaires du dessinateur
section(48, f'=IF({DS["nom"]}="","AFFAIRES DE LA RESSOURCE","AFFAIRES DE "&UPPER({DS["nom"]})&" (par charge décroissante)")')
D_TCOLS = [("N°", 2, "center", "0", "DX"), ("Projet", 9, "left", None, "DZ"), ("Sous-projet", 9, "left", None, "EA"),
           ("Produit", 9, "left", None, "EB"), ("Chef de projet", 7, "left", None, "EC"),
           ("Charge (h)", 5, "right", "#,##0", "ED"), ("Heures sur la période", 6, "right", "#,##0", "EE"),
           ("Début planifié", 4, "center", "DD/MM/YY", "EF"), ("Date de besoin", 4, "center", "DD/MM/YY", "EG"),
           ("Avancement", 4, "right", "0%", "EH")]
d_spans = []
c = 0
for title, w, al, fmt, _ in D_TCOLS:
    d_spans.append((c, c + w - 1))
    for i in range(c, c + w):
        ws.cell(50, gc(i)).fill = fill(NAVY)
    put(50, c, c + w - 1, title, font(9, True, WHITE), h="center", bg=NAVY, wrap=True)
    c += w
assert c == NG
for k in range(1, N_DESS_AFF + 1):
    r = 50 + k
    for (title, w, al, fmt, col), (s1, s2) in zip(D_TCOLS, d_spans):
        src = f"Calculs!{col}{3 + k}"
        formula = f'=IF(Calculs!$DY${3 + k}="","",{src})'
        put(r, s1, s2, formula, font(9, title == "Projet", INK if title != "N°" else GREY), h=al, fmt=fmt,
            indent=1 if al != "center" else 0)
d_tbl = f"{gl(0)}51:{gl(NG - 1)}{50 + N_DESS_AFF}"
ws.conditional_formatting.add(d_tbl, FormulaRule(formula=[f'AND(${gl(0)}51<>"",MOD(ROW(),2)=0)'],
                                                 border=Border(bottom=row_line),
                                                 fill=PatternFill("solid", start_color="F5F7FA", end_color="F5F7FA",
                                                                  bgColor="F5F7FA")))
ws.conditional_formatting.add(d_tbl, FormulaRule(formula=[f'${gl(0)}51<>""'], border=Border(bottom=row_line),
                                                 fill=PatternFill("solid", start_color=WHITE, end_color=WHITE,
                                                                  bgColor=WHITE)))
put(66, 0, NG - 1, f'=IF({DS["nbLignes"]}>{N_DESS_AFF},"Seules les {N_DESS_AFF} plus grosses lignes sont listées sur "'
                   f'&{DS["nbLignes"]}&".",IF(AND({DS["nom"]}<>"",{DS["nbLignes"]}=0),"Aucune affaire affectée à '
                   f'cette ressource dans le périmètre.",""))', font(9, True, "B45309"))

# ---- occupation de l'équipe
section(68, "OCCUPATION DE L'ÉQUIPE (charge fixe + projets, en % de la capacité)")
box(70, 91, 0, NG - 1)
put(71, 1, 57, "Vert : moins de 90 %   ·   Orange : de 90 à 100 %   ·   Rouge : au-delà de 100 % (surcharge)   ·   "
               f"les {N_DESS} premières ressources des listes Data (dessinateurs puis externes)",
    font(8, color=GREY, italic=True))
for i in range(1, NG - 1):
    ws.cell(72, gc(i)).border = Border(bottom=side("B8C2CF"))
put(72, 1, 7, "Ressource", font(8, True, MUTED))
for m in range(N_MONTHS):
    c1 = 10 + 4 * m
    put(72, c1, c1 + 3, f"=Calculs!{TEAM_H[m]}3", font(8, True, GREY), h="center", fmt="mmm yy")
for k in range(N_DESS):
    r = 73 + k
    cr = TEAM_R1 + k
    put(r, 1, 7, f"=Calculs!{TEAM_NAME}{cr}", font(9, False, INK), shrink=True)
    put(r, 8, 9, f'=IF(Calculs!{TEAM_NAME}{cr}="","",IF(Calculs!FB{cr}="Externe","ext.",""))',
        font(7.5, False, MUTED, italic=True), h="center")
    for m in range(N_MONTHS):
        c1 = 10 + 4 * m
        put(r, c1, c1 + 3, f"=Calculs!{TEAM_O[m]}{cr}", font(9, True, INK), h="center", fmt="0%")
team_rng = f"{gl(10)}73:{gl(10 + 4 * N_MONTHS - 1)}{72 + N_DESS}"
occ_cf(team_rng, f"{gl(10)}73")
ws.conditional_formatting.add(f"{gl(1)}73:{gl(7)}{72 + N_DESS}",
                              FormulaRule(formula=[f'AND(${gl(1)}73<>"",${gl(1)}73={SEL})'],
                                          font=Font(bold=True, color=ACCENT),
                                          fill=PatternFill("solid", start_color="EAF2FD", end_color="EAF2FD",
                                                           bgColor="EAF2FD")))
for j, (cr, label) in enumerate(((NA_ROW, "Charge non affectée (h)"),
                                 (OTHER_ROW, "Autres ressources (h)"))):
    r = 73 + N_DESS + j
    put(r, 1, 9, label, font(8.5, False, GREY, italic=True))
    for m in range(N_MONTHS):
        c1 = 10 + 4 * m
        put(r, c1, c1 + 3, f"=Calculs!{TEAM_H[m]}{cr}", font(8.5, False, GREY), h="center", fmt='#,##0;;"–"')
    if j == 0:
        for i in range(1, NG - 1):
            ws.cell(r, gc(i)).border = Border(top=side("D6DCE4"))

D_NOTES = [
    "Charge d'une ressource = charge fixe hors projet + heures des lignes de l'onglet Charge où elle figure (colonne K "
    "pour un dessinateur, L pour une ressource externe), réparties par mois comme sur le Dashboard.",
    "Une ligne = une ressource : plusieurs personnes sur un même sous-projet = plusieurs lignes, chacune avec ses "
    "heures.  Capacité et charge fixe : Data (F-G ou I-J), sinon valeurs par défaut du Dashboard.",
    f'=IF({DS["nbDouble"]}>0,"Attention : "&{DS["nbDouble"]}&" ligne(s) de l\'onglet Charge ont à la fois un '
    f'dessinateur (K) et une ressource externe (L) : seule la colonne K est comptée. Faire une ligne par ressource.",'
    f'"")',
    f'=IF({DS["nbListe"]}+{DS["nbExt"]}>{N_DESS},"Attention : les listes Data contiennent "&'
    f'({DS["nbListe"]}+{DS["nbExt"]})&" ressources, seules les {N_DESS} premières apparaissent dans le tableau '
    f'd\'occupation (le choix en haut de page fonctionne pour toutes).","")',
]
for i, t in enumerate(D_NOTES):
    put(93 + i, 0, NG - 1, help_formula("aide_Dessinateurs", t) if i < 2 else t,
        font(8, color=GREY if i < 2 else "B45309", italic=i < 2, bold=i >= 2))

ws.print_area = f"A1:{L(gc(NG))}{D_LAST}"
auto_hide_rows("Dessinateurs", f"${gl(0)}$51:${gl(0)}$66", f"${gl(1)}$73:${gl(1)}${72 + N_DESS}",
               f"${gl(0)}$93:${gl(0)}${92 + len(D_NOTES)}")
macro_warning(ws, MACRO_WARNING_CELLS)
# Charge et Data : ligne d'aide (ligne 2) repliée quand l'aide est masquée
auto_hide_rows("Charge", "$A$2")
auto_hide_rows("Data", "$A$2")
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A3
ws.page_setup.scale = 90
ws.row_breaks.append(Break(id=47))
ws.page_margins.left = ws.page_margins.right = 0.3
ws.page_margins.top = ws.page_margins.bottom = 0.4
ws.print_title_rows = "1:1"                 # ligne de marge répétée en haut de chaque page
ws.print_options.horizontalCentered = True

# contrôles : aucune constante texte de plus de 255 caractères dans une formule (limite Excel), aucune infobulle
for _ws in wb.worksheets:
    for _row in _ws.iter_rows():
        for _c in _row:
            if isinstance(_c.value, str) and _c.value.startswith("="):
                for _lit in re.findall(r'"((?:[^"]|"")*)"', _c.value):
                    assert len(_lit) <= 255, (_ws.title, _c.coordinate, len(_lit))
assert not any(_dv.showInputMessage for _ws in wb.worksheets for _dv in _ws.data_validations.dataValidation)

wb.active = wb.sheetnames.index("Dashboard")
for sheet in wb.worksheets:
    sheet.sheet_view.tabSelected = sheet.title == "Dashboard"
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("OK ->", OUT)

# -*- coding: utf-8 -*-
"""Ajoute aux formules du classeur généré les valeurs calculées par LibreOffice.

openpyxl écrit les formules sans résultat : Excel les calcule à l'ouverture, mais un aperçu
ou le « Mode protégé » d'Excel afficheraient des cellules vides. Ce script recopie les
résultats d'une copie recalculée (recalc.xlsx) dans le fichier d'origine, sans toucher
au reste (graphiques, mises en forme, validations).

Usage : python scripts/inject_cached_values.py source.xlsx recalcule.xlsx sortie.xlsx
"""
import posixpath
import re
import sys
import zipfile
from datetime import date, datetime, time
from xml.sax.saxutils import escape

from openpyxl import load_workbook

src, calc, out = sys.argv[1:4]

values = load_workbook(calc, data_only=True)
EPOCH = datetime(1899, 12, 30)


def serial(v):
    if isinstance(v, datetime):
        d = v - EPOCH
        return d.days + d.seconds / 86400
    if isinstance(v, date):
        return (v - EPOCH.date()).days
    if isinstance(v, time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    return v


with zipfile.ZipFile(src) as zin:
    wb_xml = zin.read("xl/workbook.xml").decode()
    rels = zin.read("xl/_rels/workbook.xml.rels").decode()
    rid_target = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
    rid_target.update({k: v for v, k in re.findall(r'Target="([^"]+)"[^>]*Id="(rId\d+)"', rels)})
    sheet_files = {}
    for name, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="(rId\d+)"', wb_xml):
        target = rid_target[rid]
        path = target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target)
        sheet_files[path] = name

    cell_re = re.compile(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v></v></c>', re.S)
    stats = {}

    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            sheet = sheet_files.get(item.filename)
            if sheet:
                ws = values[sheet]

                def fill(m):
                    ref, attrs, formula = m.groups()
                    attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
                    v = serial(ws[ref].value)
                    if v is None:
                        v = ""
                    if isinstance(v, bool):
                        return f'<c r="{ref}"{attrs} t="b"><f>{formula}</f><v>{int(v)}</v></c>'
                    if isinstance(v, (int, float)):
                        return f'<c r="{ref}"{attrs}><f>{formula}</f><v>{repr(float(v)) if isinstance(v, float) else v}</v></c>'
                    if isinstance(v, str) and v.startswith("#"):
                        return f'<c r="{ref}"{attrs} t="e"><f>{formula}</f><v>{escape(v)}</v></c>'
                    return f'<c r="{ref}"{attrs} t="str"><f>{formula}</f><v>{escape(str(v))}</v></c>'

                text = data.decode()
                text, count = cell_re.subn(fill, text)
                stats[sheet] = count
                data = text.encode()
            zout.writestr(item, data)

print("Valeurs injectées :", stats)

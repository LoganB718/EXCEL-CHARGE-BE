# -*- coding: utf-8 -*-
"""Preuve du modèle Power BI par son miroir Python : comparaison aux valeurs calculées par Excel.

Pour chaque classeur recalculé (build12/final.xlsx, les 13 scénarios build12/sc_<cas>.xlsx, build13/final.xlsx qui
porte les noms bi_*, et les scénarios de variantes de scen_variantes.py) :
  1. intégrité des tables produites par le miroir des requêtes M (colonnes déclarées, types, clés des relations,
     colonnes de tri) ;
  2. variante « Comme dans Excel » sur les 12 mois du Dashboard Excel, comparée à Excel : mois et charges
     (Calculs AO:AQ), indicateurs AL4:AL40, table Contrôle Excel (statut ✔ partout), lignes de Calculs (charge,
     rangs, ressource, répartition mensuelle W:AI), affaires (Calculs EM:FL), tableau des affaires du Dashboard,
     carte du plus gros reste, tuiles, cartes, textes, listes des chefs / projets / types, groupes de couleur des
     projets et histogramme mensuel (AS:BH), onglet Dessinateurs (occupation de l'équipe CS:DV, ressource
     sélectionnée CC:CQ, ses lignes DX:EH, tuiles) ;
  3. variantes des segments : le miroir d'un classeur avec une variante explicite (« Tous les projets »,
     « Tout sur le mois de début »...) doit reproduire Excel recalculé avec ce réglage (mêmes données).
La date du jour (AUJOURDHUI d'Excel, DateDuJour du modèle) est la date (UTC) de l'enregistrement du classeur
recalculé. Tolérance 0,01 h ; aucun écart admis.

Usage :
    python3 test_mirror.py --racine <dossier> [--dashboard build_dashboard.py] [--detail]
        suite complète : <dossier>/build12 (final.xlsx, sc_<cas>.xlsx), <dossier>/build13/final.xlsx,
        <dossier>/pbi/out/scen_variantes_{1,3,4} et <dossier>/pbi/review_calculs/scen_lo (scénarios de relecture) ;
    python3 test_mirror.py classeur.xlsx [...] [--dashboard build_dashboard.py] [--detail]
        classeurs recalculés quelconques (variante « Comme dans Excel » seulement).
--dashboard : générateur du classeur, pour vérifier les couleurs (par défaut ../build_dashboard.py).
"""
import os
import sys
import warnings
from datetime import date, datetime, time, timedelta, timezone

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ICI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ICI)
import mirror  # noqa: E402
import model_def  # noqa: E402

warnings.simplefilter("ignore")

def _option(nom, defaut=None):
    """Valeur de l'option --nom dans la ligne de commande."""
    if nom in sys.argv:
        i = sys.argv.index(nom)
        if i + 1 < len(sys.argv):
            return sys.argv[i + 1]
    return defaut


S = _option("--racine")
DASHBOARD = _option("--dashboard", os.path.join(ICI, "..", "build_dashboard.py"))
CLASSEURS = [a for i, a in enumerate(sys.argv[1:], 1)
             if a.lower().endswith(".xlsx") and sys.argv[i - 1] not in ("--racine", "--dashboard")]
B12, B13 = f"{S}/build12", f"{S}/build13"
SV, SV2, SV4 = f"{S}/pbi/out/scen_variantes_1", f"{S}/pbi/out/scen_variantes_3", f"{S}/pbi/out/scen_variantes_4"
RV = f"{S}/pbi/review_calculs/scen_lo"
CAS = ["ext", "double", "tous", "casse", "joker", "partiel", "weekend", "debut_we", "retard", "termine", "surcharge",
       "termine_tous", "cap_zero"]
# scénarios de relecture : ligne sans date de début et ligne à 0 h (r_sans_debut), trous dans les listes Data
# (r_data_trou) ; r_heure et r_jokers_noms illustrent deux défauts du classeur (LISEZMOI, limites) et sont exclus
FICHIERS = ([f"{B12}/final.xlsx"] + [f"{B12}/sc_{c}.xlsx" for c in CAS] + [f"{B13}/final.xlsx"]
            + [f"{SV}/{n}.xlsx" for n in ("v_final_tous", "v_final_debut", "v_premier_mois", "v_termine_encours",
                                          "v_retard_tous_debut", "v_surcharge_debut", "v_weekend_debut",
                                          "v_joker_tous", "v_partiel_tous")]
            + [f"{SV2}/v_arrondi.xlsx", f"{SV4}/v_seuil.xlsx"]
            + [f"{RV}/r_sans_debut.xlsx", f"{RV}/r_data_trou.xlsx"])
TOUS, EN_COURS, EXCEL = "Tous les projets", "Projets en cours", mirror.EXCEL
LISSEE, DEBUT = "Lissée jusqu'au besoin", "Tout sur le mois de début"
# (classeur dont le miroir est construit, variante Périmètre, variante Répartition, classeur Excel de référence)
VARIANTES = [
    (f"{B12}/final.xlsx", EN_COURS, LISSEE, f"{B12}/final.xlsx"),
    (f"{B12}/final.xlsx", TOUS, LISSEE, f"{SV}/v_final_tous.xlsx"),
    (f"{B12}/final.xlsx", EN_COURS, DEBUT, f"{SV}/v_final_debut.xlsx"),
    (f"{B12}/final.xlsx", TOUS, DEBUT, f"{B12}/sc_tous.xlsx"),
    (f"{B12}/sc_tous.xlsx", EN_COURS, LISSEE, f"{B12}/final.xlsx"),
    (f"{B12}/sc_tous.xlsx", TOUS, LISSEE, f"{SV}/v_final_tous.xlsx"),
    (f"{B12}/sc_termine_tous.xlsx", EN_COURS, LISSEE, f"{SV}/v_termine_encours.xlsx"),
    (f"{SV}/v_termine_encours.xlsx", TOUS, LISSEE, f"{B12}/sc_termine_tous.xlsx"),
    (f"{B12}/sc_retard.xlsx", TOUS, DEBUT, f"{SV}/v_retard_tous_debut.xlsx"),
    (f"{B12}/sc_surcharge.xlsx", EN_COURS, DEBUT, f"{SV}/v_surcharge_debut.xlsx"),
    (f"{B12}/sc_weekend.xlsx", EN_COURS, DEBUT, f"{SV}/v_weekend_debut.xlsx"),
    (f"{B12}/sc_joker.xlsx", TOUS, LISSEE, f"{SV}/v_joker_tous.xlsx"),
    (f"{B12}/sc_partiel.xlsx", TOUS, LISSEE, f"{SV}/v_partiel_tous.xlsx"),
    (f"{B12}/sc_debut_we.xlsx", EN_COURS, DEBUT, f"{B12}/sc_debut_we.xlsx"),
]
TOL = 0.01
TOL2 = 1e-6     # valeurs arrondies à 0,01 h des deux côtés (ARRONDI d'Excel / fnArrondiExcel) : égalité exigée
DETAIL = "--detail" in sys.argv
COLONNES = {t["name"]: [c.name for c in t["cols"]] for t in model_def.TABLES}
TYPES = {t["name"]: {c.name: c.kind for c in t["cols"]} for t in model_def.TABLES}


def date_recalcul(chemin):
    """Date du jour utilisée par le recalcul (AUJOURDHUI) : date UTC de l'enregistrement du fichier recalculé."""
    return datetime.fromtimestamp(os.path.getmtime(chemin), tz=timezone.utc).date()


def col(i):
    """Colonne Excel de la colonne de grille i des onglets Dashboard et Dessinateurs (grille à partir de B)."""
    return get_column_letter(2 + i)


def d(x):
    return x.date() if isinstance(x, datetime) else x


def T(x):
    """SUPPRESPACE + espaces insécables : texte affiché normalisé."""
    if x is None:
        return ""
    return " ".join(str(x).replace(" ", " ").replace(" ", " ").split())


def cle_xl(x):
    """Clé d'affaire (Calculs EN / FF) telle qu'Excel Windows la calcule : chaque partie « projet | sous-projet |
    produit » avec espaces insécables remplacés puis SUPPRESPACE ; casse ignorée."""
    return " | ".join(" ".join(p.replace(" ", " ").split()) for p in str(x).split(" | ")).casefold()


def nombre(x):
    """Valeur numérique d'une cellule qu'openpyxl a convertie en heure à cause de son format de date."""
    if isinstance(x, time):
        return (x.hour * 3600 + x.minute * 60 + x.second + x.microsecond / 1e6) / 86400
    return x


def vide(x):
    return x is None or x == ""


class Bilan:
    def __init__(self, nom):
        self.nom, self.n, self.ecarts = nom, 0, []

    def _ko(self, label, attendu, obtenu):
        self.ecarts.append(f"{label} | Excel {attendu!r} | Power BI {obtenu!r}")

    def num(self, label, attendu, obtenu, tol=TOL):
        self.n += 1
        a, o = d(attendu), d(obtenu)
        if vide(a) and vide(o):
            return
        if isinstance(a, bool) or isinstance(o, bool) or not isinstance(a, (int, float)) \
                or not isinstance(o, (int, float)) or abs(a - o) > tol:
            self._ko(label, attendu, obtenu)

    def eq(self, label, attendu, obtenu):
        self.n += 1
        a, o = d(attendu), d(obtenu)
        if not (a == o or (vide(a) and vide(o))):
            self._ko(label, attendu, obtenu)

    def txt(self, label, attendu, obtenu, casse=False):
        self.n += 1
        a, o = T(attendu), T(obtenu)
        if casse:
            a, o = a.casefold(), o.casefold()
        if a != o:
            self._ko(label, attendu, obtenu)

    def vrai(self, label, condition, info=""):
        self.n += 1
        if not condition:
            self.ecarts.append(f"{label} {info}")


# =========================================================================================== intégrité des tables
RELATIONS_UN = [("Affaires", "Clé affaire"), ("Périmètre", "Périmètre"), ("Ressources", "Clé ressource"),
                ("Lignes", "Clé ligne"), ("Répartition", "Répartition"), ("Mois", "Mois")]


def type_ok(kind, v):
    if v is None:
        return True
    return {"text": isinstance(v, str), "int": isinstance(v, int) and not isinstance(v, bool),
            "num": isinstance(v, (int, float)) and not isinstance(v, bool),
            "date": isinstance(v, date) and not isinstance(v, datetime), "bool": isinstance(v, bool),
            "datetime": isinstance(v, datetime)}[kind]


def integrite(b, modele):
    """Colonnes déclarées produites, types compatibles avec l'étape Typée, clés des relations, colonnes de tri."""
    tables = modele["tables"]
    b.vrai("colonnes déclarées produites par le miroir", not modele["colonnes_manquantes"],
           str(modele["colonnes_manquantes"]))
    for t, lignes in tables.items():
        for r in lignes:
            for c, v in r.items():
                if not type_ok(TYPES[t][c], v):
                    b.vrai(f"type {t}[{c}]", False, f"valeur {v!r} pour le type {TYPES[t][c]}")
    for t, c in RELATIONS_UN:
        vals = [r[c] for r in tables[t]]
        # le moteur compare les textes sans tenir compte de la casse
        cles = [v.casefold() if isinstance(v, str) else v for v in vals]
        b.vrai(f"clé unique {t}[{c}]", len(cles) == len(set(cles)) and None not in cles)
    for rel in model_def.RELATIONSHIPS:
        frm, fcol, to, tcol = rel[:4]
        un = {r[tcol] for r in tables[to]}
        orphelins = [r[fcol] for r in tables[frm] if r[fcol] not in un]
        b.vrai(f"intégrité {frm}[{fcol}] -> {to}[{tcol}]", not orphelins, str(orphelins[:5]))
    for t in model_def.TABLES:
        for c in t["cols"]:
            if c.sort_by:
                paires = {}
                for r in tables[t["name"]]:
                    v = r[c.name].casefold() if isinstance(r[c.name], str) else r[c.name]
                    paires.setdefault(v, set()).add(r[c.sort_by])
                multiples = {k: v for k, v in paires.items() if len(v) > 1}
                b.vrai(f"tri {t['name']}[{c.name}] par [{c.sort_by}] : une valeur de tri par valeur", not multiples,
                       str(multiples)[:200])


# =========================================================================================== comparaison à Excel
class Excel:
    def __init__(self, chemin):
        self.wb = load_workbook(chemin, data_only=True)
        self.ca, self.db, self.ds = self.wb["Calculs"], self.wb["Dashboard"], self.wb["Dessinateurs"]

    def p(self, nom):
        feuille, ref = self.wb.defined_names[nom].attr_text.split("!")
        return self.wb[feuille][ref.replace("$", "")].value

    def mois(self):
        return [d(self.ca[f"AO{4 + i}"].value) for i in range(12)]


def comparer(b, modele, xl, ctx, mode_base):
    """Compare le miroir (contexte `ctx` : variantes des segments, 12 mois du Dashboard Excel) à Excel."""
    ca, db, ds = xl.ca, xl.db, xl.ds
    T_ = modele["tables"]
    mois = xl.mois()
    M = lambda nom, c=ctx: c.mesure(nom)            # noqa: E731
    Mm = lambda nom, m: ctx.avec(mois={m}).mesure(nom)  # noqa: E731

    # ---- mois affichés et charge par mois (Calculs AO:AQ)
    tous_mois = {m["Mois"] for m in T_["Mois"]}
    b.vrai("les 12 mois du Dashboard Excel existent dans la table Mois", all(m in tous_mois for m in mois))
    if mode_base:
        fen = sorted(m["Mois"] for m in T_["Mois"] if m["Période"] == mirror.PERIODE_EXCEL)
        b.eq("Mois[Période] = 12 mois du Dashboard Excel", mois, fen)
        b.eq("Mois[N° mois Excel]", list(range(1, 13)),
             [m["N° mois Excel"] for m in sorted(T_["Mois"], key=lambda x: x["Mois"]) if m["N° mois Excel"]])
    for i, m in enumerate(mois):
        b.num(f"charge du mois {m:%m/%Y} (AP{4 + i})", ca[f"AP{4 + i}"].value, Mm("Charge répartie (h)", m))
        b.num(f"affaires actives {m:%m/%Y} (AQ{4 + i})", ca[f"AQ{4 + i}"].value, Mm("Affaires actives", m))
        b.num(f"capacité (graphique) {m:%m/%Y}", ca.cell(14, 49 + i).value,
              Mm("Capacité projets de l'équipe (h/mois)", m))

    # ---- indicateurs Calculs AL4:AL40
    al = lambda r: ca[f"AL{r}"].value              # noqa: E731
    nums = {4: "Charge identifiée (h)", 5: "Nombre d'affaires", 6: "Nombre de chefs de projet", 7: "Nombre de projets",
            8: "Nombre de types", 9: "Réalisé (h)", 10: "Reste à faire (h)", 11: "Avancement moyen simple",
            12: "Avancement moyen pondéré", 15: "Charge répartie (h)",
            16: "Charge hors période ou sans date de début (h)", 17: "Pic de charge (h)", 20: "Affaires actives au pic",
            21: "ETP au pic", 22: "ETP-mois", 24: "Lignes terminées", 26: "Part du pic dans la période",
            33: "Capacité projets de l'équipe (h/mois)", 34: "Mois en surcharge", 36: "Pic en % de la capacité",
            37: "Affaires en retard", 38: "Affaires hors délai", 39: "Affaires en écart"}
    for r, nom in nums.items():
        b.num(f"AL{r} [{nom}]", al(r), M(nom), 0.0001 if r in (11, 12, 26, 36) else TOL)
    for r, nom in {14: "Premier mois affiché", 19: "Mois du pic", 25: "Dernier mois affiché",
                   35: "Premier mois en surcharge"}.items():
        b.eq(f"AL{r} [{nom}]", al(r), M(nom))
    b.num("AL18 rang du mois de pic", al(18), mois.index(M("Mois du pic")) + 1 if M("Mois du pic") in mois else None)
    top = mirror._top(ctx)
    b.num("AL23 position du plus gros reste", al(23), None if top is None else top["N° 1re ligne"] - 3)
    b.txt("AL27 [Période affichée]", al(27), M("Période affichée"))
    b.txt("AL40 [Rappel des réglages]", al(40), M("Rappel des réglages"))
    if mode_base:
        reg = T_["Réglages"][0]
        b.eq("AL13 Réglages[1er mois automatique]", al(13), reg["1er mois automatique"])
        b.eq("AL14 Réglages[1er mois affiché (Excel)]", al(14), reg["1er mois affiché (Excel)"])
        b.eq("AL25 Réglages[Dernier mois affiché (Excel)]", al(25), reg["Dernier mois affiché (Excel)"])
        b.txt("AL27 Réglages[Période affichée (Excel)]", al(27), reg["Période affichée (Excel)"])
        b.txt("AL40 Réglages[Rappel des réglages (Excel)]", al(40), reg["Rappel des réglages (Excel)"])
        b.num("AL33 Réglages[Capacité projets de l'équipe (Excel)]", al(33), reg["Capacité projets de l'équipe (Excel)"])
        # page Contrôle : chaque ligne lue dans Excel doit concorder avec la valeur recalculée par Power BI
        for lc in T_["Contrôle Excel"]:
            cc = ctx.avec(controle=lc)
            b.eq(f"Contrôle {lc['Cellule Excel']} {lc['Code']} : {cc.mesure('Excel (affiché)')} / "
                 f"{cc.mesure('Power BI (affiché)')}", "✔", cc.mesure("Statut du contrôle"))
            if lc["Code"] != "mois":
                b.txt(f"Contrôle {lc['Cellule Excel']} libellé", ca[f"AK{lc['Ordre']}"].value, lc["Contrôle"])
        b.eq("Contrôles en écart", 0, ctx.mesure("Contrôles en écart"))
        b.eq("nombre de lignes de Contrôle Excel", len(model_def.CONTROLES) + 12, len(T_["Contrôle Excel"]))

    # ---- lignes de Calculs (une par ligne de Charge)
    lignes = {x["N° ligne"]: x for x in T_["Lignes"] if x["Périmètre"] == ctx.p}
    cm = {}
    for x in T_["Charge mensuelle"]:
        if x["Répartition"] == ctx.r:
            cm[(x["Clé ligne"], x["Mois"])] = cm.get((x["Clé ligne"], x["Mois"]), 0) + x["Heures (h)"]
    toutes = {x["N° ligne"]: x for x in modele["intermediaires"]["DonnéesCharge"]}
    dans_excel = set()
    for r in range(4, 504):
        if ca[f"B{r}"].value != 1:
            b.vrai(f"ligne {r} non valide absente de DonnéesCharge", r not in toutes)
            continue
        x0 = toutes.get(r)
        b.vrai(f"ligne {r} valide présente dans DonnéesCharge", x0 is not None)
        if x0 is None:
            continue
        b.num(f"ligne {r} avancement (C)", ca[f"C{r}"].value, x0["Avancement"])
        b.num(f"ligne {r} heures toutes lignes (FG)", ca[f"FG{r}"].value, x0["Heures (h)"])
        b.eq(f"ligne {r} dessinateur et externe (FD)", ca[f"FD{r}"].value == 1, x0["Dessinateur et externe"])
        # clés d'affaire, casse ignorée. Le recalcul LibreOffice ne remplace pas les espaces insécables (son CAR(160)
        # n'est pas U+00A0) : SUBSTITUE(...;CAR(160);" ") puis SUPPRESPACE d'Excel Windows sont rejoués par cle_xl
        b.eq(f"ligne {r} clé affaire (FF)", cle_xl(ca[f"FF{r}"].value), x0["Clé affaire (texte)"])
        if ca[f"D{r}"].value != 1:
            continue
        dans_excel.add(r)
        x = lignes.get(r)
        b.vrai(f"ligne {r} présente dans Lignes ({ctx.p})", x is not None)
        if x is None:
            continue
        b.num(f"ligne {r} heures (E)", ca[f"E{r}"].value, x["Heures (h)"])
        b.txt(f"ligne {r} chef (F)", ca[f"F{r}"].value, x["Chef de projet"], casse=True)
        b.txt(f"ligne {r} projet (G)", ca[f"G{r}"].value, x["Projet"], casse=True)
        b.txt(f"ligne {r} type (H)", ca[f"H{r}"].value, x["Type"], casse=True)
        b.eq(f"ligne {r} début (I)", None if vide(ca[f"I{r}"].value) else ca[f"I{r}"].value, x["Début planifié"])
        b.num(f"ligne {r} rang (L)", ca[f"L{r}"].value, x["Rang ligne"])
        b.num(f"ligne {r} groupe couleur (V)", ca[f"V{r}"].value, x["Ordre groupe projet"])
        b.txt(f"ligne {r} ressource (BZ)", ca[f"BZ{r}"].value, x["Ressource de la ligne"], casse=True)
        b.eq(f"ligne {r} type de ressource (EM)", ca[f"EM{r}"].value, x["Type de ressource"])
        b.eq(f"ligne {r} clé affaire (EN)", cle_xl(ca[f"EN{r}"].value), x["Clé affaire (texte)"])
        for k, (rc, nom) in enumerate((("O", "Rang chef"), ("R", "Rang projet"), ("U", "Rang type"))):
            if ca[f"{'MPS'[k]}{r}"].value == 1:
                b.num(f"ligne {r} {nom} ({rc})", ca[f"{rc}{r}"].value, x[nom])
        tot = 0
        for i, m in enumerate(mois):
            v = cm.get((x["Clé ligne"], m), 0)
            tot += v
            b.num(f"ligne {r} heures {m:%m/%Y} ({get_column_letter(23 + i)})", ca.cell(r, 23 + i).value, v)
        b.num(f"ligne {r} heures sur la période (AI)", ca[f"AI{r}"].value, tot)
        b.num(f"ligne {r} [Charge répartie (h)] de la ligne", ca[f"AI{r}"].value,
              ctx.avec(visuel={"Clé ligne": x["Clé ligne"]}).mesure("Charge répartie (h)"))
    b.eq(f"lignes du périmètre ({ctx.p})", sorted(dans_excel), sorted(lignes))

    # ---- affaires (Calculs EO:FL, une par 1re ligne)
    aff = {a["N° 1re ligne"]: a for a in T_["Affaires"] if a["Périmètre"] == ctx.p}
    premieres = [r for r in range(4, 504) if ca[f"EO{r}"].value == 1]
    b.eq(f"affaires du périmètre ({ctx.p}) : 1res lignes", premieres, sorted(aff))
    for r in premieres:
        a = aff.get(r)
        if a is None:
            continue
        b.num(f"affaire {r} charge (EP)", ca[f"EP{r}"].value, a["Charge (h)"], TOL2)
        b.num(f"affaire {r} rang (EQ)", ca[f"EQ{r}"].value, a["N°"])
        b.num(f"affaire {r} internes (ER)", ca[f"ER{r}"].value, a["Personnes internes"])
        b.num(f"affaire {r} externes (ES)", ca[f"ES{r}"].value, a["Personnes externes"])
        b.eq(f"affaire {r} ressources (ET)", ca[f"ET{r}"].value, a["Ressources"])
        b.num(f"affaire {r} avancement (EU)", ca[f"EU{r}"].value, a["Avancement"], 1e-9)
        ew, ey = d(ca[f"EW{r}"].value), d(ca[f"EY{r}"].value)
        b.eq(f"affaire {r} début (EW)", None if ew in (2958465, date(9999, 12, 31)) else ew, a["Début planifié"])
        b.eq(f"affaire {r} besoin (EY)", None if ey in (0, None) else ey, a["Date de besoin"])
        b.num(f"affaire {r} heures toutes lignes (FH)", ca[f"FH{r}"].value, a["Heures toutes lignes (h)"], TOL2)
        b.num(f"affaire {r} réalisé (FI)", ca[f"FI{r}"].value, a["Réalisé toutes lignes (h)"], TOL2)
        b.num(f"affaire {r} reste (FJ)", ca[f"FJ{r}"].value, a["Reste à faire de l'affaire (h)"], TOL2)
        b.num(f"affaire {r} attendu (FK)", None if vide(ca[f"FK{r}"].value) else ca[f"FK{r}"].value,
              a["Avancement attendu"], 1e-9)
        b.num(f"affaire {r} retard (FL)", ca[f"FL{r}"].value, a["Retard"])
        b.eq(f"affaire {r} couleur d'alerte", "#B42318" if ca[f"FL{r}"].value else "#1F2937", a["Couleur d'alerte"])
        b.eq(f"affaire {r} [Couleur avancement]", "#B42318" if ca[f"FL{r}"].value else "#1F2937",
             ctx.avec(affaire=a["Clé affaire"]).mesure("Couleur avancement"))
        b.num(f"affaire {r} [Charge identifiée (h)] de l'affaire", ca[f"EP{r}"].value,
              ctx.avec(affaire=a["Clé affaire"]).mesure("Charge identifiée (h)"))

    # ---- tableau des affaires du Dashboard (lignes 60 à 159) et colonne de retard (Calculs BY)
    tab = sorted(aff.values(), key=lambda a: a["N°"])
    total = M("Charge identifiée (h)")
    for k in range(100):
        r = 60 + k
        if k >= len(tab):
            b.vrai(f"tableau ligne {r} vide", vide(db.cell(r, 2).value))
            continue
        a = tab[k]
        b.num(f"tableau {r} N°", db.cell(r, 2).value, a["N°"])
        b.txt(f"tableau {r} projet", db.cell(r, 4).value, a["Projet"], casse=True)
        b.txt(f"tableau {r} sous-projet", db.cell(r, 12).value, a["Sous-projet"])
        b.txt(f"tableau {r} produit", db.cell(r, 20).value, a["Produit"])
        chef = T(db.cell(r, 28).value) or "(non renseigné)"
        b.txt(f"tableau {r} chef", chef, a["Chef de projet"], casse=True)
        b.eq(f"tableau {r} ressources", db.cell(r, 35).value, a["Ressources"])
        b.num(f"tableau {r} charge", db.cell(r, 39).value, a["Charge (h)"], TOL2)
        b.num(f"tableau {r} % de la charge", db.cell(r, 44).value, a["% de la charge totale"], 1e-9)
        # la mesure ne reprend pas l'arrondi à 0,01 h de la charge de l'affaire (EP) : tolérance 0,01 point
        b.num(f"tableau {r} % ([Part de la charge (%)])", db.cell(r, 44).value,
              ctx.avec(affaire=a["Clé affaire"]).mesure("Part de la charge (%)"), 1e-4)
        b.num(f"tableau {r} avancement", db.cell(r, 50).value, a["Avancement"], 1e-9)
        deb = d(db.cell(r, 56).value)
        b.eq(f"tableau {r} début", deb, a["Début planifié"] or "à planifier")
        b.eq(f"tableau {r} début (affiché)", deb.strftime("%d/%m/%Y") if isinstance(deb, date) else deb,
             a["Début (affiché)"])
        b.num(f"tableau {r} retard (BY)", ca[f"BY{4 + k}"].value, a["Retard"])
        b.num(f"tableau {r} retard (BJ masquée)", db[f"BJ{r}"].value, a["Retard"])

    # ---- carte « plus gros reste à faire » (AV12:BA22)
    pg = lambda n: M(f"Plus gros reste à faire – {n}")   # noqa: E731
    b.num("plus gros reste (AV12)", db["AV12"].value, M("Plus gros reste à faire (h)"), TOL2)
    b.num("plus gros reste : total de l'affaire (AV13)", db["AV13"].value,
          pg("total de l'affaire (h)") if top else 0, TOL2)
    if top:
        for cell, n, casse in (("BA15", "projet", True), ("BA16", "sous-projet", False), ("BA17", "produit", False),
                               ("BA19", "ressources", False)):
            b.txt(f"plus gros reste : {n} ({cell})", db[cell].value, pg(n), casse)
        b.txt("plus gros reste : chef (BA18)", T(db["BA18"].value) or "(non renseigné)", pg("chef de projet"), True)
        for cell, n in (("BA20", "début planifié"), ("BA21", "date de besoin")):
            v = d(db[cell].value)
            b.eq(f"plus gros reste : {n} ({cell})", None if v == "—" else v, pg(n))
        b.num("plus gros reste : avancement (BA22)", db["BA22"].value, pg("avancement"), 1e-9)
        b.txt("plus gros reste : sous-titre", f"sur {mirror.format_fr(db['AV13'].value, '#,0')} h au total pour "
              "cette affaire", M("Sous-titre plus gros reste"))
    else:
        b.eq("plus gros reste : aucune affaire (BA15)", "—", db["BA15"].value)

    # ---- tuiles, cartes et textes du Dashboard
    for cell, nom in (("C12", "Nombre de chefs de projet"), ("R12", "Nombre de projets"), ("AG12", "Charge identifiée (h)"),
                      ("AG13", "ETP-mois"), ("AP15", "Réalisé (h)"), ("AP16", "Reste à faire (h)"),
                      ("C29", "Mois en surcharge"), ("C30", "Capacité projets de l'équipe (h/mois)"),
                      ("O29", "Pic de charge (h)"), ("AA29", "Affaires en retard"), ("AA30", "Affaires hors délai"),
                      ("AA31", "Affaires en écart"), ("AM29", "Nombre d'affaires"), ("AM30", "Nombre de projets"),
                      ("AY29", "Avancement moyen pondéré"), ("AY31", "Avancement moyen simple"),
                      ("U37", "Capacité projets de l'équipe (h/mois)"), ("AI45", "Charge identifiée (h)")):
        b.num(f"Dashboard {cell} [{nom}]", db[cell].value, M(nom), 0.0001 if cell.startswith("AY") else TOL)
    b.eq("Dashboard O30 [Mois du pic]", db["O30"].value, M("Mois du pic"))
    b.eq("Dashboard C38 [Premier mois affiché]", db["C38"].value, M("Premier mois affiché"))
    b.num("Dashboard Q38 [Charge hors période ou sans date de début (h)]", db["Q38"].value,
          M("Charge hors période ou sans date de début (h)"))
    hors = db["Q38"].value
    b.txt("Dashboard Q38 [Texte hors période]", f"Hors période ou sans date de début : "
          f"{mirror.format_fr(hors, '#,0')} h" if hors > 0 else "Toute la charge est visible sur la période",
          M("Texte hors période"))
    if al(34):        # branche choisie d'après Excel (mois en surcharge)
        c31 = d(db["C31"].value)
        b.eq("Dashboard C31 premier mois en surcharge", c31, M("Premier mois en surcharge"))
        b.txt("Dashboard C31 [Sous-titre surcharge]",
              "Dès " + mirror._mois_long(c31) if isinstance(c31, date) else repr(c31), M("Sous-titre surcharge"))
    else:
        pc = nombre(db["C31"].value)
        b.num("Dashboard C31 pic / capacité", pc, M("Pic en % de la capacité"), 0.0001)
        b.txt("Dashboard C31 [Sous-titre surcharge]", "" if not isinstance(pc, (int, float)) else
              f"Pic à {mirror.format_fr(mirror.round_dax(pc * 100, 0), '0')} % de la capacité",
              M("Sous-titre surcharge"))
    b.txt("Dashboard O30 [Sous-titre pic (mois)]", "en " + mirror._mois_long(d(db["O30"].value)),
          M("Sous-titre pic (mois)"))
    # recalcul LibreOffice en anglais : séparateur décimal « . » dans les textes
    b.txt("Dashboard O31 [Sous-titre pic (affaires)]", str(db["O31"].value).replace(".", ","),
          M("Sous-titre pic (affaires)"))
    b.txt("Dashboard AM31 [Sous-titre lignes terminées]", db["AM31"].value, M("Sous-titre lignes terminées"))
    b.txt("Dashboard B8 [Titre du périmètre]", db["B8"].value, M("Titre du périmètre"))
    b.txt("Dashboard C13 [Sous-titre chefs]", db["C13"].value, M("Sous-titre chefs"))
    b.txt("Dashboard C4 [Rappel des réglages]", db["C4"].value, M("Rappel des réglages"))
    b.txt("Dashboard AF4 [Période affichée]", db["AF4"].value, M("Période affichée"))
    b.txt("Dashboard AY31 [Sous-titre avancement]", "Moyenne simple par affaire : "
          f"{mirror.format_fr(mirror.round_dax(db['AY31'].value * 100, 0), '0')} %", M("Sous-titre avancement"))
    b.txt("Dashboard AA30 [Sous-titre hors délai]", f"dont {db['AA30'].value} hors délai (besoin dépassé)",
          M("Sous-titre hors délai"))
    b.txt("Dashboard AA31 [Sous-titre écart]", f"{db['AA31'].value} en retard de plus de 15 pts", M("Sous-titre écart"))
    b.txt("Dashboard AM30 [Sous-titre projets]", f"réparties sur {db['AM30'].value} projet(s)", M("Sous-titre projets"))
    b.txt("Dashboard AG13 [Sous-titre charge identifiée]", f"soit ≈ {mirror.format_fr(db['AG13'].value, '0.0')} "
          "ETP-mois", M("Sous-titre charge identifiée"))
    b.txt("Dashboard C30 [Sous-titre capacité]", f"Capacité projets : {mirror.format_fr(db['C30'].value, '#,0')} "
          "h/mois", M("Sous-titre capacité"))
    b.txt("Dashboard C38 [Sous-titre graphique mensuel]", "12 mois à partir de " + mirror._mois_long(d(db["C38"].value)),
          M("Sous-titre graphique mensuel"))

    # ---- listes des chefs, projets et types (Calculs BK:BU) : rangs des colonnes Rang chef / projet / type
    for (ck, cn, chh), n, colonne, rang, compte in ((("BK", "BL", "BM"), 8, "Chef de projet", "Rang chef",
                                                     "Nombre de chefs de projet"),
                                                    (("BO", "BP", "BQ"), 8, "Projet", "Rang projet",
                                                     "Nombre de projets"),
                                                    (("BS", "BT", "BU"), 4, "Type", "Rang type", "Nombre de types")):
        groupes = {}
        for x in lignes.values():
            groupes[x[rang]] = x[colonne]
        nb = M(compte)
        b.eq(f"{colonne} : nombre de groupes de rang", nb, len(groupes))
        cumul = 0
        for k in range(1, n + 1):
            nom_xl, h_xl = ca[f"{cn}{3 + k}"].value, ca[f"{chh}{3 + k}"].value
            if k > nb:
                b.vrai(f"{colonne} rang {k} vide", vide(nom_xl))
                continue
            if k == n and nb > n:
                b.vrai(f"{colonne} rang {k} : « + {nb - n + 1} autres »", T(nom_xl).startswith(f"+ {nb - n + 1} "),
                       repr(nom_xl))
                b.num(f"{colonne} rang {k} : charge des autres", h_xl, total - cumul)
                continue
            b.txt(f"{colonne} rang {k}", nom_xl, groupes.get(k), casse=True)
            h = ctx.avec(visuel={colonne: groupes.get(k)}).mesure("Charge identifiée (h)")
            cumul += h
            b.num(f"{colonne} rang {k} charge", h_xl, h)

    # ---- groupes de couleur des projets et histogramme mensuel (Calculs AS:BH)
    grp = {}
    for x in lignes.values():
        grp.setdefault(x["Ordre groupe projet"], (x["Groupe projet"], x["Couleur projet"]))
    for k in range(1, 10):
        r = 3 + k
        if vide(ca[f"AT{r}"].value):
            b.vrai(f"groupe projet {k} absent", k not in grp)
            b.eq(f"[Groupe {k} – libellé] vide", None, M(f"Groupe {k} – libellé"))
            b.num(f"[Groupe {k} – charge (h)] nulle", 0, M(f"Groupe {k} – charge (h)"))
            continue
        g = grp.get(k, (None, None))
        b.txt(f"groupe projet {k} libellé (AT{r})", ca[f"AT{r}"].value, g[0], casse=True)
        b.eq(f"groupe projet {k} couleur", "#" + model_def.SLOT_COLORS[k - 1], g[1])
        cg = ctx.avec(visuel={"Ordre groupe projet": k})
        b.num(f"groupe projet {k} charge (AU{r})", ca[f"AU{r}"].value, cg.mesure("Charge identifiée (h)"))
        b.num(f"groupe projet {k} part (AV{r})", ca[f"AV{r}"].value, cg.mesure("Part de la charge (%)"), 1e-9)
        b.txt(f"[Groupe {k} – libellé] (AT{r})", ca[f"AT{r}"].value, M(f"Groupe {k} – libellé"), casse=True)
        b.num(f"[Groupe {k} – charge (h)] (AU{r})", ca[f"AU{r}"].value, M(f"Groupe {k} – charge (h)"))
        for i, m in enumerate(mois):
            b.num(f"groupe projet {k} {m:%m/%Y}", ca.cell(r, 49 + i).value,
                  cg.avec(mois={m}).mesure("Charge répartie (h)"))
            b.num(f"[Groupe {k} – charge répartie (h)] {m:%m/%Y}", ca.cell(r, 49 + i).value,
                  Mm(f"Groupe {k} – charge répartie (h)", m))
    for i, m in enumerate(mois):
        b.num(f"total arrondi {m:%m/%Y} (ligne 13)", ca.cell(13, 49 + i).value,
              mirror.round_dax(Mm("Charge répartie (h)", m), 0), 0)

    # ---- onglet Dessinateurs : occupation de l'équipe (Calculs CS:DV, lignes 4 à 21)
    res = T_["Ressources"]
    par_pos = {x["Position Excel"]: x for x in res if x["Position Excel"] is not None}
    team_h = [get_column_letter(101 + i) for i in range(12)]
    team_o = [get_column_letter(113 + i) for i in range(12)]
    somme_equipe = [0.0] * 12
    for k in range(1, 17):
        r = 3 + k
        x = par_pos.get(k)
        if vide(ca[f"CT{r}"].value):
            b.vrai(f"équipe position {k} vide", x is None)
            continue
        b.vrai(f"équipe position {k} présente", x is not None, repr(ca[f"CT{r}"].value))
        if x is None:
            continue
        b.eq(f"équipe {k} nom", ca[f"CT{r}"].value, x["Ressource"])
        b.vrai(f"équipe {k} dans le tableau Excel", x["Dans le tableau Excel"] and x["Dans l'équipe"])
        b.num(f"équipe {k} capacité (CU)", ca[f"CU{r}"].value, x["Capacité (h/mois)"])
        b.num(f"équipe {k} charge fixe (CV)", ca[f"CV{r}"].value, x["Charge fixe (h/mois)"])
        b.eq(f"équipe {k} mention", ds.cell(72 + k, 10).value or "", x["Mention"])
        cr = ctx.avec(ressource=x["Clé ressource"])
        for i, m in enumerate(mois):
            h = cr.avec(mois={m}).mesure("Charge répartie (h)")
            somme_equipe[i] += h
            b.num(f"équipe {x['Ressource']} heures {m:%m/%Y}", ca[f"{team_h[i]}{r}"].value, h)
            b.num(f"équipe {x['Ressource']} occupation {m:%m/%Y}", ca[f"{team_o[i]}{r}"].value,
                  cr.avec(mois={m}).mesure("Occupation (%)"), 1e-9)
            b.num(f"équipe {x['Ressource']} occupation affichée {m:%m/%Y}", ds.cell(72 + k, 12 + 4 * i).value,
                  cr.avec(mois={m}).mesure("Occupation (%)"), 1e-9)
        b.num(f"équipe {x['Ressource']} occupation maximale (DU)", ca[f"DU{r}"].value,
              cr.mesure("Occupation maximale (%)"), 1e-9)
        b.num(f"équipe {x['Ressource']} charge sur la période (DV)", ca[f"DV{r}"].value, cr.mesure("Charge répartie (h)"))
    na = next((x for x in res if x["Type de ressource"] == "Non affecté"), None)
    for i, m in enumerate(mois):
        h_na = 0 if na is None else ctx.avec(ressource=na["Clé ressource"], mois={m}).mesure("Charge répartie (h)")
        b.num(f"charge non affectée {m:%m/%Y} (ligne 20)", ca[f"{team_h[i]}20"].value, h_na)
        autres = sum(ctx.avec(ressource=x["Clé ressource"], mois={m}).mesure("Charge répartie (h)") for x in res
                     if not x["Dans le tableau Excel"] and x["Type de ressource"] != "Non affecté")
        b.num(f"autres ressources {m:%m/%Y} (ligne 21)", ca[f"{team_h[i]}21"].value,
              max(0, mirror.round_dax(autres, 1)), 0.051)

    # ---- ressource sélectionnée (Calculs CC:CQ, DX:EH ; tuiles de l'onglet Dessinateurs)
    sel = T(xl.p("p_Dessinateur"))
    x = next((y for y in res if y["Ressource"].casefold() == sel.casefold()), None) if sel else None
    b.num("CD27 lignes à corriger", ca["CD27"].value, M("Lignes avec dessinateur et externe"))
    # lignes terminées / à corriger par chef (segment Chef de projet) : lignes valides du chef, somme = total
    toutes = [y for y in T_["Lignes"] if y["Périmètre"] == TOUS]
    for nom, colonne in (("Lignes terminées", "Terminée"), ("Lignes avec dessinateur et externe", "Dessinateur et externe")):
        par_chef = {}
        for y in toutes:
            par_chef.setdefault(y["Chef de projet"], 0)
            par_chef[y["Chef de projet"]] += 1 if y[colonne] else 0
        for chef, n in par_chef.items():
            b.num(f"[{nom}] avec le segment Chef = {chef}", n, ctx.avec(lignes={"Chef de projet": chef}).mesure(nom))
        b.num(f"[{nom}] : somme par chef = total", M(nom), sum(par_chef.values()))
    if sel and x is None:
        b.vrai(f"ressource sélectionnée {sel} absente des données (aucune ligne, hors listes)",
               ca["CD9"].value == 0 and ca["CD28"].value == 0)
    if x is not None:
        cr = ctx.avec(ressource=x["Clé ressource"])
        b.eq("CD25 type de la ressource", ca["CD25"].value, x["Type de ressource"])
        for cell, nom, tol in (("CD6", "Capacité de la ressource (h/mois)", TOL),
                               ("CD7", "Charge fixe de la ressource (h/mois)", TOL),
                               ("CD8", "Capacité projets de la ressource (h/mois)", TOL),
                               ("CD9", "Charge identifiée (h)", TOL), ("CD10", "Nombre d'affaires", TOL),
                               ("CD11", "Charge répartie (h)", TOL), ("CD12", "Charge au pic de la ressource (h)", TOL),
                               ("CD15", "Occupation au pic de la ressource (%)", 1e-9),
                               ("CD16", "Mois en surcharge de la ressource", TOL), ("CD17", "Dépassement cumulé (h)", TOL),
                               ("CD18", "Heures encore disponibles (h)", TOL), ("CD19", "Occupation (%)", 1e-9),
                               ("CD28", "Nombre de lignes", TOL)):
            b.num(f"sélection {sel} {cell} [{nom}]", ca[cell].value, cr.mesure(nom), tol)
        b.eq(f"sélection {sel} CD14 [Mois le plus chargé de la ressource]", ca["CD14"].value,
             cr.mesure("Mois le plus chargé de la ressource"))
        b.eq(f"sélection {sel} CD20 [Premier mois en surcharge de la ressource]",
             None if vide(ca["CD20"].value) else ca["CD20"].value, cr.mesure("Premier mois en surcharge de la ressource"))
        for i, m in enumerate(mois):
            r, cm_ = 4 + i, cr.avec(mois={m})
            for c_, nom in (("CH", "Charge fixe (h)"), ("CI", "Charge répartie (h)"),
                            ("CJ", "Charge totale de la ressource (h)"), ("CK", "Capacité de la ressource (h/mois)"),
                            ("CL", "Projets dans la capacité (h)"), ("CM", "Dépassement de capacité (h)"),
                            ("CN", "Marge de la ressource (h)"), ("CO", "Occupation (%)")):
                b.num(f"sélection {sel} {m:%m/%Y} {c_} [{nom}]", ca[f"{c_}{r}"].value, cm_.mesure(nom),
                      1e-9 if c_ == "CO" else TOL)
        # tuiles de l'onglet Dessinateurs
        b.txt("Dessinateurs titre", f"CHARGE MENSUELLE DE {sel.upper()} (HEURES)", cr.mesure("Titre de la ressource"))
        c14 = ds["C14"].value
        b.txt("Dessinateurs C14 [Sous-titre capacité de la ressource]",
              c14 if isinstance(c14, str) else f"soit {mirror.format_fr(c14, '#,0')} h/mois pour les projets",
              cr.mesure("Sous-titre capacité de la ressource"))
        b.txt("Dessinateurs C13 [Sous-titre charge fixe de la ressource]",
              f"dont {mirror.format_fr(ds['C13'].value, '#,0')} h de charge fixe",
              cr.mesure("Sous-titre charge fixe de la ressource"))
        b.txt("Dessinateurs O13 [Sous-titre affaires de la ressource]", f"sur {ds['O13'].value} affaire(s)",
              cr.mesure("Sous-titre affaires de la ressource"))
        b.txt("Dessinateurs O14 [Sous-titre charge sur la période]",
              f"dont {mirror.format_fr(ds['O14'].value, '#,0')} h sur la période",
              cr.mesure("Sous-titre charge sur la période"))
        aa12 = d(ds["AA12"].value)
        b.txt("Dessinateurs AA12 [Mois le plus chargé de la ressource (texte)]",
              aa12 if isinstance(aa12, str) else mirror._mois_long(aa12),
              cr.mesure("Mois le plus chargé de la ressource (texte)"))
        b.txt("Dessinateurs AA13 [Sous-titre pic de la ressource]",
              f"{mirror.format_fr(ds['AA13'].value, '#,0')} h au total (fixe + projets)",
              cr.mesure("Sous-titre pic de la ressource"))
        b.txt("Dessinateurs AA14 [Sous-titre occupation au pic]",
              f"{mirror.format_fr(mirror.round_dax((nombre(ds['AA14'].value) or 0) * 100, 0), '0')} % de la capacité",
              cr.mesure("Sous-titre occupation au pic"))
        b.num("Dessinateurs AM12 [Mois en surcharge de la ressource]", ds["AM12"].value,
              cr.mesure("Mois en surcharge de la ressource"))
        b.txt("Dessinateurs AM13 [Sous-titre dépassement cumulé]",
              f"Dépassement cumulé : {mirror.format_fr(ds['AM13'].value, '#,0')} h",
              cr.mesure("Sous-titre dépassement cumulé"))
        am14 = d(ds["AM14"].value)
        b.txt("Dessinateurs AM14 [Sous-titre surcharge de la ressource]",
              am14 if isinstance(am14, str) else "Dès " + mirror._mois_long(am14),
              cr.mesure("Sous-titre surcharge de la ressource"))
        b.num("Dessinateurs AY12 [Occupation (%)]", ds["AY12"].value, cr.mesure("Occupation (%)"), 1e-9)
        b.txt("Dessinateurs AY13 [Sous-titre heures disponibles]",
              f"{mirror.format_fr(ds['AY13'].value, '#,0')} h encore disponibles",
              cr.mesure("Sous-titre heures disponibles"))
        if not vide(ds["G7"].value):
            b.txt("Dessinateurs G7 [Origine de la ressource]", ds["G7"].value, cr.mesure("Origine de la ressource"))
        # lignes de la ressource (DX:EH), par rang de ligne ; message sans ligne (Dessinateurs, ligne 66)
        n_l = cr.mesure("Nombre de lignes")
        if n_l == 0:
            b.txt("Dessinateurs B66 [Sous-titre lignes de la ressource]", ds["B66"].value,
                  cr.mesure("Sous-titre lignes de la ressource"))
        else:
            b.vrai("[Sous-titre lignes de la ressource] annonce le nombre de lignes",
                   cr.mesure("Sous-titre lignes de la ressource").startswith(f"{n_l} ligne(s) dans le périmètre"))
        aff_lib = {a["Clé affaire"]: a["Affaire"] for a in T_["Affaires"]}
        lr = sorted((y for y in lignes.values() if y["Clé ressource"] == x["Clé ressource"]),
                    key=lambda y: y["Rang ligne"])
        for k in range(15):
            r = 4 + k
            if k >= len(lr):
                b.vrai(f"sélection {sel} ligne {k + 1} vide", vide(ca[f"DY{r}"].value))
                continue
            y = lr[k]
            b.num(f"sélection {sel} ligne {k + 1} position (DY)", ca[f"DY{r}"].value, y["N° ligne"] - 3)
            b.txt(f"sélection {sel} ligne {k + 1} projet", ca[f"DZ{r}"].value, y["Projet"], casse=True)
            b.txt(f"sélection {sel} ligne {k + 1} sous-projet", ca[f"EA{r}"].value, y["Sous-projet"])
            b.txt(f"sélection {sel} ligne {k + 1} produit", ca[f"EB{r}"].value, y["Produit"])
            b.txt(f"sélection {sel} ligne {k + 1} chef", T(ca[f"EC{r}"].value) or "(non renseigné)",
                  y["Chef de projet"], casse=True)
            b.num(f"sélection {sel} ligne {k + 1} charge", ca[f"ED{r}"].value, y["Heures (h)"])
            # cellule du tableau « Lignes de la ressource choisie » : ligne × affaire de la ligne
            cl = cr.avec(visuel={"N° ligne": y["N° ligne"]}, affaires_f={"Affaire": aff_lib[y["Clé affaire"]]})
            b.num(f"sélection {sel} ligne {k + 1} heures sur la période", ca[f"EE{r}"].value,
                  cl.mesure("Heures de la ligne sur la période (h)"))
            autres = [a for a in aff_lib.values() if a != aff_lib[y["Clé affaire"]]]
            if autres:
                b.eq(f"sélection {sel} ligne {k + 1} × autre affaire : combinaison vide", None,
                     cr.avec(visuel={"N° ligne": y["N° ligne"]}, affaires_f={"Affaire": autres[0]})
                     .mesure("Heures de la ligne sur la période (h)"))
            for c_, nom in (("EF", "Début planifié"), ("EG", "Date de besoin")):
                v = d(ca[f"{c_}{r}"].value)
                b.eq(f"sélection {sel} ligne {k + 1} {nom}", None if v == "—" else v, y[nom])
            v = d(ca[f"EF{r}"].value)
            b.txt(f"sélection {sel} ligne {k + 1} début affiché", v if v == "—" else f"{v:%d/%m/%Y}",
                  y["Début (affiché)"])
            b.num(f"sélection {sel} ligne {k + 1} avancement", ca[f"EH{r}"].value, y["Avancement"], 1e-9)


# =========================================================================================== heure de Paris
def test_heure_paris(b):
    """HeureParis (M) comparée au fuseau Europe/Paris de la bibliothèque standard, de 2024 à 2031, toutes les
    30 minutes autour des changements d'heure et toutes les 7 h 17 min le reste du temps."""
    try:
        from zoneinfo import ZoneInfo
    except ImportError:
        b.vrai("zoneinfo disponible", False)
        return
    paris = ZoneInfo("Europe/Paris")
    t = datetime(2024, 1, 1)
    while t < datetime(2032, 1, 1):
        attendu = t.replace(tzinfo=timezone.utc).astimezone(paris).replace(tzinfo=None)
        b.eq(f"HeureParis({t:%Y-%m-%d %H:%M} UTC)", attendu, mirror.HeureParis(t))
        proche = t.month in (3, 10) and t.day >= 24
        t += timedelta(minutes=30) if proche else timedelta(hours=7, minutes=17)


# =========================================================================================== programme principal
_CACHE = {}


def modele(chemin, jour):
    cle = (chemin, jour)
    if cle not in _CACHE:
        _CACHE[cle] = mirror.construire(chemin, jour, colonnes=COLONNES)
    return _CACHE[cle]


def couverture(b):
    """Chaque mesure du modèle a son miroir et s'évalue sans erreur."""
    noms = [m[0] for m in model_def.MEASURES]
    b.eq("mesures sans miroir", [], [n for n in noms if n not in mirror.MESURES])
    b.eq("miroirs sans mesure", [], [n for n in mirror.MESURES if n not in noms])
    # constantes du miroir identiques à celles des requêtes M
    m = model_def.m_expressions()
    b.vrai("seuil de retard du miroir = requête Affaires",
           f"fnArrondiExcel([Avancement] - [Avancement attendu], 4) < {mirror.SEUIL_RETARD}" in m["Affaires"])
    b.vrai("palette des projets du miroir = requête LignesPérimètres",
           "{" + ", ".join(f'"{c}"' for c in mirror.COULEURS) + "}" in m["LignesPérimètres"])
    # fonctions M (miroir) : règles d'Excel sur des cas limites absents des classeurs recalculés
    b.vrai("fnCléAffaire (M) remplace l'espace insécable avant SUPPRESPACE",
           'Text.Replace(fnTexteExcel(valeur), "#(00A0)", " ")' in m["fnCléAffaire"])
    for args, attendu in ((("VEGA", "Mise en  conformité ", " Passerelle"), "vega | mise en conformité | passerelle"),
                          (("A*", "b?", "c~d"), "a+ | b! | c-d"), (("ÉTÉ", None, 12.5), "été |  | 12,5"),
                          (("x" * 250, "y", "z"), "x" * 240)):
        b.eq(f"fnCléAffaire{args!r:.60}", attendu, mirror.fnCléAffaire(*args))
    b.eq("fnSupprEspaces", "a b c", mirror.fnSupprEspaces("  a   b c  "))
    b.eq("fnSupprEspaces garde l'espace insécable (SUPPRESPACE d'Excel)", "a b", mirror.fnSupprEspaces(" a b "))
    for x, n, attendu in ((0.125, 2, 0.13), (-0.125, 2, -0.13), (70.125, 2, 70.13), (2.675, 2, 2.68), (1.005, 2, 1.01),
                          (0.625, 2, 0.63), (2.5, 0, 3), (-2.5, 0, -3), (0.12345, 4, 0.1235), (-0.15005, 4, -0.1501)):
        b.eq(f"fnArrondiExcel({x}, {n}) = ARRONDI d'Excel", attendu, mirror.fnArrondiExcel(x, n))
        b.eq(f"ROUND DAX({x}, {n})", attendu, mirror.round_dax(x, n))
    import random
    alea = random.Random(20261020)
    feries = sorted({date(2026, 1, 1) + timedelta(days=alea.randrange(0, 900)) for _ in range(40)})
    feries_semaine = [j for j in feries if j.weekday() < 5]
    for _ in range(3000):
        a = date(2026, 1, 1) + timedelta(days=alea.randrange(0, 900))
        z = a + timedelta(days=alea.randrange(-40, 400))
        lo, hi = min(a, z), max(a, z)
        n = sum(1 for k in range((hi - lo).days + 1) if (lo + timedelta(days=k)).weekday() < 5
                and lo + timedelta(days=k) not in feries)
        b.eq(f"fnJoursOuvrés({a}, {z}) = NB.JOURS.OUVRES", n if z >= a else -n,
             mirror.fnJoursOuvrés(a, z, feries_semaine))
    b.eq("fnDateExcel(numéro de série)", date(2026, 10, 20), mirror.fnDateExcel(46315))
    b.eq("fnDateExcel(datetime)", date(2026, 10, 20), mirror.fnDateExcel(datetime(2026, 10, 20, 15, 30)))
    b.eq("fnDateExcel(texte)", None, mirror.fnDateExcel("20/10/2026"))
    b.eq("fnNExcel(VRAI, texte, date)", (1, 0, 46315), (mirror.fnNExcel(True), mirror.fnNExcel("12"),
                                                         mirror.fnNExcel(date(2026, 10, 20))))
    b.eq("fnTexteExcel(12,5 ; VRAI ; date)", ("12,5", "VRAI", "46315"),
         (mirror.fnTexteExcel(12.5), mirror.fnTexteExcel(True), mirror.fnTexteExcel(date(2026, 10, 20))))
    # couleurs du modèle = constantes de build_dashboard.py (lues dans le source, sans l'exécuter)
    import ast
    import re
    src = open(DASHBOARD, encoding="utf-8").read()
    slots = ast.literal_eval(re.search(r"^SLOT_COLORS = (\[.*?\])", src, re.M).group(1))
    b.eq("palette des projets = SLOT_COLORS de build_dashboard.py", ["#" + c for c in slots], mirror.COULEURS)
    b.eq("SLOT_COLORS de model_def", slots, model_def.SLOT_COLORS)
    regles = ast.literal_eval(re.search(r"^OCC_RULES = (\[.*?\])\n", src, re.M | re.S).group(1)
                              .replace("RED_T", '"B42318"').replace("GREEN_T", '"1E7B34"'))
    b.eq("couleurs d'occupation = OCC_RULES de build_dashboard.py",
         [("#" + bg, "#" + fg) for _, bg, fg in regles], [model_def.OCC_ROUGE, model_def.OCC_ORANGE, model_def.OCC_VERT])
    b.vrai("RED_T et GREEN_T de build_dashboard.py",
           'RED_T, GREEN_T = "B42318", "1E7B34"' in src and 'INK = "1F2937"' in src)
    b.eq("couleur de retard et encre", ("#B42318", "#1F2937"), (model_def.RETARD_COULEUR, model_def.ENCRE))


def executer(fichiers=FICHIERS, variantes=VARIANTES, generaux=True, afficher=True):
    """Exécute les contrôles ; renvoie (nombre de contrôles, liste des bilans)."""
    _CACHE.clear()
    bilans = []
    if generaux:
        b = Bilan("couverture des mesures et heure de Paris")
        couverture(b)
        test_heure_paris(b)
        bilans.append(b)
    for f in fichiers:
        jour = date_recalcul(f)
        mo = modele(f, jour)
        b = Bilan(f"{os.path.basename(os.path.dirname(f))}/{os.path.basename(f)} (Comme dans Excel, "
                  f"aujourd'hui = {jour:%d/%m/%Y}{', noms bi_* de repli' if mo['intermediaires']['noms_de_repli'] else ''})")
        integrite(b, mo)
        ctx = mirror.contexte_excel(mo)
        xl = Excel(f)
        b.eq("12 mois du contexte Excel", xl.mois(), sorted(ctx.mois))
        comparer(b, mo, xl, ctx, True)
        # les mesures de la page s'évaluent toutes dans ce contexte et pour chaque ressource
        for n in mirror.MESURES:
            if n not in ("Valeur Power BI", "Mois Power BI", "Texte Power BI", "Valeur Excel", "Écart",
                         "Excel (affiché)", "Power BI (affiché)", "Statut du contrôle"):
                ctx.mesure(n)
                for x in mo["tables"]["Ressources"]:
                    ctx.avec(ressource=x["Clé ressource"]).mesure(n)
        bilans.append(b)
    for base, perim, rep, cible in variantes:
        jour = date_recalcul(cible)
        mo = modele(base, jour)
        xl = Excel(cible)
        b = Bilan(f"variante {perim} / {rep} sur {os.path.basename(base)} = Excel {os.path.basename(cible)} "
                  f"(aujourd'hui = {jour:%d/%m/%Y})")
        ctx = mirror.Contexte(mo, perim, rep, mois=xl.mois())
        b.eq("périmètre Excel de la référence", T(xl.p("p_Perimetre")).casefold(),
             ctx.mesure("Périmètre en vigueur").casefold())
        comparer(b, mo, xl, ctx, False)
        bilans.append(b)
    if afficher:
        total, ecarts = 0, 0
        for b in bilans:
            total += b.n
            ecarts += len(b.ecarts)
            print(f"{'OK ' if not b.ecarts else 'KO '} {b.n:6d} contrôles, {len(b.ecarts):3d} écart(s)  {b.nom}")
            for e in b.ecarts[:None if DETAIL else 15]:
                print("      ÉCART", e)
        print(f"TOTAL : {total} contrôles sur {len(fichiers)} classeurs et {len(variantes)} variantes, "
              f"{ecarts} écart(s)")
        print("MIROIR = EXCEL : 0 écart" if ecarts == 0 else "MIROIR ≠ EXCEL")
    return bilans


def main():
    if CLASSEURS:
        bilans = executer(fichiers=CLASSEURS, variantes=[])
    elif S:
        bilans = executer()
    else:
        print(__doc__)
        return 2
    return 1 if any(b.ecarts for b in bilans) else 0


if __name__ == "__main__":
    sys.exit(main())

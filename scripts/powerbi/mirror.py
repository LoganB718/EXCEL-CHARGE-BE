# -*- coding: utf-8 -*-
"""Miroir Python du modèle sémantique : requêtes Power Query (mêmes étapes, mêmes noms) et mesures DAX.

Power BI, le moteur M et le moteur DAX ne s'exécutent pas ici : ce miroir les réimplémente pas à pas pour
comparer le résultat aux valeurs calculées par Excel (test_mirror.py).

    modele = construire("classeur.xlsx", aujourdhui=date(2026, 10, 19))
    ctx = Contexte(modele)                       # segments « Comme dans Excel », aucun autre filtre
    ctx.mesure("Charge identifiée (h)")

Tables : listes de dictionnaires {colonne: valeur} ; None = null (M) ou BLANK (DAX).
"""
import math
import warnings
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal

from openpyxl import load_workbook

warnings.simplefilter("ignore")

EXCEL = "Comme dans Excel"
PERIODE_EXCEL = "12 mois du Dashboard Excel"
# Plages des noms bi_* (classeurs générés avant l'ajout de ces noms : mêmes cellules)
REPLI_NOMS = {"bi_Charge": ("Charge", "A4:O503"), "bi_Data": ("Data", "A4:K203"),
              "bi_Indicateurs": ("Calculs", "AK4:AL40"), "bi_Mois": ("Calculs", "AO4:AP15")}
ERREUR = object()   # valeur d'erreur d'une cellule (#N/A...) : remplacée par null comme Table.ReplaceErrorValues
SEUIL_RETARD = -0.15  # écart d'avancement au-delà duquel une affaire est en retard (constante de la requête Affaires)


# ============================================================================================ lecture du classeur
class Classeur:
    """Équivalent de Excel.Workbook(..., null, true) : valeurs enregistrées des plages des noms définis."""

    def __init__(self, chemin):
        self.chemin = chemin
        self.wb = load_workbook(chemin, data_only=True)
        self.repli = []

    @staticmethod
    def _valeur(c):
        if c.data_type == "e":
            return ERREUR
        v = c.value
        if isinstance(v, time):          # heure seule : type time de M
            return v
        return v

    def fnLireNom(self, nom):
        dn = self.wb.defined_names.get(nom)
        if dn is None:
            if nom not in REPLI_NOMS:
                raise KeyError(f"Le classeur ne contient pas le nom défini {nom}")
            self.repli.append(nom)
            feuille, plage = REPLI_NOMS[nom]
        else:
            dests = list(dn.destinations)
            assert len(dests) == 1, dests
            feuille, plage = dests[0][0], dests[0][1].replace("$", "")
        cel = self.wb[feuille][plage]
        if not isinstance(cel, tuple):
            cel = ((cel,),)
        return [{f"Column{j + 1}": self._valeur(c) for j, c in enumerate(row)} for row in cel]

    def fnValeurNom(self, nom):
        plage = self.fnLireNom(nom)
        if not plage or not plage[0]:
            return None
        v = plage[0]["Column1"]
        return None if v is ERREUR else v


def _sans_erreurs(lignes, colonnes):
    """Table.SelectColumns(..., MissingField.UseNull) puis Table.ReplaceErrorValues(..., null)."""
    out = []
    for r in lignes:
        out.append({c: (None if r.get(c) is ERREUR else r.get(c)) for c in colonnes})
    return out


# ============================================================================================ fonctions M
def _est_nombre(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _est_date(v):
    return isinstance(v, date) and not isinstance(v, datetime)


def _g15(x):
    """Number.ToText(x, "G", "fr-FR") : 15 chiffres significatifs, virgule décimale."""
    if x == int(x) and abs(x) < 1e15:
        return str(int(x))
    return format(x, ".15G").replace(".", ",")


def fnSérieExcel(valeur):
    if _est_date(valeur):
        return (valeur - date(1899, 12, 30)).days
    if isinstance(valeur, datetime):
        return (valeur - datetime(1899, 12, 30)).total_seconds() / 86400
    return None


def fnNombreExcel(valeur):
    if _est_nombre(valeur):
        return valeur
    if isinstance(valeur, (date, datetime)):
        return fnSérieExcel(valeur)
    return None


def fnNExcel(valeur):
    nombre = fnNombreExcel(valeur)
    if isinstance(valeur, bool):
        return 1 if valeur else 0
    return 0 if nombre is None else nombre


def fnDateExcel(valeur):
    if _est_date(valeur):
        return valeur
    if isinstance(valeur, datetime):
        return valeur.date()
    if _est_nombre(valeur):
        try:   # Date.From(nombre) : date OLE Automation
            return (datetime(1899, 12, 30) + timedelta(days=valeur)).date()
        except (OverflowError, ValueError):
            return None
    return None


def fnTexteExcel(valeur):
    if valeur is None:
        return ""
    if isinstance(valeur, str):
        return valeur
    if isinstance(valeur, bool):
        return "VRAI" if valeur else "FAUX"
    if _est_nombre(valeur):
        return _g15(valeur)
    if isinstance(valeur, (date, datetime)):
        return _g15(fnSérieExcel(valeur))
    return str(valeur)


def fnSupprEspaces(texte):
    return " ".join(x for x in texte.split(" ") if x != "")


def fnCléAffaire(projet, sous_projet, produit):
    def partie(v):
        return fnSupprEspaces(fnTexteExcel(v).replace("\u00a0", " "))
    brute = partie(projet) + " | " + partie(sous_projet) + " | " + partie(produit)
    sans_jokers = brute.replace("~", "-").replace("*", "+").replace("?", "!")
    return sans_jokers[:240].lower()


def _round_dotnet(valeur, chiffres):
    """Math.Round(double, chiffres, MidpointRounding.AwayFromZero) de .NET (Number.Round de M)."""
    if abs(valeur) >= 1e16:
        return valeur
    p = 10.0 ** chiffres
    v = valeur * p
    entier = math.trunc(v)
    fraction = v - entier
    if abs(fraction) >= 0.5:
        entier += 1 if fraction > 0 else -1
    return entier / p


def fnArrondiExcel(nombre, decimales):
    puissance = 10.0 ** decimales
    decale = _round_dotnet(nombre * puissance, 6)
    return _round_dotnet(decale, 0) / puissance


def fnJoursOuvrés(debut, fin, jours_non_travailles):
    sens = -1 if fin < debut else 1
    premier, dernier = (debut, fin) if sens == 1 else (fin, debut)
    nb_jours = (dernier - premier).days + 1
    semaines, reste = divmod(nb_jours, 7)
    jour_premier = premier.weekday()          # Date.DayOfWeek(d, Day.Monday)
    ouvres_reste = sum(1 for k in range(jour_premier, jour_premier + reste) if k % 7 < 5)
    retires = sum(1 for j in jours_non_travailles if premier <= j <= dernier)
    return sens * (semaines * 5 + ouvres_reste - retires)


def _debut_mois(d):
    return date(d.year, d.month, 1)


def _fin_mois(d):
    suivant = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    return suivant - timedelta(days=1)


def _ajout_mois(d, n):
    k = d.year * 12 + d.month - 1 + n
    return date(k // 12, k % 12 + 1, 1)


NOMS_COURTS = ["janv.", "févr.", "mars", "avr.", "mai", "juin", "juil.", "août", "sept.", "oct.", "nov.", "déc."]
NOMS_LONGS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre",
              "novembre", "décembre"]
COULEURS = ["#2A78D6", "#EB6834", "#1BAF7A", "#EDA100", "#E87BA4", "#008300", "#6250D6", "#E34948", "#9CA3AF"]


def _mois_long(d):
    return f"{NOMS_LONGS[d.month - 1]} {d.year}"


def _mois_court(d):
    return "" if d is None else f"{NOMS_COURTS[d.month - 1]} {d.year}"


# ============================================================================================ requêtes M
def ParamètresExcel(cl):
    valeur_perimetre = cl.fnValeurNom("p_Perimetre")
    valeur_repartition = cl.fnValeurNom("p_Repartition")
    premier_mois = cl.fnValeurNom("p_PremierMois")
    heures_etp = cl.fnValeurNom("p_HeuresETP")
    charge_fixe = cl.fnValeurNom("p_ChargeFixe")
    premier_mois_date = fnDateExcel(premier_mois)
    return {
        "Périmètre Excel": fnTexteExcel(valeur_perimetre),
        "Répartition Excel": fnTexteExcel(valeur_repartition),
        "Tous les projets": fnTexteExcel(valeur_perimetre).lower() == "tous les projets",
        "Lissée": fnTexteExcel(valeur_repartition).lower() == "lissée jusqu'à la date de besoin",
        "Premier mois saisi": None if premier_mois_date is None else _debut_mois(premier_mois_date),
        "Capacité par défaut": fnNExcel(heures_etp),
        "Charge fixe par défaut": fnNExcel(charge_fixe),
        "Capacité (texte)": fnTexteExcel(heures_etp),
        "Charge fixe (texte)": fnTexteExcel(charge_fixe),
    }


def DonnéesData(cl):
    source = cl.fnLireNom("bi_Data")
    colonnes = [f"Column{i}" for i in range(1, 12)]
    sans_erreurs = _sans_erreurs(source, colonnes)
    for i, r in enumerate(sans_erreurs):
        r["N° ligne"] = 4 + i
    return sans_erreurs


def JoursNonTravaillés(donnees_data):
    dates = [fnDateExcel(r["Column11"]) for r in donnees_data]
    jours_de_semaine = [d for d in dates if d is not None and d.weekday() < 5]
    return sorted(set(jours_de_semaine))


def DonnéesCharge(cl):
    source = cl.fnLireNom("bi_Charge")
    colonnes = [f"Column{i}" for i in range(1, 16)]
    lettres = list("ABCDEFGHIJKLMNO")
    complete = _sans_erreurs(source, colonnes)
    renommee = [{lt: r[c] for c, lt in zip(colonnes, lettres)} for r in complete]
    for i, r in enumerate(renommee):
        r["N° ligne"] = 4 + i
    valides = [r for r in renommee if fnSupprEspaces(fnTexteExcel(r["B"])) != ""]
    out = []
    for r in valides:
        x = {"N° ligne": r["N° ligne"]}
        t = fnSupprEspaces(fnTexteExcel(r["A"]))
        x["Chef de projet (ligne)"] = "(non renseigné)" if t == "" else t
        x["Projet (ligne)"] = fnSupprEspaces(fnTexteExcel(r["B"]))
        x["Sous-projet"] = fnSupprEspaces(fnTexteExcel(r["C"]))
        x["Produit"] = fnSupprEspaces(fnTexteExcel(r["D"]))
        t = fnSupprEspaces(fnTexteExcel(r["E"]))
        x["Type (ligne)"] = "(non renseigné)" if t == "" else t
        x["Date de besoin"] = fnDateExcel(r["F"])
        x["Description"] = fnSupprEspaces(fnTexteExcel(r["G"]))
        x["Contraintes"] = fnSupprEspaces(fnTexteExcel(r["H"]))
        x["Autres informations"] = fnSupprEspaces(fnTexteExcel(r["I"]))
        x["Coordinateur"] = fnSupprEspaces(fnTexteExcel(r["J"]))
        x["Dessinateur"] = fnSupprEspaces(fnTexteExcel(r["K"]))
        x["Ressource externe"] = fnSupprEspaces(fnTexteExcel(r["L"]))
        x["Heures (h)"] = max(0, fnNExcel(r["M"]))
        x["Début planifié"] = fnDateExcel(r["N"])
        x["Avancement"] = min(1, max(0, fnNExcel(r["O"])))
        x["Ressource de la ligne"] = (x["Dessinateur"] if x["Dessinateur"] != "" else
                                      x["Ressource externe"] if x["Ressource externe"] != "" else "(non affecté)")
        x["Type de ressource"] = ("Interne" if x["Dessinateur"] != "" else
                                  "Externe" if x["Ressource externe"] != "" else "Non affecté")
        x["Dessinateur et externe"] = x["Dessinateur"] != "" and x["Ressource externe"] != ""
        x["Clé ressource (texte)"] = x["Ressource de la ligne"].lower()
        x["Clé affaire (texte)"] = fnCléAffaire(r["B"], r["C"], r["D"])
        x["Réalisé de la ligne (h)"] = x["Heures (h)"] * x["Avancement"]
        x["Terminée"] = x["Avancement"] >= 1
        out.append(x)
    return out


def RépartitionLignes(donnees_charge, jnt):
    source = [r for r in donnees_charge if r["Début planifié"] is not None and r["Heures (h)"] != 0]
    out = []
    for r in source:
        for code in ("Lissée", "Début"):
            debut, besoin, h = r["Début planifié"], r["Date de besoin"], r["Heures (h)"]
            fin = besoin if (code == "Lissée" and besoin is not None and besoin >= debut) else debut
            jours = fnJoursOuvrés(debut, fin, jnt)
            premier = _debut_mois(debut)
            nb_mois = (fin.year - premier.year) * 12 + fin.month - premier.month + 1
            for i in range(nb_mois):
                debut_mois = _ajout_mois(premier, i)
                fin_mois = _fin_mois(debut_mois)
                if jours == 0:
                    hm = h if debut_mois <= debut <= fin_mois else 0
                else:
                    hm = h * fnJoursOuvrés(max(debut, debut_mois), min(fin, fin_mois), jnt) / jours
                if hm > 0:
                    out.append({"N° ligne": r["N° ligne"], "Code répartition": code, "Mois": debut_mois,
                                "Heures du mois": hm})
    return out


def Périmètre(param):
    tous = param["Tous les projets"]
    applique = "Tous les projets" if tous else "Projets en cours"
    return [
        {"Périmètre": EXCEL, "Ordre": 1, "Périmètre appliqué": applique, "Code périmètre": "Tous" if tous else "En cours",
         "Description": "Réglage du classeur Excel : " + applique},
        {"Périmètre": "Projets en cours", "Ordre": 2, "Périmètre appliqué": "Projets en cours",
         "Code périmètre": "En cours", "Description": "Lignes dont l'avancement est inférieur à 100 %"},
        {"Périmètre": "Tous les projets", "Ordre": 3, "Périmètre appliqué": "Tous les projets",
         "Code périmètre": "Tous", "Description": "Toutes les lignes, y compris les lignes terminées"},
    ]


def Répartition(param):
    lissee = param["Lissée"]
    appliquee = "Lissée jusqu'à la date de besoin" if lissee else "Tout sur le mois de début"
    return [
        {"Répartition": EXCEL, "Ordre": 1, "Répartition appliquée": appliquee,
         "Code répartition": "Lissée" if lissee else "Début", "Description": "Réglage du classeur Excel : " + appliquee},
        {"Répartition": "Lissée jusqu'au besoin", "Ordre": 2, "Répartition appliquée": "Lissée jusqu'à la date de besoin",
         "Code répartition": "Lissée",
         "Description": "Prorata des jours ouvrés entre le début planifié et la date de besoin"},
        {"Répartition": "Tout sur le mois de début", "Ordre": 3, "Répartition appliquée": "Tout sur le mois de début",
         "Code répartition": "Début", "Description": "Toute la charge sur le mois du début planifié"},
    ]


def _grouper(lignes, colonne, nom):
    groupes = {}
    for r in lignes:
        groupes.setdefault(r[colonne].lower(), []).append(r)
    rows = []
    for cle, g in groupes.items():
        premiere = min(g, key=lambda x: x["N° ligne"])
        rows.append({"Clé groupe": cle, nom: premiere[colonne], "Première ligne": premiere["N° ligne"],
                     "Total": fnArrondiExcel(sum(x["Heures (h)"] for x in g), 2)})
    rows.sort(key=lambda x: (-x["Total"], x["Première ligne"]))
    for i, x in enumerate(rows):
        x["Rang"] = i + 1
    return rows


def _associer(lignes, colonne, nom, rang):
    groupes = {g["Clé groupe"]: g for g in _grouper(lignes, colonne, nom)}
    out = []
    for r in lignes:
        g = groupes[r[colonne].lower()]
        x = dict(r)
        x[nom], x[rang] = g[nom], g["Rang"]
        out.append(x)
    return out


def LignesPérimètres(donnees_charge, options):
    tables = []
    for option in options:
        ordre = option["Ordre"]
        lignes_option = (list(donnees_charge) if option["Code périmètre"] == "Tous"
                         else [r for r in donnees_charge if r["Avancement"] < 1])
        t = []
        for r in lignes_option:
            x = dict(r)
            x["Ordre périmètre"], x["Périmètre"] = ordre, option["Périmètre"]
            x["Clé ligne"] = ordre * 10000 + x["N° ligne"]
            t.append(x)
        premieres = {}
        for x in t:
            k = x["Clé affaire (texte)"]
            premieres[k] = min(premieres.get(k, x["N° ligne"]), x["N° ligne"])
        for x in t:
            x["Première ligne de l'affaire"] = premieres[x["Clé affaire (texte)"]]
            x["Clé affaire"] = ordre * 10000 + x["Première ligne de l'affaire"]
        t = _associer(t, "Chef de projet (ligne)", "Chef de projet", "Rang chef")
        t = _associer(t, "Type (ligne)", "Type", "Rang type")
        nb_projets = len(_grouper(t, "Projet (ligne)", "Projet"))
        t = _associer(t, "Projet (ligne)", "Projet", "Rang projet")
        for x in t:
            x["Ordre groupe projet"] = min(x["Rang projet"], 9) if nb_projets > 9 else x["Rang projet"]
            x["Groupe projet"] = (f"Autres projets ({nb_projets - 8})" if nb_projets > 9 and x["Ordre groupe projet"] == 9
                                  else x["Projet"])
            x["Couleur projet"] = COULEURS[x["Ordre groupe projet"] - 1]
        tri = sorted(t, key=lambda x: (-x["Heures (h)"], x["N° ligne"]))
        for i, x in enumerate(tri):
            x["Rang ligne"] = i + 1
        tables.append(sorted(tri, key=lambda x: x["N° ligne"]))
    return [x for tb in tables for x in tb]


def FenêtreExcel(lignes_perimetres, param, aujourdhui):
    lignes_excel = [r for r in lignes_perimetres if r["Ordre périmètre"] == 1]
    debuts = [r["Début planifié"] for r in lignes_excel if r["Début planifié"] is not None]
    debut_auto = _debut_mois(aujourdhui) if not debuts else _debut_mois(min(debuts))
    saisi = param["Premier mois saisi"]
    debut = saisi if saisi is not None else debut_auto
    return {"1er mois automatique": debut_auto, "1er mois affiché": debut, "Dernier mois affiché": _ajout_mois(debut, 11)}


def Ressources(donnees_data, param, donnees_charge):
    defauts = param

    def compter(col):
        return sum(1 for r in donnees_data if r[col] is not None)

    nb_internes, nb_externes = compter("Column5"), compter("Column8")

    def lire(lignes, nom, cap, fixe, type_ressource):
        out = []
        for r in lignes:
            c, f = fnNombreExcel(r[cap]), fnNombreExcel(r[fixe])
            out.append({"Ressource": fnSupprEspaces(fnTexteExcel(r[nom])), "Type de ressource": type_ressource,
                        "Capacité (h/mois)": defauts["Capacité par défaut"] if c is None else c,
                        "Charge fixe (h/mois)": defauts["Charge fixe par défaut"] if f is None else f})
        return out

    liste = (lire(donnees_data[:nb_internes], "Column5", "Column6", "Column7", "Interne")
             + lire(donnees_data[:nb_externes], "Column8", "Column9", "Column10", "Externe"))
    for i, r in enumerate(liste):
        r["Position Excel"] = i + 1
    non_vides = [r for r in liste if r["Ressource"] != ""]
    for r in non_vides:
        r["Clé ressource (texte)"] = r["Ressource"].lower()
    cles_liste = [r["Clé ressource (texte)"] for r in non_vides]
    uniques = [r for i, r in enumerate(non_vides) if cles_liste.index(r["Clé ressource (texte)"]) == i]
    equipe = [dict(r, **{"Dans l'équipe": True}) for r in uniques]
    cles_equipe = [r["Clé ressource (texte)"] for r in equipe]
    lignes_hors = [r for r in donnees_charge if r["Clé ressource (texte)"] != "(non affecté)"
                   and r["Clé ressource (texte)"] not in cles_equipe]
    premieres_hors = {}
    for r in lignes_hors:
        k = r["Clé ressource (texte)"]
        if k not in premieres_hors or r["N° ligne"] < premieres_hors[k]["N° ligne"]:
            premieres_hors[k] = r
    hors = [{"Ressource": r["Ressource de la ligne"], "Type de ressource": "Hors liste",
             "Capacité (h/mois)": defauts["Capacité par défaut"], "Charge fixe (h/mois)": defauts["Charge fixe par défaut"],
             "Position Excel": None, "Clé ressource (texte)": k, "Dans l'équipe": False}
            for k, r in sorted(premieres_hors.items(), key=lambda kv: kv[1]["N° ligne"])]
    non_affectee = ([] if "(non affecté)" in cles_equipe else
                    [{"Ressource": "(non affecté)", "Type de ressource": "Non affecté", "Capacité (h/mois)": None,
                      "Charge fixe (h/mois)": None, "Position Excel": None, "Clé ressource (texte)": "(non affecté)",
                      "Dans l'équipe": False}])
    toutes = equipe + hors + non_affectee
    for i, r in enumerate(toutes):
        r["Clé ressource"] = i + 1
        r["Ordre"] = r["Clé ressource"]
        r["Capacité projets (h/mois)"] = (None if r["Capacité (h/mois)"] is None
                                          else max(0, r["Capacité (h/mois)"] - r["Charge fixe (h/mois)"]))
        r["Dans le tableau Excel"] = r["Position Excel"] is not None and r["Position Excel"] <= 16
        r["Mention"] = ("ext." if r["Type de ressource"] == "Externe" else
                        "hors liste" if r["Type de ressource"] == "Hors liste" else "")
        r["Ressource (avec mention)"] = r["Ressource"] if r["Mention"] == "" else r["Ressource"] + " · " + r["Mention"]
    return toutes


def Lignes(lignes_perimetres, ressources):
    cles = {r["Clé ressource (texte)"]: r["Clé ressource"] for r in ressources}
    out = []
    for r in lignes_perimetres:
        x = dict(r)
        x["Clé ressource"] = cles.get(r["Clé ressource (texte)"])
        x["Reste de la ligne (h)"] = x["Heures (h)"] - x["Réalisé de la ligne (h)"]
        x["Sans date de début"] = x["Début planifié"] is None
        x["Début (affiché)"] = "—" if x["Début planifié"] is None else x["Début planifié"].strftime("%d/%m/%Y")
        out.append(x)
    return out


def _personnes(lignes):
    avec_personne = [r for r in lignes if r["Clé ressource (texte)"] != "(non affecté)"]
    par_personne = {}
    for r in avec_personne:
        k = r["Clé ressource (texte)"]
        if k not in par_personne or r["N° ligne"] < par_personne[k]["N° ligne"]:
            par_personne[k] = r
    internes = sum(1 for r in par_personne.values() if r["Type de ressource"] == "Interne")
    externes = sum(1 for r in par_personne.values() if r["Type de ressource"] == "Externe")
    non_affectees = sum(1 for r in lignes if r["Type de ressource"] == "Non affecté")
    if internes + externes == 0:
        texte = "à affecter"
    else:
        texte = ((f"{externes} ext." if internes == 0 else
                  f"{internes + externes}" + (f" dont {externes} ext." if externes > 0 else ""))
                 + (" + à affecter" if non_affectees > 0 else ""))
    return {"Ressources": texte, "Personnes internes": internes, "Personnes externes": externes,
            "À affecter": non_affectees > 0}


def Affaires(lignes_perimetres, donnees_charge, jnt, aujourdhui):
    toutes = {}
    for r in donnees_charge:
        toutes.setdefault(r["Clé affaire (texte)"], []).append(r)
    toutes = {k: {"Heures toutes lignes (h)": fnArrondiExcel(sum(x["Heures (h)"] for x in g), 2),
                  "Réalisé toutes lignes (h)": fnArrondiExcel(sum(x["Réalisé de la ligne (h)"] for x in g), 2),
                  "Moyenne des avancements": sum(x["Avancement"] for x in g) / len(g)} for k, g in toutes.items()}
    groupes = {}
    for r in lignes_perimetres:
        k = (r["Ordre périmètre"], r["Périmètre"], r["Clé affaire"], r["Clé affaire (texte)"])
        groupes.setdefault(k, []).append(r)
    totaux = {}
    for r in lignes_perimetres:
        totaux[r["Ordre périmètre"]] = totaux.get(r["Ordre périmètre"], 0) + r["Heures (h)"]
    rows = []
    for (ordre, perim, cle, cle_texte), g in groupes.items():
        premiere = min(g, key=lambda x: x["N° ligne"])
        debuts = [x["Début planifié"] for x in g if x["Début planifié"] is not None]
        besoins = [x["Date de besoin"] for x in g if x["Date de besoin"] is not None]
        a = {"Ordre périmètre": ordre, "Périmètre": perim, "Clé affaire": cle, "Clé affaire (texte)": cle_texte,
             "Charge (h)": fnArrondiExcel(sum(x["Heures (h)"] for x in g), 2),
             "Début planifié": min(debuts) if debuts else None, "Date de besoin": max(besoins) if besoins else None,
             "Lignes de l'affaire": len(g)}
        a.update(_personnes(g))
        a.update(toutes[cle_texte])
        a["Total du périmètre"] = totaux[ordre]
        a["Projet"], a["Sous-projet"], a["Produit"] = premiere["Projet"], premiere["Sous-projet"], premiere["Produit"]
        a["Chef de projet"] = premiere["Chef de projet"]
        a["N° 1re ligne"] = premiere["N° ligne"]
        a["Affaire"] = " · ".join(v if v != "" else "—" for v in (a["Projet"], a["Sous-projet"], a["Produit"]))
        a["% de la charge totale"] = 0 if a["Total du périmètre"] == 0 else a["Charge (h)"] / a["Total du périmètre"]
        a["Avancement"] = (a["Réalisé toutes lignes (h)"] / a["Heures toutes lignes (h)"]
                           if a["Heures toutes lignes (h)"] > 0 else a["Moyenne des avancements"])
        a["Reste à faire de l'affaire (h)"] = max(0, fnArrondiExcel(a["Heures toutes lignes (h)"] - a["Réalisé toutes lignes (h)"], 2))
        debut, besoin = a["Début planifié"], a["Date de besoin"]
        total = 0 if (debut is None or besoin is None or besoin <= debut) else fnJoursOuvrés(debut, besoin, jnt)
        if debut is None or besoin is None or besoin <= debut:
            attendu = None
        elif aujourdhui <= debut:
            attendu = 0
        elif aujourdhui > besoin:
            attendu = 1
        elif total == 0:
            attendu = None
        else:
            attendu = fnJoursOuvrés(debut, aujourdhui - timedelta(days=1), jnt) / total
        a["Avancement attendu"] = attendu
        if besoin is not None and besoin < aujourdhui and fnArrondiExcel(a["Avancement"], 4) < 1:
            a["Retard"] = 2
        elif attendu is None:
            a["Retard"] = 0
        elif fnArrondiExcel(a["Avancement"] - attendu, 4) < SEUIL_RETARD:
            a["Retard"] = 1
        else:
            a["Retard"] = 0
        a["Statut"] = ("Hors délai (besoin dépassé)" if a["Retard"] == 2 else
                       "En retard (écart > 15 pts)" if a["Retard"] == 1 else "Dans les temps")
        a["Couleur d'alerte"] = "#B42318" if a["Retard"] > 0 else "#1F2937"
        a["Début (affiché)"] = "à planifier" if debut is None else debut.strftime("%d/%m/%Y")
        rows.append(a)
    out = []
    for ordre in sorted({a["Ordre périmètre"] for a in rows}):
        part = sorted((a for a in rows if a["Ordre périmètre"] == ordre), key=lambda a: (-a["Charge (h)"], a["N° 1re ligne"]))
        for i, a in enumerate(part):
            a["N°"] = i + 1
        out += part
    return out


def ChargeMensuelle(lignes_perimetres, options_repartition, repartition_lignes):
    index = {}
    for r in repartition_lignes:
        index.setdefault((r["N° ligne"], r["Code répartition"]), []).append(r)
    out = []
    for l in lignes_perimetres:
        for o in options_repartition:
            for r in index.get((l["N° ligne"], o["Code répartition"]), []):
                out.append({"Clé ligne": l["Clé ligne"], "Clé affaire": l["Clé affaire"], "Répartition": o["Répartition"],
                            "Mois": r["Mois"], "Heures (h)": r["Heures du mois"]})
    return out


def Mois(fenetre, repartition_lignes):
    mois_repartis = {r["Mois"] for r in repartition_lignes}
    premier = min([fenetre["1er mois affiché"]] + list(mois_repartis))
    dernier = max([fenetre["Dernier mois affiché"]] + list(mois_repartis))
    nb = (dernier.year - premier.year) * 12 + dernier.month - premier.month + 1
    out = []
    f0 = fenetre["1er mois affiché"]
    for i in range(nb):
        m = _ajout_mois(premier, i)
        n = (m.year - f0.year) * 12 + m.month - f0.month + 1
        num = n if 1 <= n <= 12 else None
        periode = PERIODE_EXCEL if num is not None else ("Avant" if m < f0 else "Après")
        out.append({"Mois": m, "Libellé mois": _mois_court(m), "Mois (long)": _mois_long(m),
                    "Clé de tri": m.year * 100 + m.month, "Année": m.year, "Fin de mois": _fin_mois(m),
                    "N° mois Excel": num, "Période": periode,
                    "Ordre période": 1 if periode == "Avant" else 3 if periode == "Après" else 2})
    return out


def Réglages(param, fenetre, ressources, donnees_charge):
    # List.Sum d'une liste vide = null, remplacé par 0
    somme = sum(r["Capacité projets (h/mois)"] for r in ressources if r["Dans l'équipe"])
    rappel_capacite = param["Capacité (texte)"] + " h/pers. dont " + param["Charge fixe (texte)"] + " h hors projet"
    rappel = (param["Périmètre Excel"] + "  ·  "
              + ("lissage en jours ouvrés jusqu'au besoin" if param["Lissée"] else "tout sur le mois de début")
              + "  ·  " + rappel_capacite)
    return [{
        "Périmètre Excel": param["Périmètre Excel"], "Répartition Excel": param["Répartition Excel"],
        "Périmètre appliqué (Excel)": "Tous les projets" if param["Tous les projets"] else "Projets en cours",
        "Répartition appliquée (Excel)": ("Lissée jusqu'à la date de besoin" if param["Lissée"]
                                          else "Tout sur le mois de début"),
        "1er mois saisi": param["Premier mois saisi"], "1er mois automatique": fenetre["1er mois automatique"],
        "1er mois affiché (Excel)": fenetre["1er mois affiché"],
        "Dernier mois affiché (Excel)": fenetre["Dernier mois affiché"],
        "Période affichée (Excel)": ("Période " + _mois_long(fenetre["1er mois affiché"]) + " – "
                                     + _mois_long(fenetre["Dernier mois affiché"])),
        "Capacité par défaut (h/mois)": param["Capacité par défaut"],
        "Charge fixe par défaut (h/mois)": param["Charge fixe par défaut"],
        "Capacité projets par personne (Excel)": param["Capacité par défaut"] - param["Charge fixe par défaut"],
        "Capacité projets de l'équipe (Excel)": somme,
        "Rappel capacité": rappel_capacite, "Rappel des réglages (Excel)": rappel,
        "Lignes valides (Excel)": len(donnees_charge),
        "Lignes terminées (Excel)": sum(1 for r in donnees_charge if r["Terminée"]),
        "Lignes à corriger (Excel)": sum(1 for r in donnees_charge if r["Dessinateur et externe"]),
    }]


def Actualisation(heure):
    return [{"Actualisé le": heure, "Aujourd'hui": heure.date(),
             "Libellé": f"Données lues le {heure:%d/%m/%Y} à {heure:%H:%M} (heure de Paris)"}]


def ContrôleExcel(cl, controles):
    indicateurs = _sans_erreurs(cl.fnLireNom("bi_Indicateurs"), ["Column1", "Column2"])
    par_ligne = {4 + i: r for i, r in enumerate(indicateurs)}
    out = []
    for ligne, code, nature, tol in controles:
        r = par_ligne.get(ligne, {"Column1": None, "Column2": None})
        brute = r["Column2"]
        out.append({
            "Ordre": ligne, "Catégorie": "Indicateurs",
            "Contrôle": fnTexteExcel(r["Column1"]) or "Indicateur " + code, "Code": code,
            "Nature": nature, "Cellule Excel": f"Calculs!AL{ligne}",
            "Mois": fnDateExcel(brute) if nature == "Mois" else None,
            "Valeur lue": None if nature in ("Mois", "Texte") else fnNombreExcel(brute),
            "Valeur lue (texte)": (fnTexteExcel(brute) if nature == "Texte" else
                                     _mois_court(fnDateExcel(brute)) if nature == "Mois" else None),
            "Tolérance": tol})
    mois = _sans_erreurs(cl.fnLireNom("bi_Mois"), ["Column1", "Column2"])
    for i, r in enumerate(mois):
        n = i + 1
        out.append({"Ordre": 100 + n, "Catégorie": "Charge par mois",
                    "Contrôle": f"Mois {n:02d} : {_mois_court(fnDateExcel(r['Column1']))}", "Code": "mois",
                    "Nature": "Heures", "Cellule Excel": f"Calculs!AP{n + 3}", "Mois": fnDateExcel(r["Column1"]),
                    "Valeur lue": fnNombreExcel(r["Column2"]), "Valeur lue (texte)": None, "Tolérance": 0.01})
    return out


# ============================================================================================ assemblage
def construire(chemin, aujourdhui, heure=None, colonnes=None, controles=None):
    """Toutes les requêtes du modèle pour un classeur. `aujourdhui` = DateDuJour (date de l'actualisation à Paris).
    `colonnes` (facultatif) = {table: [colonnes du modèle]} : étapes Sélection/Typée des tables chargées."""
    cl = Classeur(chemin)
    heure = heure or datetime.combine(aujourdhui, time(12, 0))
    param = ParamètresExcel(cl)
    donnees_data = DonnéesData(cl)
    jnt = JoursNonTravaillés(donnees_data)
    donnees_charge = DonnéesCharge(cl)
    repartition_lignes = RépartitionLignes(donnees_charge, jnt)
    perimetre = Périmètre(param)
    repartition = Répartition(param)
    lignes_perimetres = LignesPérimètres(donnees_charge, perimetre)
    fenetre = FenêtreExcel(lignes_perimetres, param, aujourdhui)
    ressources = Ressources(donnees_data, param, donnees_charge)
    if controles is None:
        import model_def
        controles = model_def.CONTROLES
    brut = {
        "Périmètre": perimetre, "Répartition": repartition, "Ressources": ressources,
        "Lignes": Lignes(lignes_perimetres, ressources),
        "Affaires": Affaires(lignes_perimetres, donnees_charge, jnt, aujourdhui),
        "Charge mensuelle": ChargeMensuelle(lignes_perimetres, repartition, repartition_lignes),
        "Mois": Mois(fenetre, repartition_lignes),
        "Réglages": Réglages(param, fenetre, ressources, donnees_charge),
        "Actualisation": Actualisation(heure),
        "Contrôle Excel": ContrôleExcel(cl, controles),
        "Mesures": [],
    }
    intermediaires = {"ParamètresExcel": param, "DonnéesData": donnees_data, "JoursNonTravaillés": jnt,
                      "DonnéesCharge": donnees_charge, "RépartitionLignes": repartition_lignes,
                      "LignesPérimètres": lignes_perimetres, "FenêtreExcel": fenetre, "DateDuJour": aujourdhui,
                      "noms_de_repli": cl.repli}
    manquantes = {}
    if colonnes:
        for t, cols in colonnes.items():
            for r in brut[t]:
                absentes = [c for c in cols if c not in r]
                if absentes:
                    manquantes[t] = absentes
            brut[t] = [{c: r.get(c) for c in cols} for r in brut[t]]
    return {"tables": brut, "intermediaires": intermediaires, "colonnes_manquantes": manquantes}


def HeureParis(utc):
    """Requête HeureParis : heure UTC de l'actualisation convertie en heure de Paris (heure d'été du dernier dimanche
    de mars 01:00 UTC au dernier dimanche d'octobre 01:00 UTC)."""
    def dernier_dimanche(mois):
        fin = _fin_mois(date(utc.year, mois, 1))
        return fin - timedelta(days=(fin.weekday() + 1) % 7)       # Date.DayOfWeek(fin, Day.Sunday)
    debut_ete = datetime.combine(dernier_dimanche(3), time(1, 0))
    fin_ete = datetime.combine(dernier_dimanche(10), time(1, 0))
    decalage = 2 if debut_ete <= utc < fin_ete else 1
    return utc + timedelta(hours=decalage)


# ============================================================================================ mesures DAX
def round_dax(x, n):
    """ROUND de DAX (et d'Excel) : 15 chiffres significatifs, demi-unité arrondie en s'éloignant de zéro."""
    if x is None:
        x = 0
    d = Decimal(format(x, ".15g"))
    return float(d.quantize(Decimal(1).scaleb(-n), rounding=ROUND_HALF_UP))


def divide(a, b):
    if a is None or b is None or b == 0:
        return None
    return a / b


def _plus0(x):
    return 0 if x is None else x


def _sum(vals):
    vals = [v for v in vals if v is not None]
    return sum(vals) if vals else None


def _distinct(vals):
    s = {v for v in vals}
    return len(s) if s else None


ESPACE_MILLIERS = " "    # séparateur de milliers de la culture fr-FR


def format_fr(v, motif):
    """FORMAT ( v, motif, "fr-FR" ) pour les motifs utilisés par les mesures ; BLANK -> ""."""
    if v is None:
        return ""
    if motif == "0.0 %":
        v, motif, suffixe = v * 100, "0.0", " %"
    else:
        suffixe = ""
    dec = {"0": 0, "#,0": 0, "0.0": 1, "#,0.0": 1, "0.00": 2}[motif]
    q = Decimal(format(v, ".15g")).quantize(Decimal(1).scaleb(-dec), rounding=ROUND_HALF_UP)
    signe = "-" if q < 0 else ""
    entier, _, frac = f"{abs(q):f}".partition(".")
    if motif.startswith("#,"):
        groupes = []
        while len(entier) > 3:
            groupes.insert(0, entier[-3:])
            entier = entier[:-3]
        entier = ESPACE_MILLIERS.join([entier] + groupes)
    return signe + entier + ("," + frac if dec else "") + suffixe


def texte_dax(v):
    """Conversion implicite d'un nombre entier en texte par l'opérateur & de DAX."""
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return str(v)


class Contexte:
    """Contexte de filtre DAX simplifié.
    perimetre / repartition : valeur sélectionnée dans le segment (None = aucune -> « Comme dans Excel ») ;
    mois : ensemble des mois visibles (segment Période, axe d'un graphique) ou None = tous ;
    lignes : filtres de segments {colonne de Lignes: valeur} (Chef de projet, Projet, Type...) ;
    visuel : filtres de lignes / axes d'un visuel sur des colonnes de Lignes (retirés par ALLSELECTED) ;
    affaires_f : filtres {colonne d'Affaires: valeur} ; affaire : Clé affaire (ligne d'un tableau, extraction) ;
    ressource : Clé ressource ; ressources_f : filtres {colonne de Ressources: valeur} ;
    controle : ligne courante de la table Contrôle Excel."""

    CHAMPS = ("perimetre", "repartition", "mois", "lignes", "visuel", "affaires_f", "affaire", "ressource",
              "ressources_f", "controle")

    def __init__(self, modele, perimetre=None, repartition=None, mois=None, lignes=None, visuel=None,
                 affaires_f=None, affaire=None, ressource=None, ressources_f=None, controle=None):
        self.modele = modele
        self.M = modele["tables"]
        self.perimetre, self.repartition = perimetre, repartition
        self.mois = None if mois is None else set(mois)
        self.lignes_f, self.visuel = dict(lignes or {}), dict(visuel or {})
        self.affaires_f, self.affaire = dict(affaires_f or {}), affaire
        self.ressource, self.ressources_f = ressource, dict(ressources_f or {})
        self.controle = controle

    def avec(self, **kw):
        d = dict(perimetre=self.perimetre, repartition=self.repartition, mois=self.mois, lignes=self.lignes_f,
                 visuel=self.visuel, affaires_f=self.affaires_f, affaire=self.affaire, ressource=self.ressource,
                 ressources_f=self.ressources_f, controle=self.controle)
        d.update(kw)
        return Contexte(self.modele, **d)

    # SELECTEDVALUE ( 'Périmètre'[Périmètre], "Comme dans Excel" )
    @property
    def p(self):
        return self.perimetre or EXCEL

    @property
    def r(self):
        return self.repartition or EXCEL

    # --- propagation des filtres : Périmètre -> Affaires <-> Lignes <- Ressources ; Lignes, Répartition, Mois -> CM
    def ressources(self):
        R = self.M["Ressources"]
        if self.ressource is not None:
            R = [x for x in R if x["Clé ressource"] == self.ressource]
        for col, val in self.ressources_f.items():
            R = [x for x in R if x[col] == val]
        return R

    def _filtre_ressources(self):
        return self.ressource is not None or bool(self.ressources_f)

    def _lignes_directes(self):
        L = [x for x in self.M["Lignes"] if x["Périmètre"] == self.p]
        for col, val in list(self.lignes_f.items()) + list(self.visuel.items()):
            L = [x for x in L if x[col] == val]
        if self._filtre_ressources():
            cles = {x["Clé ressource"] for x in self.ressources()}
            L = [x for x in L if x["Clé ressource"] in cles]
        return L

    def affaires(self):
        A = [a for a in self.M["Affaires"] if a["Périmètre"] == self.p]
        if self.affaire is not None:
            A = [a for a in A if a["Clé affaire"] == self.affaire]
        for col, val in self.affaires_f.items():
            A = [a for a in A if a[col] == val]
        if self.lignes_f or self.visuel or self._filtre_ressources():     # relation bidirectionnelle Lignes -> Affaires
            cles = {x["Clé affaire"] for x in self._lignes_directes()}
            A = [a for a in A if a["Clé affaire"] in cles]
        return A

    def lignes(self):
        cles = {a["Clé affaire"] for a in self.affaires()}
        return [x for x in self._lignes_directes() if x["Clé affaire"] in cles]

    def cm(self):
        cles = {x["Clé ligne"] for x in self.lignes()}
        return [c for c in self.M["Charge mensuelle"] if c["Clé ligne"] in cles and c["Répartition"] == self.r
                and (self.mois is None or c["Mois"] in self.mois)]

    def mois_visibles(self):
        return [m["Mois"] for m in self.M["Mois"] if self.mois is None or m["Mois"] in self.mois]

    def reglages(self):
        return self.M["Réglages"][0]

    def ressource_unique(self):
        """HASONEVALUE ( Ressources[Clé ressource] ) : la ressource, sinon None."""
        R = self.ressources()
        return R[0] if len(R) == 1 else None

    def mesure(self, nom):
        return MESURES[nom](self)


MESURES = {}


def mesure(nom):
    def deco(f):
        MESURES[nom] = f
        return f
    return deco


def _lookup_mois(c, d, col="Mois (long)"):
    """LOOKUPVALUE ( Mois[col], Mois[Mois], d ) ; BLANK si d est vide ou absent."""
    if d is None:
        return None
    return next((m[col] for m in c.M["Mois"] if m["Mois"] == d), None)


def _t(v):
    """Texte d'une valeur dans une concaténation DAX (BLANK -> "")."""
    return "" if v is None else v


# ---- Charge
# ALLSELECTED () : le contexte du miroir ne contient que les filtres extérieurs au visuel
mesure("Affaire choisie")(lambda c: 1 if len({a["Affaire"] for a in c.affaires()}) == 1 else 0)
mesure("Message fiche affaire")(lambda c: "" if c.mesure("Affaire choisie") == 1 else
                                "Aucune affaire choisie : ouvrir cette fiche par clic droit sur une affaire > Extraire > "
                                "Fiche affaire.")
mesure("Charge identifiée (h)")(lambda c: _plus0(_sum(x["Heures (h)"] for x in c.lignes())))
mesure("Réalisé (h)")(lambda c: _plus0(_sum(x["Réalisé de la ligne (h)"] for x in c.lignes())))
mesure("Reste à faire (h)")(lambda c: c.mesure("Charge identifiée (h)") - c.mesure("Réalisé (h)"))
mesure("Capacité projets par personne (h/mois)")(lambda c: c.reglages()["Capacité projets par personne (Excel)"])


@mesure("ETP-mois")
def _etp(c):
    cap = c.mesure("Capacité projets par personne (h/mois)")
    return 0 if cap <= 0 else divide(c.mesure("Charge identifiée (h)"), cap)


# DISTINCTCOUNT : comparaison de textes insensible à la casse (classement du moteur)
mesure("Nombre de chefs de projet")(lambda c: _plus0(_distinct(x["Chef de projet"].lower() for x in c.lignes())))
mesure("Nombre de projets")(lambda c: _plus0(_distinct(x["Projet"].lower() for x in c.lignes())))
mesure("Nombre de types")(lambda c: _plus0(_distinct(x["Type"].lower() for x in c.lignes())))
mesure("Nombre de lignes")(lambda c: _plus0(len(c.lignes()) or None))
mesure("Part de la charge (%)")(lambda c: divide(c.mesure("Charge identifiée (h)"),
                                                 c.avec(visuel={}, affaire=None).mesure("Charge identifiée (h)")))

# ---- Affaires
mesure("Nombre d'affaires")(lambda c: _plus0(len(c.affaires()) or None))


@mesure("Avancement moyen simple")
def _avg_simple(c):
    A = c.affaires()
    return 0 if not A else divide(sum(a["Avancement"] for a in A), len(A))


@mesure("Avancement moyen pondéré")
def _avg_pond(c):
    A = c.affaires()
    h = sum(a["Heures toutes lignes (h)"] for a in A)
    return c.mesure("Avancement moyen simple") if h == 0 else divide(sum(a["Réalisé toutes lignes (h)"] for a in A), h)


mesure("Affaires en retard")(lambda c: sum(1 for a in c.affaires() if a["Retard"] > 0))
mesure("Affaires hors délai")(lambda c: sum(1 for a in c.affaires() if a["Retard"] == 2))
mesure("Affaires en écart")(lambda c: sum(1 for a in c.affaires() if a["Retard"] == 1))
mesure("Plus gros reste à faire (h)")(lambda c: max((a["Reste à faire de l'affaire (h)"] for a in c.affaires()),
                                                    default=0))


def _top(c):
    """TOPN ( 1, Affaires, reste DESC, N° 1re ligne ASC )."""
    A = sorted(c.affaires(), key=lambda a: (-a["Reste à faire de l'affaire (h)"], a["N° 1re ligne"]))
    return A[0] if A else None


for _nom, _col in (("affaire", "Affaire"), ("projet", "Projet"), ("sous-projet", "Sous-projet"),
                   ("produit", "Produit"), ("chef de projet", "Chef de projet"), ("ressources", "Ressources"),
                   ("total de l'affaire (h)", "Heures toutes lignes (h)"), ("début planifié", "Début planifié"),
                   ("date de besoin", "Date de besoin"), ("avancement", "Avancement"), ("retard", "Retard")):
    # CONCATENATEX sur une table vide = "" ; SUMX / MAXX sur une table vide = BLANK
    _vide = "" if _col in ("Affaire", "Projet", "Sous-projet", "Produit", "Chef de projet", "Ressources") else None
    MESURES[f"Plus gros reste à faire – {_nom}"] = (
        lambda col, vide: lambda c: (_top(c) or {}).get(col, vide))(_col, _vide)


@mesure("Couleur avancement")
def _couleur_av(c):
    valeurs = {a["Retard"] for a in c.affaires()}
    retard = next(iter(valeurs)) if len(valeurs) == 1 else 0     # SELECTEDVALUE ( Affaires[Retard], 0 )
    return "#B42318" if retard > 0 else "#1F2937"


# CALCULATE ( COUNTROWS ( Lignes ), REMOVEFILTERS ( Affaires ), 'Périmètre'[Périmètre] = "Tous les projets",
#             KEEPFILTERS ( condition ) ) + 0 : filtres sur les colonnes de Lignes et de Ressources conservés
def _lignes_tous(c, colonne):
    return sum(1 for x in c.avec(perimetre="Tous les projets", affaires_f={}, affaire=None).lignes() if x[colonne])


mesure("Lignes terminées")(lambda c: _lignes_tous(c, "Terminée"))
mesure("Lignes avec dessinateur et externe")(lambda c: _lignes_tous(c, "Dessinateur et externe"))

# ---- Mois et capacité
mesure("Charge répartie (h)")(lambda c: _plus0(_sum(x["Heures (h)"] for x in c.cm())))
mesure("Charge répartie sans les mois vides (h)")(lambda c: _sum(x["Heures (h)"] for x in c.cm()))
# IF ( NOT ( ISEMPTY ( Lignes ) ), [Charge répartie (h)] )
mesure("Heures de la ligne sur la période (h)")(lambda c: c.mesure("Charge répartie (h)") if c.lignes() else None)
mesure("Affaires actives")(lambda c: _plus0(_distinct(x["Clé affaire"] for x in c.cm())))
mesure("Charge hors période ou sans date de début (h)")(
    lambda c: max(0, round_dax(c.mesure("Charge identifiée (h)") - c.mesure("Charge répartie (h)"), 1)))
mesure("Capacité projets de l'équipe (h/mois)")(lambda c: c.reglages()["Capacité projets de l'équipe (Excel)"])


def _mois_surcharge(c):
    cap = round_dax(c.mesure("Capacité projets de l'équipe (h/mois)"), 0)
    return [m for m in c.mois_visibles() if round_dax(c.avec(mois={m}).mesure("Charge répartie (h)"), 0) > cap]


mesure("Mois en surcharge")(lambda c: len(_mois_surcharge(c)))
mesure("Premier mois en surcharge")(lambda c: min(_mois_surcharge(c), default=None))
mesure("Pic de charge (h)")(lambda c: max((c.avec(mois={m}).mesure("Charge répartie (h)") for m in c.mois_visibles()),
                                          default=0))


@mesure("Mois du pic")
def _mois_pic(c):
    pic = c.mesure("Pic de charge (h)")
    return min((m for m in c.mois_visibles() if c.avec(mois={m}).mesure("Charge répartie (h)") == pic), default=None)


@mesure("Affaires actives au pic")
def _actives_pic(c):
    m = c.mesure("Mois du pic")
    return 0 if m is None else c.avec(mois={m} & set(c.mois_visibles())).mesure("Affaires actives")


@mesure("ETP au pic")
def _etp_pic(c):
    cap = c.mesure("Capacité projets par personne (h/mois)")
    return 0 if cap <= 0 else divide(c.mesure("Pic de charge (h)"), cap)


@mesure("Pic en % de la capacité")
def _pic_cap(c):
    cap = c.mesure("Capacité projets de l'équipe (h/mois)")
    return None if cap <= 0 else divide(c.mesure("Pic de charge (h)"), cap)


@mesure("Part du pic dans la période")
def _part_pic(c):
    t = c.mesure("Charge répartie (h)")
    return 0 if t == 0 else divide(c.mesure("Pic de charge (h)"), t)


mesure("Premier mois affiché")(lambda c: min(c.mois_visibles(), default=None))
mesure("Dernier mois affiché")(lambda c: max(c.mois_visibles(), default=None))


# ---- Groupes de projets : CALCULATE ( ..., Lignes[Ordre groupe projet] = k ) remplace le filtre de cette colonne
def _groupe(c, k):
    return c.avec(visuel=dict(c.visuel, **{"Ordre groupe projet": k}))


def _libelle_groupe(c, k):
    valeurs = {x["Groupe projet"].casefold(): x["Groupe projet"] for x in _groupe(c, k).lignes()}
    return next(iter(valeurs.values())) if len(valeurs) == 1 else None      # SELECTEDVALUE


for _k in range(1, 10):
    MESURES[f"Groupe {_k} – libellé"] = (lambda k: lambda c: _libelle_groupe(c, k))(_k)
    MESURES[f"Groupe {_k} – charge (h)"] = (lambda k: lambda c: _groupe(c, k).mesure("Charge identifiée (h)"))(_k)
    MESURES[f"Groupe {_k} – charge répartie (h)"] = (
        lambda k: lambda c: _groupe(c, k).mesure("Charge répartie (h)"))(_k)


# ---- Ressource
def _res(c, col):
    r = c.ressource_unique()
    return None if r is None else r[col]


mesure("Capacité de la ressource (h/mois)")(lambda c: _res(c, "Capacité (h/mois)"))
mesure("Charge fixe de la ressource (h/mois)")(lambda c: _res(c, "Charge fixe (h/mois)"))
mesure("Capacité projets de la ressource (h/mois)")(lambda c: _res(c, "Capacité projets (h/mois)"))


def _par_mois_ressource(c, fonction, agreger=sum):
    """IF ( NOT ISBLANK ( _f ), <agrégat>X ( VALUES ( Mois[Mois] ), fonction(_f, _c, [Charge répartie (h)]) ) )."""
    f, cap = c.mesure("Charge fixe de la ressource (h/mois)"), c.mesure("Capacité de la ressource (h/mois)")
    if f is None:
        return None
    vals = [fonction(f, cap, c.avec(mois={m}).mesure("Charge répartie (h)")) for m in c.mois_visibles()]
    vals = [v for v in vals if v is not None]
    return agreger(vals) if vals else None


def _dans_capacite(f, cap, h):
    marge = cap - round_dax(f + h, 0)
    return max(0, min(h, cap - f)) if marge < 0 else h


def _occupation(f, cap, h):
    t = round_dax(f + h, 0)
    return (9.99 if t > 0 else 0) if cap <= 0 else divide(t, cap)


mesure("Charge fixe (h)")(lambda c: None if c.mesure("Charge fixe de la ressource (h/mois)") is None
                         else c.mesure("Charge fixe de la ressource (h/mois)") * len(c.mois_visibles()))
mesure("Charge totale de la ressource (h)")(lambda c: _par_mois_ressource(c, lambda f, cap, h: f + h))
mesure("Projets dans la capacité (h)")(lambda c: _par_mois_ressource(c, _dans_capacite))
mesure("Dépassement de capacité (h)")(lambda c: _par_mois_ressource(c, lambda f, cap, h: h - _dans_capacite(f, cap, h)))
mesure("Marge de la ressource (h)")(lambda c: _par_mois_ressource(c, lambda f, cap, h: cap - round_dax(f + h, 0)))
mesure("Dépassement cumulé (h)")(lambda c: _par_mois_ressource(
    c, lambda f, cap, h: -(cap - round_dax(f + h, 0)) if cap - round_dax(f + h, 0) < 0 else 0))
mesure("Heures encore disponibles (h)")(lambda c: _par_mois_ressource(
    c, lambda f, cap, h: (cap - round_dax(f + h, 0)) if cap - round_dax(f + h, 0) > 0 else 0))
mesure("Occupation (%)")(lambda c: _par_mois_ressource(c, _occupation, lambda v: sum(v) / len(v)))
mesure("Occupation maximale (%)")(lambda c: _par_mois_ressource(c, _occupation, max))


def _couleur_occ(c, i):
    o = c.mesure("Occupation (%)")
    if o is None:
        return None
    return (("#F8D7D7", "#B42318") if o > 1 else ("#FFF1CC", "#8A5A00") if o >= 0.9 else ("#E3F4E8", "#1E7B34"))[i]


mesure("Couleur de fond (occupation)")(lambda c: _couleur_occ(c, 0))
mesure("Couleur du texte (occupation)")(lambda c: _couleur_occ(c, 1))


def _mois_surcharge_res(c):
    f, cap = c.mesure("Charge fixe de la ressource (h/mois)"), c.mesure("Capacité de la ressource (h/mois)")
    if f is None:
        return None
    return [m for m in c.mois_visibles() if cap - round_dax(f + c.avec(mois={m}).mesure("Charge répartie (h)"), 0) < 0]


mesure("Mois en surcharge de la ressource")(lambda c: None if _mois_surcharge_res(c) is None
                                            else len(_mois_surcharge_res(c)))
mesure("Premier mois en surcharge de la ressource")(lambda c: None if _mois_surcharge_res(c) is None
                                                    else min(_mois_surcharge_res(c), default=None))
mesure("Charge au pic de la ressource (h)")(lambda c: _par_mois_ressource(c, lambda f, cap, h: f + h, max))


@mesure("Mois le plus chargé de la ressource")
def _mois_pic_res(c):
    f = c.mesure("Charge fixe de la ressource (h/mois)")
    if f is None:
        return None
    pic = c.mesure("Charge au pic de la ressource (h)")
    return min((m for m in c.mois_visibles() if f + c.avec(mois={m}).mesure("Charge répartie (h)") == pic), default=None)


@mesure("Occupation au pic de la ressource (%)")
def _occ_pic_res(c):
    m = c.mesure("Mois le plus chargé de la ressource")
    return None if m is None else c.avec(mois={m} & set(c.mois_visibles())).mesure("Occupation (%)")


# ---- Textes
mesure("Périmètre en vigueur")(lambda c: next((p["Périmètre appliqué"] for p in c.M["Périmètre"]
                                               if p["Périmètre"] == c.p), None))
mesure("Répartition en vigueur")(lambda c: next((r["Répartition appliquée"] for r in c.M["Répartition"]
                                                 if r["Répartition"] == c.r), None))
mesure("Rappel des réglages")(lambda c: _t(c.mesure("Périmètre en vigueur")) + "  ·  "
                              + ("tout sur le mois de début" if c.mesure("Répartition en vigueur")
                                 == "Tout sur le mois de début" else "lissage en jours ouvrés jusqu'au besoin")
                              + "  ·  " + c.reglages()["Rappel capacité"])
mesure("Données lues le")(lambda c: c.M["Actualisation"][0]["Libellé"])
mesure("Titre du périmètre")(lambda c: "TOUS LES PROJETS" if c.mesure("Périmètre en vigueur") == "Tous les projets"
                             else "PROJETS EN COURS")
mesure("Sous-titre chefs")(lambda c: "sur l'ensemble des projets" if c.mesure("Périmètre en vigueur")
                           == "Tous les projets" else "sur les projets en cours")
mesure("Sous-titre charge identifiée")(lambda c: "soit ≈ " + format_fr(c.mesure("ETP-mois"), "0.0") + " ETP-mois")
mesure("Sous-titre capacité")(lambda c: "Capacité projets : "
                              + format_fr(c.mesure("Capacité projets de l'équipe (h/mois)"), "#,0") + " h/mois")


@mesure("Sous-titre surcharge")
def _st_surcharge(c):
    m = c.mesure("Premier mois en surcharge")
    if m is not None:
        return "Dès " + _t(_lookup_mois(c, m))
    pc = c.mesure("Pic en % de la capacité")
    return "" if pc is None else "Pic à " + format_fr(round_dax(pc * 100, 0), "0") + " % de la capacité"


mesure("Sous-titre pic (mois)")(lambda c: None if c.mesure("Mois du pic") is None
                                 else "en " + _t(_lookup_mois(c, c.mesure("Mois du pic"))))


@mesure("Sous-titre pic (affaires)")
def _st_pic_aff(c):
    e = round_dax(c.mesure("ETP au pic"), 1)
    n = c.mesure("Affaires actives au pic")
    return (texte_dax(n) + " affaire(s) active(s)  ·  ≈ "
            + (format_fr(e, "0") if e == int(e) else format_fr(e, "0.0")) + " ETP")


mesure("Sous-titre hors délai")(lambda c: "dont " + texte_dax(c.mesure("Affaires hors délai"))
                                + " hors délai (besoin dépassé)")
mesure("Sous-titre écart")(lambda c: texte_dax(c.mesure("Affaires en écart")) + " en retard de plus de 15 pts")
mesure("Sous-titre projets")(lambda c: "réparties sur " + texte_dax(c.mesure("Nombre de projets")) + " projet(s)")
mesure("Sous-titre lignes terminées")(lambda c: (f"dont {texte_dax(c.mesure('Lignes terminées'))} ligne(s) terminée(s)"
                                                 if c.mesure("Périmètre en vigueur") == "Tous les projets" else
                                                 f"{texte_dax(c.mesure('Lignes terminées'))} ligne(s) terminée(s) "
                                                 "non comptée(s)"))
mesure("Sous-titre avancement")(lambda c: "Moyenne simple par affaire : "
                                + format_fr(round_dax(c.mesure("Avancement moyen simple") * 100, 0), "0") + " %")
mesure("Période affichée")(lambda c: None if not c.mois_visibles() else
                           "Période " + _t(_lookup_mois(c, min(c.mois_visibles()))) + " – "
                           + _t(_lookup_mois(c, max(c.mois_visibles()))))


@mesure("Texte hors période")
def _texte_hors(c):
    h = c.mesure("Charge hors période ou sans date de début (h)")
    return ("Hors période ou sans date de début : " + format_fr(h, "#,0") + " h" if h > 0
            else "Toute la charge est visible sur la période")


@mesure("Titre de la ressource")
def _titre_res(c):
    r = c.ressource_unique()
    return ("CHARGE MENSUELLE DE " + r["Ressource"].upper() + " (HEURES)" if r is not None
            else "CHARGE MENSUELLE (HEURES) — choisir une ressource")


@mesure("Sous-titre graphique mensuel")
def _st_graph(c):
    mv = c.mois_visibles()
    return (texte_dax(len(mv)) + " mois à partir de " + _t(_lookup_mois(c, min(mv)))) if mv else None


mesure("Sous-titre plus gros reste")(lambda c: "sur " + format_fr(_plus0(c.mesure(
    "Plus gros reste à faire – total de l'affaire (h)")), "#,0") + " h au total pour cette affaire")


@mesure("Origine de la ressource")
def _origine(c):
    types = {x["Type de ressource"] for x in c.ressources()}
    t = next(iter(types)) if len(types) == 1 else None            # SELECTEDVALUE
    return {"Interne": "Dessinateur interne (Data, colonne E)", "Externe": "Ressource externe (Data, colonne H)",
            "Hors liste": "Absent des listes Data : réglages par défaut",
            "Non affecté": "Lignes sans dessinateur ni ressource externe"}.get(t)


@mesure("Sous-titre charge fixe de la ressource")
def _st_fixe(c):
    f = c.mesure("Charge fixe de la ressource (h/mois)")
    return None if f is None else "dont " + format_fr(f, "#,0") + " h de charge fixe"


@mesure("Sous-titre capacité de la ressource")
def _st_cap_res(c):
    cap, f = c.mesure("Capacité de la ressource (h/mois)"), c.mesure("Charge fixe de la ressource (h/mois)")
    if cap is None:
        return None
    if f >= cap:
        return "Attention : charge fixe ≥ capacité"
    return "soit " + format_fr(c.mesure("Capacité projets de la ressource (h/mois)"), "#,0") + " h/mois pour les projets"


mesure("Sous-titre affaires de la ressource")(lambda c: "sur " + texte_dax(c.mesure("Nombre d'affaires"))
                                              + " affaire(s)")
mesure("Sous-titre charge sur la période")(lambda c: "dont " + format_fr(c.mesure("Charge répartie (h)"), "#,0")
                                           + " h sur la période")


@mesure("Sous-titre lignes de la ressource")
def _st_lignes_res(c):
    n = c.mesure("Nombre de lignes")
    if n == 0:
        return "Aucune affaire affectée à cette ressource dans le périmètre."
    return texte_dax(n) + " ligne(s) dans le périmètre · heures sur la période = mois sélectionnés (0 sans date de début)"


@mesure("Titre info-bulle projet")
def _titre_ib_projet(c):
    mv = c.mois_visibles()        # HASONEVALUE ( Mois[Mois] )
    if len(mv) != 1:
        return "Projet survolé · période affichée"
    return "Projet survolé · " + _t(_lookup_mois(c, mv[0]))


@mesure("Mois le plus chargé de la ressource (texte)")
def _mois_pic_res_txt(c):
    m = c.mesure("Mois le plus chargé de la ressource")
    if m is None or c.mesure("Charge répartie (h)") == 0:
        return "—"
    return _lookup_mois(c, m)


@mesure("Sous-titre pic de la ressource")
def _st_pic_res(c):
    p = c.mesure("Charge au pic de la ressource (h)")
    return None if p is None else format_fr(p, "#,0") + " h au total (fixe + projets)"


@mesure("Sous-titre occupation au pic")
def _st_occ_pic(c):
    o = c.mesure("Occupation au pic de la ressource (%)")
    return None if o is None else format_fr(round_dax(o * 100, 0), "0") + " % de la capacité"


@mesure("Sous-titre dépassement cumulé")
def _st_dep(c):
    d = c.mesure("Dépassement cumulé (h)")
    return None if d is None else "Dépassement cumulé : " + format_fr(d, "#,0") + " h"


@mesure("Sous-titre surcharge de la ressource")
def _st_sur_res(c):
    n = c.mesure("Mois en surcharge de la ressource")
    if n is None:
        return None
    m = c.mesure("Premier mois en surcharge de la ressource")
    return "Aucun dépassement prévu" if n == 0 else "Dès " + _t(_lookup_mois(c, m))


@mesure("Sous-titre heures disponibles")
def _st_dispo(c):
    d = c.mesure("Heures encore disponibles (h)")
    return None if d is None else format_fr(d, "#,0") + " h encore disponibles"


# ---- Contrôle (réglages Excel, 12 mois du Dashboard Excel, aucun autre filtre)
_CODES = {
    "total": "Charge identifiée (h)", "nbAff": "Nombre d'affaires", "nbChefs": "Nombre de chefs de projet",
    "nbProj": "Nombre de projets", "nbTypes": "Nombre de types", "realise": "Réalisé (h)", "reste": "Reste à faire (h)",
    "avgSimple": "Avancement moyen simple", "avgPond": "Avancement moyen pondéré", "periode": "Charge répartie (h)",
    "hors": "Charge hors période ou sans date de début (h)", "pic": "Pic de charge (h)",
    "picActives": "Affaires actives au pic", "picETP": "ETP au pic", "totalETP": "ETP-mois",
    "nbTermines": "Lignes terminées", "picPart": "Part du pic dans la période",
    "capEquipe": "Capacité projets de l'équipe (h/mois)", "nbSurcharge": "Mois en surcharge",
    "picCap": "Pic en % de la capacité", "nbRetard": "Affaires en retard", "nbDepasse": "Affaires hors délai",
    "nbEcart": "Affaires en écart",
}
_CODES_MOIS = {"start": "Premier mois affiché", "picMois": "Mois du pic", "finPeriode": "Dernier mois affiché",
               "premSurcharge": "Premier mois en surcharge"}
_CODES_TEXTE = {"periodeTxt": "Période affichée"}


def contexte_excel(modele, controle=None):
    """CALCULATE ( ..., REMOVEFILTERS (...), segments « Comme dans Excel », Mois[Période] = « 12 mois... » )."""
    fen = [m["Mois"] for m in modele["tables"]["Mois"] if m["Période"] == PERIODE_EXCEL]
    return Contexte(modele, EXCEL, EXCEL, mois=fen, controle=controle)


def _ligne_controle(c):
    if c.controle is None:
        raise ValueError("mesure de contrôle hors d'une ligne de la table Contrôle Excel")
    return c.controle


@mesure("Valeur Power BI")
def _valeur_pbi(c):
    lc = _ligne_controle(c)
    x = contexte_excel(c.modele, lc)
    code = lc["Code"]
    if code == "mois":
        return x.avec(mois={lc["Mois"]} & set(x.mois_visibles())).mesure("Charge répartie (h)")
    return x.mesure(_CODES[code]) if code in _CODES else None


@mesure("Mois Power BI")
def _mois_pbi(c):
    lc = _ligne_controle(c)
    x = contexte_excel(c.modele, lc)
    if lc["Code"] == "autoStart":
        return x.reglages()["1er mois automatique"]
    return x.mesure(_CODES_MOIS[lc["Code"]]) if lc["Code"] in _CODES_MOIS else None


@mesure("Texte Power BI")
def _texte_pbi(c):
    lc = _ligne_controle(c)
    x = contexte_excel(c.modele, lc)
    if lc["Code"] == "reglagesTxt":
        return x.reglages()["Rappel des réglages (Excel)"]
    return x.mesure(_CODES_TEXTE[lc["Code"]]) if lc["Code"] in _CODES_TEXTE else None


mesure("Valeur Excel")(lambda c: _ligne_controle(c)["Valeur lue"])


@mesure("Écart")
def _ecart(c):
    if _ligne_controle(c)["Nature"] not in ("Heures", "Nombre", "Décimal", "Pourcentage"):
        return None
    pbi, xl = c.mesure("Valeur Power BI"), c.mesure("Valeur Excel")
    return (pbi or 0) - (xl or 0) if (pbi is not None or xl is not None) else None


def _affiche(c, nature, v, texte):
    if nature in ("Texte", "Mois"):
        return texte
    if nature == "Pourcentage":
        return "" if v is None else format_fr(v, "0.0 %")
    return format_fr(v, {"Heures": "#,0.0", "Décimal": "0.00"}.get(nature, "0"))


mesure("Excel (affiché)")(lambda c: _affiche(c, _ligne_controle(c)["Nature"], c.mesure("Valeur Excel"),
                                             _ligne_controle(c)["Valeur lue (texte)"]))


@mesure("Power BI (affiché)")
def _pbi_affiche(c):
    n = _ligne_controle(c)["Nature"]
    if n == "Mois":
        m = c.mesure("Mois Power BI")
        return "" if m is None else _mois_court(m)          # FORMAT ( _m, "mmm yyyy", "fr-FR" )
    return _affiche(c, n, c.mesure("Valeur Power BI"), c.mesure("Texte Power BI"))


@mesure("Statut du contrôle")
def _statut(c):
    lc = _ligne_controle(c)
    n = lc["Nature"]
    if n == "Texte":
        ok = c.mesure("Texte Power BI") == lc["Valeur lue (texte)"]
    elif n == "Mois":
        ok = c.mesure("Mois Power BI") == lc["Mois"]
    else:
        pbi, xl = c.mesure("Valeur Power BI"), c.mesure("Valeur Excel")
        if pbi is None and xl is None:
            ok = True
        elif pbi is None or xl is None:
            ok = False
        else:
            ok = abs(pbi - xl) <= lc["Tolérance"]
    return "✔" if ok else "✘"


mesure("Contrôles en écart")(lambda c: sum(1 for lc in c.M["Contrôle Excel"]
                                           if c.avec(controle=lc).mesure("Statut du contrôle") == "✘"))
